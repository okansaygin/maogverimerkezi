# -*- coding: utf-8 -*-
"""
VERİ MERKEZİ - DOSYA TÜRÜ TESPİTİ (veri_merkezi_dosya_tespit.py)
=====================================================================
Tek "İçe Aktar" sihirbazı için: yüklenen Excel'in bir PERSONEL BİLGİLERİ
(aylık kimlik) dosyası mı, yoksa GÜNLÜK SAP MESAİ dosyası mı olduğunu
başlıklara bakarak anlar ve arayüzde gösterilecek sütun eşleştirme tablosunu
hazırlar.

MOTOR TEKRARLANMAZ: başlık kuralları doğrudan veri_merkezi_dogrulama
(kimlik) ve veri_merkezi_mesai_ice_aktarma (mesai) modüllerinden okunur -
bir kural orada değişirse burası da otomatik değişir.
"""

__version__ = "2026-09-25.1"

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import puantaj_engine as engine
import veri_merkezi_sema as sema
import veri_merkezi_dogrulama as dogrulama
import veri_merkezi_mesai_ice_aktarma as mie
from veri_merkezi_ice_aktarma import _pick_sheet

KIMLIK_ALAN_ETIKET = {
    "tc": "tc", "ad_soyad": "ad_soyad", "grup": "grup", "alt_ekip": "alt_ekip",
    "gorevi": "gorevi", "tip": "tip (normal/fazla)", "giris_tarihi": "giris_tarihi", "cikis_tarihi": "cikis_tarihi",
}
MESAI_ALANLAR = ["tc", "ad_soyad", "tarih", "durum", "toplam_saat"]


def _header_texts(ws, rows):
    texts = {}
    for r in rows:
        for c in range(1, ws.max_column + 1):
            if c not in texts:
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    texts[c] = " ".join(str(v).split())
    return texts


ALAN_ETIKET = {
    "tc": "TC Kimlik No", "ad_soyad": "Ad Soyad", "grup": "Grup", "alt_ekip": "Bağlı olduğu ekip",
    "gorevi": "Görevi", "giris_tarihi": "İşe giriş tarihi", "cikis_tarihi": "İşten çıkış tarihi",
    "tarih": "Tarih", "durum": "Çalışma durumu", "toplam_saat": "Çalışma saati (toplam)",
}


def sutun_harfi_coz(harf):
    """'b' / ' B ' -> 2. Geçersizse ValueError. Boş -> None."""
    from openpyxl.utils import column_index_from_string
    harf = (harf or "").strip().upper()
    if not harf:
        return None
    if not harf.isalpha() or len(harf) > 3:
        raise ValueError(f"'{harf}' geçerli bir sütun harfi değil")
    return column_index_from_string(harf)


def _ornek(ws, col, bas, bit):
    """Sütundaki ilk dolu veri hücresi (kullanıcının doğru sütunu seçtiğini görmesi için)."""
    if not col:
        return ""
    for r in range(bas, min(bit, ws.max_row) + 1):
        v = ws.cell(row=r, column=col).value
        if v not in (None, ""):
            if hasattr(v, "strftime"):
                return v.strftime("%d.%m.%Y")
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            return str(v)[:28]
    return ""


def dosya_turu_tespit(dosya_yolu, sheet_name=None, tur_zorla=None, override=None):
    """Döner: {"tur": 'kimlik'|'mesai'|None, "tahmin", "sayfa", "satir_sayisi", "son_veri_satiri",
    "eslesme": [{"alan", "etiket", "baslik", "sutun", "zorunlu", "bulundu", "elle", "ornek"}],
    "colmap", "bulunan", "beklenen", "boyut_kb"}

    tur_zorla : kullanıcı dosya türünü elle seçtiyse ('kimlik' | 'mesai').
    override  : {alan: sütun_no} - kullanıcının elle girdiği sütunlar, başlık eşleşmesine üstün gelir."""
    dosya_yolu = Path(dosya_yolu)
    wb = openpyxl.load_workbook(dosya_yolu, data_only=True, read_only=False)
    ws = _pick_sheet(wb, sheet_name)
    override = {k: v for k, v in (override or {}).items() if v}

    kimlik_map = dogrulama.build_column_map(ws)
    mesai_map = mie.build_column_map(ws)
    mesai_tamam = all(f in mesai_map for f in mie.REQUIRED_FIELDS)
    kimlik_tamam = all(f in kimlik_map for f in dogrulama.REQUIRED_FIELDS)
    if mesai_tamam and (not kimlik_tamam or len(mesai_map) >= len(kimlik_map)):
        tahmin = "mesai"
    elif kimlik_tamam:
        tahmin = "kimlik"
    else:
        tahmin = "mesai" if len(mesai_map) > len(kimlik_map) else "kimlik"

    secilen = tur_zorla or tahmin
    if secilen == "mesai":
        colmap, rows, alanlar, zorunlu = dict(mesai_map), (1, 2, 3), MESAI_ALANLAR, set(mie.REQUIRED_FIELDS)
        veri_bas = 2
    else:
        colmap, rows = dict(kimlik_map), (3, 2, 1, 4)
        alanlar = [k for k, _a, _n in dogrulama.FIELD_HEADER_RULES if k != "tip"]
        zorunlu = set(dogrulama.REQUIRED_FIELDS)
        veri_bas = 4
    otomatik = dict(colmap)
    colmap.update({k: v for k, v in override.items() if k in alanlar})
    tamam = all(f in colmap for f in zorunlu)
    # Tür belli: zorunlu sütunların hepsi (başlıktan ya da elle) bulunduysa
    tur = secilen if tamam else None

    texts = _header_texts(ws, rows)
    eslesme = []
    for alan in alanlar:
        c = colmap.get(alan)
        eslesme.append({
            "alan": alan,
            "etiket": ALAN_ETIKET.get(alan, alan),
            "baslik": texts.get(c, "—") if c else "—",
            "sutun": get_column_letter(c) if c else "",
            "zorunlu": alan in zorunlu,
            "bulundu": c is not None,
            "elle": alan in override and override[alan] != otomatik.get(alan),
            "ornek": _ornek(ws, c, veri_bas, veri_bas + 40),
        })

    son_satir = 0
    tc_col = colmap.get("tc")
    if tc_col:
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=tc_col).value not in (None, ""):
                son_satir = r
    return {
        "tur": tur,
        "tahmin": secilen,
        "sayfa": ws.title,
        "satir_sayisi": ws.max_row,
        "son_veri_satiri": son_satir,
        "eslesme": eslesme,
        "colmap": colmap,
        "override": override,
        "bulunan": sum(1 for e in eslesme if e["bulundu"]),
        "beklenen": len(eslesme),
        "boyut_kb": round(dosya_yolu.stat().st_size / 1024, 1),
    }


def personel_onizleme_sayilari(preview, db_path=None):
    """preview_import() sonucu için: kaç yeni / güncellenecek / değişmeyen kişi
    olduğunu veri setiyle karşılaştırarak hesaplar (commit_import ile AYNI
    kural: boş gelen alan mevcut değeri silmez). Çakışma kararları (mevcudu
    koru) hesaba katılmaz - bu yüzden 'güncellenecek' bir ÜST sınırdır."""
    kayitlar = preview.get("personel_kayitlari") or {}
    if not kayitlar:
        return {"yeni": 0, "guncellenecek": 0, "degismeyen": 0}
    conn = sema.get_connection(db_path)
    try:
        tcs = list(kayitlar)
        mevcut = {}
        for i in range(0, len(tcs), 900):
            parca = tcs[i:i + 900]
            for r in conn.execute(f"SELECT * FROM personnel WHERE tc IN ({','.join('?' * len(parca))})", parca):
                mevcut[r["tc"]] = dict(r)
    finally:
        conn.close()
    yeni = gunc = ayni = 0
    for tc, b in kayitlar.items():
        row = mevcut.get(tc)
        if row is None:
            yeni += 1
            continue
        alanlar = {
            "ad_soyad": b["ad_soyad"], "grup": b["grup"], "bagli_oldugu_ekip": b["alt_ekip"],
            "gorevi": b["gorevi"],
            "ise_giris_tarihi": b["giris_tarihi"].isoformat() if b["giris_tarihi"] else None,
            "isten_cikis_tarihi": b["cikis_tarihi"].isoformat() if b["cikis_tarihi"] else None,
        }
        if any(v and (row.get(k) or "") != v for k, v in alanlar.items()):
            gunc += 1
        else:
            ayni += 1
    return {"yeni": yeni, "guncellenecek": gunc, "degismeyen": ayni}


def mesai_onizleme_sayilari(preview, db_path=None):
    """preview_mesai_import() sonucu için yeni / güncellenecek / değişmeyen günlük kayıt."""
    kayitlar = preview.get("gunluk_kayitlar") or []
    if not kayitlar:
        return {"yeni": 0, "guncellenecek": 0, "degismeyen": 0, "kisi": 0}
    conn = sema.get_connection(db_path)
    try:
        tarihler = sorted({k["tarih"] for k in kayitlar})
        mevcut = {}
        for r in conn.execute("SELECT tc, tarih, normal_saat, fazla_saat, durum FROM gunluk_puantaj "
                              "WHERE tarih BETWEEN ? AND ?", (tarihler[0], tarihler[-1])):
            mevcut[(r["tc"], r["tarih"])] = r
    finally:
        conn.close()
    yeni = gunc = ayni = 0
    for k in kayitlar:
        m = mevcut.get((k["tc"], k["tarih"]))
        if m is None:
            yeni += 1
        elif (round(m["normal_saat"], 2) != k["normal_saat"] or round(m["fazla_saat"], 2) != k["fazla_saat"]
              or (m["durum"] or None) != k["durum"]):
            gunc += 1
        else:
            ayni += 1
    return {"yeni": yeni, "guncellenecek": gunc, "degismeyen": ayni,
            "kisi": len({k["tc"] for k in kayitlar})}
