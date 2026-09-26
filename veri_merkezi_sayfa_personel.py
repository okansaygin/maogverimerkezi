# -*- coding: utf-8 -*-
"""VERİ MERKEZİ - SAYFALAR: Personel Kayıtları, Personel Kartı (360°), Personel Düzenle."""

import datetime

from dash import dcc, html, Input, Output, State, no_update, callback, ctx
import dash_bootstrap_components as dbc
import openpyxl

import personnel_db as pdb
import team_report as report
import veri_merkezi_ekip_listesi as el
import veri_merkezi_hafta_kurali as hk
import veri_merkezi_ik as ik
import veri_merkezi_personel_karti as pk
import veri_merkezi_sorgu as sorgu
import veri_merkezi_verimlilik as vm
import veri_merkezi_ui as ui

SAYFA_BOYU = 25
ALAN_ETIKET = {"ad_soyad": "Ad soyad", "grup": "Grup", "bagli_oldugu_ekip": "Bağlı olduğu ekip", "gorevi": "Görevi",
               "ise_giris_tarihi": "İşe giriş tarihi", "isten_cikis_tarihi": "İşten çıkış tarihi", "notlar": "Notlar",
               "kayit": "Kayıt"}


def _durum_chip(p):
    return ui.chip("Ayrılmış", "gray") if (p.get("isten_cikis_tarihi") or "").strip() else ui.chip("Aktif", "green")


def _fc(etiket, n):
    return html.Span([html.Span(etiket), html.Span(ui.sayi(n), className="fc-n")], style={"display": "flex", "width": "100%"})


# ============================================================ PERSONEL LİSTESİ

def sayfa_liste():
    f = sorgu.get_personnel_facets()
    grup_opts = [{"label": _fc("(Grup belirtilmemiş)" if g["grup"] == "__bos__" else g["grup"], g["adet"]), "value": g["grup"]}
                 for g in f["gruplar"]]
    return ui.sayfa("Personel Kayıtları", "Kayıtlar / Personel", html.Div([
        html.Div([
            html.Div([html.P("Durum", className="vm-navsec", style={"color": "#5A6572", "padding": 0}),
                      dbc.Checklist(id="pl-durum", className="vm-facet", value=["aktif"], options=[
                          {"label": _fc("Aktif", f["aktif"]), "value": "aktif"},
                          {"label": _fc("Ayrılmış", f["ayrilmis"]), "value": "ayrilmis"}])], className="col", style={"gap": 0}),
            html.Div([html.P("Grup", className="vm-navsec", style={"color": "#5A6572", "padding": 0}),
                      dbc.Checklist(id="pl-grup", className="vm-facet", value=[], options=grup_opts)],
                     className="col", style={"gap": 0}),
            html.Div([html.P("Veri kalitesi", className="vm-navsec", style={"color": "#5A6572", "padding": 0}),
                      dbc.Checklist(id="pl-kalite", className="vm-facet", value=[], options=[
                          {"label": _fc("Açık uyarısı olanlar", f["uyarili"]), "value": "uyarili"},
                          {"label": _fc("Puantajı olmayanlar", f["puantajsiz"]), "value": "puantajsiz"}])],
                     className="col", style={"gap": 0}),
        ], className="vm-pane-l", style={"gap": "18px"}),
        html.Div([
            html.Div([
                html.Label([ui.ikon("search"), dcc.Input(id="pl-ara", type="text", debounce=True,
                                                         placeholder="Ad, TC, ekip veya görev")],
                           className="vm-searchbox", style={"width": "320px"}),
                html.Span(id="pl-sayi", className="grow", style={"fontSize": "13px", "color": "#5A6572"}),
                html.Button([ui.ikon("file-earmark-excel"), "Excel'e aktar"], id="pl-excel", className="vm-btn", n_clicks=0),
                dcc.Link([ui.ikon("person-plus"), "Yeni personel"], href="/personel-duzenle/yeni", className="vm-btn pri"),
            ], className="row", style={"gap": "10px"}),
            html.Div(id="pl-mesaj"),
            html.Section([
                html.Div(id="pl-tablo"),
                html.Div([html.Span(id="pl-sayfa-bilgi", className="grow"),
                          html.Button("Önceki", id="pl-onceki", className="vm-btn sm", n_clicks=0),
                          html.Button("Sonraki", id="pl-sonraki", className="vm-btn sm", n_clicks=0)],
                         className="row", style={"padding": "12px 16px", "borderTop": "1px solid #EFEBE4",
                                                 "fontSize": "13px", "color": "#5A6572", "gap": "10px"}),
            ], className="vm-card clip"),
        ], className="vm-pane-c"),
        dcc.Store(id="pl-sayfa", data=0),
    ], className="vm-split", style={"gridTemplateColumns": "250px minmax(0,1fr)"}), govde_cls=None)


def _liste_sorgu(durum, grup, kalite, ara, limit, offset):
    return sorgu.get_personnel_page(search=ara, durumlar=durum or None, gruplar=grup or None,
                                    sadece_uyarili="uyarili" in (kalite or []),
                                    sadece_puantajsiz="puantajsiz" in (kalite or []), limit=limit, offset=offset)


@callback(
    Output("pl-tablo", "children"),
    Output("pl-sayi", "children"),
    Output("pl-sayfa-bilgi", "children"),
    Output("pl-sayfa", "data"),
    Output("pl-onceki", "disabled"),
    Output("pl-sonraki", "disabled"),
    Input("pl-durum", "value"),
    Input("pl-grup", "value"),
    Input("pl-kalite", "value"),
    Input("pl-ara", "value"),
    Input("pl-onceki", "n_clicks"),
    Input("pl-sonraki", "n_clicks"),
    State("pl-sayfa", "data"),
)
def liste(durum, grup, kalite, ara, _o, _s, sayfa_no):
    sayfa_no = sayfa_no or 0
    if ctx.triggered_id == "pl-onceki":
        sayfa_no = max(0, sayfa_no - 1)
    elif ctx.triggered_id == "pl-sonraki":
        sayfa_no += 1
    else:
        sayfa_no = 0
    sonuc = _liste_sorgu(durum, grup, kalite, ara, SAYFA_BOYU, sayfa_no * SAYFA_BOYU)
    toplam = sonuc["toplam"]
    if sayfa_no * SAYFA_BOYU >= toplam and sayfa_no > 0:
        sayfa_no = max(0, (toplam - 1) // SAYFA_BOYU)
        sonuc = _liste_sorgu(durum, grup, kalite, ara, SAYFA_BOYU, sayfa_no * SAYFA_BOYU)
    uyarili = sorgu.get_open_alert_tcs()
    satirlar = []
    for p in sonuc["kayitlar"]:
        satirlar.append([
            html.Td(ui.kisi_link(p["tc"], p["ad_soyad"])),
            html.Td(ui.maske_tc(p["tc"]), className="mono", style={"fontSize": "13px"}),
            html.Td([p["grup"] or "—", html.Span(" › ", className="muted"), p["bagli_oldugu_ekip"] or "—"],
                    style={"fontSize": "13px"}),
            html.Td(p["gorevi"] or "—", style={"fontSize": "13px"}),
            html.Td(ui.tarih_tr(p["ise_giris_tarihi"]), className="mono small-12"),
            html.Td(_durum_chip(p)),
            html.Td(html.Div([html.Span(ui.tarih_tr(p["updated_at"]), className="mono small-12"),
                              html.Span((p["kaynak"] or "").replace("veri_merkezi:", "").replace("excel:", ""),
                                        className="vm-sub", style={"fontSize": "11px"})], className="col", style={"gap": 0})),
            html.Td(ui.ikon("exclamation-triangle-fill", className="c-orange", title="Açık uyarı var")
                    if p["tc"] in uyarili else None),
        ])
    tablo = ui.tablo(["Personel", "TC", "Grup › Alt ekip", "Görevi", "İşe giriş", "Durum", "Son güncelleme", ""],
                     satirlar) if satirlar else ui.bos("Bu filtrelerle eşleşen personel yok.")
    bas = sayfa_no * SAYFA_BOYU + 1 if toplam else 0
    bit = min(toplam, (sayfa_no + 1) * SAYFA_BOYU)
    return (tablo, [html.Strong(ui.sayi(toplam), style={"color": "#16202B"}), " kişi"],
            f"{ui.sayi(bas)}–{ui.sayi(bit)} / {ui.sayi(toplam)}", sayfa_no, sayfa_no == 0, bit >= toplam)


@callback(
    Output("pl-mesaj", "children"),
    Input("pl-excel", "n_clicks"),
    State("pl-durum", "value"), State("pl-grup", "value"), State("pl-kalite", "value"), State("pl-ara", "value"),
    prevent_initial_call=True,
)
def liste_excel(n, durum, grup, kalite, ara):
    if not n:
        return no_update
    kayitlar = _liste_sorgu(durum, grup, kalite, ara, 100000, 0)["kayitlar"]
    if not kayitlar:
        return ui.banner("Aktarılacak kayıt yok.", "warn")
    yol = report.OUTPUT_DIR / f"personel_listesi_{datetime.date.today():%Y%m%d}.xlsx"
    try:
        pdb.export_to_excel(kayitlar, yol)
    except PermissionError:
        return ui.banner("Dosya yazılamadı: aynı adlı Excel şu an açık olabilir.", "err")
    el.open_path(yol)
    return ui.banner(["Excel oluşturuldu: ", html.Code(str(yol))], "ok", "check2-circle")


# ============================================================ PERSONEL KARTI (360°)

def _ay_sec(tc):
    aylar = sorgu.get_person_months(tc)
    bu_ay = datetime.date.today().strftime("%Y-%m")
    if not aylar:
        aylar = [bu_ay]
    return aylar


def _gun_hucreleri(p, ay):
    bas, bit = sorgu.ay_araligi(ay)
    kayit = {r["tarih"]: r for r in sorgu.get_person_daily(p["tc"], bas, bit)}
    # Hafta tatili / pazar kuralı: ay sınırındaki haftalar da tam okunur (bkz. veri_merkezi_hafta_kurali).
    kural = hk.kural_uygula(hk.db_kayitlari(bas, bit, tc=p["tc"]))
    bugun = datetime.date.today()
    giris = (p.get("ise_giris_tarihi") or "")[:10]
    cikis = (p.get("isten_cikis_tarihi") or "")[:10]
    d1, d2 = datetime.date.fromisoformat(bas), datetime.date.fromisoformat(bit)
    hucreler = [html.Span(g, className="vm-cal-h", style={"textAlign": "left", "paddingLeft": "4px"}) for g in ui.GUN_KISA]
    hucreler += [html.Div(className="vm-pday bos") for _ in range(d1.weekday())]
    top = {"normal": 0.0, "fazla": 0.0, "kn": 0.0, "kf": 0.0, "katli_gun": 0, "izin": 0, "eksik": 0, "gun": 0}
    d = d1
    while d <= d2:
        iso = d.isoformat()
        r = kayit.get(iso)
        pazar = d.weekday() == 6
        cls, v, s, s_renk, nw, fw = "vm-pday", "", "", "#5A6572", 0, 0
        baslik_ek = ""
        if r:
            n, f = r["normal_saat"] or 0.0, r["fazla_saat"] or 0.0
            durum = (r["durum"] or "").strip()
            calisti = ui.tr_upper(durum) in ("", "ÇALIŞTI", "CALISTI")
            k_kat, katli = kural.get((p["tc"], iso), (None, False))
            ham_kat = vm.kategori(r["durum"], n + f)
            if n + f == 0 and not calisti:
                cls, s, s_renk = "vm-pday izin", ui.tr_upper(durum), "#5B3386"
                top["izin"] += 1
                if k_kat == "hafta_tatili" and ham_kat != "hafta_tatili":
                    s = "HT'YE DÖNDÜ"
                    baslik_ek = " · pazar çalışıldığı için hafta tatili sayıldı"
                elif k_kat == "ucretsiz_izin" and ham_kat != "ucretsiz_izin":
                    s = "ÜCRETSİZ SAYILDI"
                    baslik_ek = " · haftada mazeretsiz izin var, hafta tatili hak edilmedi"
            else:
                top["normal"] += n
                top["fazla"] += f
                if katli:
                    top["kn"] += n
                    top["kf"] += f
                    top["katli_gun"] += 1
                top["gun"] += 1
                v = ui.saat(n + f)
                nw, fw = n / 13 * 100, f / 13 * 100
                if pazar and katli:
                    s, s_renk = "PAZAR ×2,5", "#9A3A0B"
                    baslik_ek = " · normal + fazla mesai ×2,5"
                elif pazar:
                    s, s_renk = "PAZAR ×1", "#5A6572"
                    baslik_ek = (" · haftada mazeretsiz izin ya da kaydırılmış hafta tatili var; pazar normal "
                                 "iş günü sayıldı")
                elif f:
                    s, s_renk = f"FM {ui.saat(f)}", "#9A3A0B"
                else:
                    s, s_renk = (ui.tr_upper(durum) or "ÇALIŞTI"), "#2F6B3A"
        elif d >= bugun:
            cls, baslik_ek = "vm-pday gelecek", ""
        elif (giris and iso < giris) or (cikis and iso > cikis):
            cls, s, baslik_ek = "vm-pday gelecek", "KAYIT DIŞI", ""
        else:
            cls, s, s_renk, baslik_ek = "vm-pday yok", "VERİ YOK", "#9A3A0B", ""
            top["eksik"] += 1
        hucreler.append(html.Div([
            html.Div([html.Span(str(d.day), className="vm-pday-n"), html.Span(v, className="vm-pday-v")], className="vm-pday-top"),
            ui.bar([(nw, "bg-navy"), (fw, "bg-orange")], "thin") if r and v else html.Div(style={"height": "6px"}),
            html.Span(s, className="vm-pday-s", style={"color": s_renk}),
        ], className=cls, title=f"{d:%d.%m.%Y}" + (f" · {r['durum'] or ''}" if r else "") + baslik_ek))
        d += datetime.timedelta(days=1)
    return hucreler, top


def _ozet_ay(p, ay):
    hucreler, top = _gun_hucreleri(p, ay)
    yev = sorgu.yevmiye_hesapla(top["normal"], top["fazla"], top["kn"], top["kf"])
    return [
        html.Section([
            ui.kpi("Normal", f"{ui.saat(top['normal'])} sa", "navy"),
            ui.kpi("Fazla mesai", f"{ui.saat(top['fazla'])} sa", "orange",
                   f"Pazar ×2,5: {ui.saat(top['kn'] + top['kf'])} sa ({top['katli_gun']} gün)" if top["katli_gun"] else None),
            ui.kpi("Yevmiye günü", ui.sayi(yev, 1), None, f"{top['gun']} çalışılan gün"),
            ui.kpi("İzin / rapor", f"{top['izin']} gün", "purple",
                   f"{top['eksik']} gün veri yok" if top["eksik"] else None),
        ], className="g g-4 gap-12"),
        ui.kart([
            html.Div([ui.h2(f"{ui.ay_etiketi(ay)} · günlük puantaj", buyuk=True),
                      ui.lejant("#1F4E78", "Normal"), ui.lejant("#E8762C", "Fazla"),
                      html.Span([html.I(style={"background": "#EDE3F5", "border": "1px solid #7B4FA8"}), "İzin / HT"],
                                className="vm-legend")], className="vm-card-head"),
            html.Div(hucreler, className="vm-pcal"),
        ], "pad-l col", style={"gap": "12px"}),
    ]


def _gecmis_listesi(tc, limit=None):
    h = pdb.get_history(tc)
    p = pdb.get_personnel(tc)
    olaylar = [{"tarih": x["changed_at"], "alan": ALAN_ETIKET.get(x["field"], x["field"]), "eski": x["old_value"] or "—",
                "yeni": x["new_value"] or "—", "kaynak": x["source"] or "", "renk": "#E8762C" if i == 0 else "#1F4E78"}
               for i, x in enumerate(h)]
    if p:
        olaylar.append({"tarih": p["created_at"], "alan": "Kayıt", "eski": "—", "yeni": "Veri setine eklendi",
                        "kaynak": p.get("kaynak") if not h else "", "renk": "#8C99A6"})
    if limit:
        olaylar = olaylar[:limit]
    return html.Ol([html.Li([
        html.Div([html.Span(style={"background": o["renk"]}), html.Span()], className="vm-tl-rail"),
        html.Div([
            html.Span([html.Span(ui.tarih_tr(o["tarih"]), className="mono"), f" · {o['alan']}"],
                      style={"fontSize": "13px", "color": "#5A6572"}),
            html.Span(html.Strong(o["yeni"]) if o["alan"] == "Kayıt" else
                      [html.Span(o["eski"], className="vm-strike"), " → ", html.Strong(o["yeni"])], style={"fontSize": "14px"}),
            html.Span(o["kaynak"], className="vm-sub", style={"fontSize": "11px"}) if o["kaynak"] else None,
        ], className="vm-tl-body"),
    ]) for o in olaylar], className="vm-tl"), len(h)


def _uyari_kartlari(tc, sadece_acik=True):
    acik = sorgu.get_alerts(tc=tc, cozuldu=False)
    liste = acik if sadece_acik else acik + sorgu.get_alerts(tc=tc, cozuldu=True)
    kartlar = []
    for a in liste:
        sev = sorgu.uyari_seviyesi(a["tip"], a["onem"])
        kartlar.append(html.Div([
            html.Div([ui.seviye_chip(sev), html.Span(sorgu.UYARI_ETIKET.get(a["tip"], a["tip"]), style={"fontWeight": 600}),
                      ui.chip("çözüldü", "green") if a["cozuldu"] else None], className="row", style={"gap": "8px"}),
            html.Span(a["detay"], style={"fontSize": "13px", "color": "#3E4751"}),
            dcc.Link("Veri Kalitesi'nde incele →", href=f"/kalite/{a['tip']}", style={"fontSize": "13px", "fontWeight": 600}),
        ], className="vm-card pad col " + ("danger" if sev == "kritik" and not a["cozuldu"] else ""), style={"gap": "8px"}))
    return kartlar, len(acik)


def sayfa_kart(tc):
    """Personel Kartı: kimlik bandı (yaka kartı + acil durum), İSG şeridi ve sekmeler.
    Yerleşim parçaları veri_merkezi_personel_karti'nda; bu fonksiyon yalnız birleştirir."""
    p = pdb.get_personnel(tc)
    if not p:
        gunluk = sorgu.get_person_months(tc)
        return ui.sayfa("Personel Kartı", "Kayıtlar / Personel", [
            ui.banner([html.Strong(f"TC {tc} veri setinde kayıtlı değil. "),
                       f"Bu TC için {len(gunluk)} aylık puantaj kaydı var." if gunluk else ""], "warn"),
            html.Div([dcc.Link([ui.ikon("person-plus"), "Bu TC ile personel kaydı aç"], href=f"/personel-duzenle/{tc}",
                               className="vm-btn pri"),
                      dcc.Link("Personel listesine dön", href="/personel", className="vm-btn")], className="row"),
        ])
    ikb = ik.getir(tc)
    aylar = _ay_sec(tc)
    ay = aylar[-1]
    gecmis, gecmis_n = _gecmis_listesi(tc)
    gecmis_kisa, _ = _gecmis_listesi(tc, 5)
    uyarilar, acik_n = _uyari_kartlari(tc)
    tum_uyarilar, _ = _uyari_kartlari(tc, sadece_acik=False)
    kritik = sum(1 for a in sorgu.get_alerts(tc=tc) if sorgu.uyari_seviyesi(a["tip"], a["onem"]) == "kritik")

    ay_secici = html.Div(dcc.Dropdown(id="pk-ay", options=[{"label": ui.ay_etiketi(a), "value": a} for a in reversed(aylar)],
                                      value=ay, clearable=False, searchable=False, style={"width": "170px"}), className="vm-dd")
    ozet = html.Div([
        html.Div([
            html.Div(ay_secici, className="row"),
            html.Div(_ozet_ay(p, ay), id="pk-ozet-ay", className="col", style={"gap": "16px"}),
        ], className="col", style={"gap": "12px", "minWidth": 0}),
        html.Div([
            pk.calisma_ozeti(p, ikb),
            *(uyarilar or [ui.banner("Bu kişi için açık veri kalitesi uyarısı yok.", "ok", "check2-circle")]),
            ui.kart([html.Div(html.H3("Son değişiklikler", className="pk-kart-baslik"), className="pk-kart-bas"),
                     gecmis_kisa], "pad-l col pk-kart", style={"gap": "12px"}),
        ], className="col", style={"gap": "16px"}),
    ], className="g", style={"gridTemplateColumns": "minmax(0,1fr) 380px", "gap": "16px"})

    sekmeler = ui.sekmeler("pk", [
        ("ozet", "Özet", ozet),
        ("puantaj", "Puantaj", html.Div(_puantaj_tablosu(tc, ay), id="pk-puantaj")),
        ("kisisel", "Kişisel bilgiler", pk.sekme_kisisel(p, ikb)),
        ("is", "İş ve sözleşme", pk.sekme_is(p, ikb)),
        ("isg", "İSG ve evrak", pk.sekme_isg(ikb)),
        ("gecmis", f"Değişiklik geçmişi · {gecmis_n}", ui.kart(gecmis, "pad-l")),
        ("pdegisim", f"Puantaj değişiklikleri · {sorgu.get_puantaj_history_count(tc)}", _puantaj_degisiklikleri(tc)),
        ("uyari", f"Uyarılar · {acik_n}",
         html.Div(tum_uyarilar or [ui.banner("Bu kişi için uyarı kaydı yok.", "ok", "check2-circle")],
                  className="col", style={"gap": "12px"})),
    ], aktif="ozet")

    return ui.sayfa("Personel Kartı", f"Kayıtlar / Personel / {p['ad_soyad']}", [
        pk.kimlik_bandi(p, ikb, _durum_chip(p), kritik),
        pk.isg_seridi(ikb),
        html.Div(id="pk-mesaj"),
        sekmeler,
        dcc.Store(id="pk-tc-store", data=tc),
    ])


ISLEM_ETIKET = {"ekle": ("eklendi", "green"), "guncelle": ("değişti", "orange"), "sil": ("silindi", "red")}


def _puantaj_degisiklikleri(tc, limit=500):
    """v5: bu kişinin günlük puantaj kayıtlarında olan her değişiklik (eski → yeni), en yeni önce."""
    kayitlar = sorgu.get_puantaj_history(tc, limit=limit)
    if not kayitlar:
        return ui.bos("Bu kişinin puantaj kayıtlarında henüz bir değişiklik izi yok. (İz v5 ile tutulmaya başlandı; "
                      "önceki içe aktarmalar burada görünmez.)")

    def deger(n, f, d):
        if n is None and f is None and not d:
            return html.Span("—", className="muted")
        return [html.Span(f"{ui.saat(n or 0)} + {ui.saat(f or 0)}", className="mono"),
                html.Span(f" {d}", className="small-12 muted") if d else None]

    satirlar = []
    for r in kayitlar:
        etk, ton = ISLEM_ETIKET.get(r["islem"], (r["islem"], "gray"))
        satirlar.append([
            html.Td(ui.tarih_tr(r["tarih"]), className="mono"),
            html.Td([ui.chip(etk, ton), ui.chip("geri alma", "purple") if r["geri_alma"] else None],
                    className="row", style={"gap": "4px"}),
            html.Td(deger(r["eski_normal"], r["eski_fazla"], r["eski_durum"])),
            html.Td(deger(r["yeni_normal"], r["yeni_fazla"], r["yeni_durum"])),
            html.Td([html.Span(f"#{r['import_batch_id']}", className="mono"),
                     html.Span(f" {r['dosya_adi']}", className="vm-sub") if r.get("dosya_adi") else None]),
            html.Td(ui.tarih_tr(r["zaman"], saatli=True), className="mono small-12 muted"),
        ])
    not_ = html.Span(f"Son {limit} değişiklik gösteriliyor.", className="small-12 muted") if len(kayitlar) >= limit else None
    return html.Div([
        html.Span("Saatler 'normal + fazla' biçimindedir. İçe aktarmalar Denetim › İçe Aktarma Geçmişi'nden geri alınabilir.",
                  className="small-12 muted"),
        html.Section(ui.tablo(["Gün", "İşlem", "Önce", "Sonra", "İçe aktarma", "Zaman"], satirlar), className="vm-card clip"),
        not_,
    ], className="col", style={"gap": "10px"})


def _puantaj_tablosu(tc, ay):
    bas, bit = sorgu.ay_araligi(ay)
    kayitlar = sorgu.get_person_daily(tc, bas, bit)
    if not kayitlar:
        return ui.bos(f"{ui.ay_etiketi(ay)} için puantaj kaydı yok.")
    satirlar = [[html.Td(ui.tarih_tr(r["tarih"]), className="mono"),
                 html.Td(ui.GUN_KISA[datetime.date.fromisoformat(r["tarih"]).weekday()], className="muted"),
                 html.Td(r["durum"] or "—"),
                 html.Td(ui.saat(r["normal_saat"]), className="mono r"),
                 html.Td(ui.saat(r["fazla_saat"]), className="mono r c-orange"),
                 html.Td(ui.saat((r["normal_saat"] or 0) + (r["fazla_saat"] or 0)), className="mono r", style={"fontWeight": 600}),
                 html.Td(f"#{r['import_batch_id']}" if r["import_batch_id"] else "—", className="vm-sub")] for r in kayitlar]
    return html.Section(ui.tablo(["Tarih", "Gün", "Çalışma durumu", ("Normal", "r"), ("Fazla", "r"), ("Toplam", "r"),
                                  "İçe aktarma"], satirlar), className="vm-card clip")


@callback(
    Output("pk-ozet-ay", "children"),
    Output("pk-puantaj", "children"),
    Input("pk-ay", "value"),
    State("pk-tc-store", "data"),
    prevent_initial_call=True,
)
def kart_ay_degisti(ay, tc):
    p = pdb.get_personnel(tc)
    if not p or not ay:
        return no_update, no_update
    return _ozet_ay(p, ay), _puantaj_tablosu(tc, ay)


@callback(
    Output("pk-tc", "children"),
    Input("pk-tc-goster", "n_clicks"),
    State("pk-tc-store", "data"),
    prevent_initial_call=True,
)
def tc_goster(n, tc):
    return tc if (n or 0) % 2 == 1 else ui.maske_tc(tc)


@callback(
    Output("pk-mesaj", "children"),
    Input("pk-excel", "n_clicks"),
    State("pk-tc-store", "data"),
    prevent_initial_call=True,
)
def kart_excel(n, tc):
    if not n:
        return no_update
    p = pdb.get_personnel(tc) or {"ad_soyad": tc, "tc": tc}
    kayitlar = sorgu.get_daily_hours(tc=tc)
    wb = openpyxl.Workbook()
    # ---- 1. sayfa: kişi bilgileri (Veri Merkezi + İK)
    wk = wb.active
    wk.title = "Kişi bilgileri"
    report.banner(wk, "A1:B1", f"{p['ad_soyad']} — PERSONEL KARTI", height=26)
    r = 3
    for etiket, deger in pk.excel_satirlari(p, ik.getir(tc)):
        if etiket and deger is None and etiket.isupper():
            wk.cell(row=r, column=1, value=etiket)
            report.style_header_row(wk, r, 1, 2)
        elif etiket:
            wk.cell(row=r, column=1, value=etiket)
            wk.cell(row=r, column=2, value=deger)
            report.style_data_row(wk, r, 1, 2, zebra=(r % 2 == 0))
        r += 1
    wk.column_dimensions["A"].width = 34
    wk.column_dimensions["B"].width = 60
    # ---- 2. sayfa: puantaj
    ws = wb.create_sheet("Puantaj")
    report.banner(ws, "A1:F1", f"{p['ad_soyad']} — PUANTAJ", height=26)
    for c, h in enumerate(["Tarih", "Gün", "Çalışma Durumu", "Normal (sa)", "Fazla (sa)", "Toplam (sa)"], start=1):
        ws.cell(row=3, column=c, value=h)
    report.style_header_row(ws, 3, 1, 6)
    for i, r in enumerate(kayitlar):
        d = datetime.date.fromisoformat(r["tarih"])
        for c, v in enumerate([d, ui.GUN_KISA[d.weekday()], r["durum"], r["normal_saat"], r["fazla_saat"], r["toplam_saat"]], 1):
            cell = ws.cell(row=4 + i, column=c, value=v)
            if c == 1:
                cell.number_format = "DD.MM.YYYY"
            elif c >= 4:
                cell.number_format = report.HOURS_FMT
        report.style_data_row(ws, 4 + i, 1, 6, zebra=(i % 2 == 1))
    for col, w in zip("ABCDEF", (13, 8, 22, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    report.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    yol = report.OUTPUT_DIR / f"{report._slugify(p['ad_soyad'])}_personel_karti.xlsx"
    try:
        wb.save(yol)
    except PermissionError:
        return ui.banner("Dosya yazılamadı: aynı adlı Excel şu an açık olabilir.", "err")
    el.open_path(yol)
    return ui.banner(["Excel oluşturuldu: ", html.Code(str(yol))], "ok", "check2-circle")


# ============================================================ PERSONEL DÜZENLE

FORM_ALANLARI = ["ad_soyad", "grup", "bagli_oldugu_ekip", "gorevi", "ise_giris_tarihi", "isten_cikis_tarihi", "notlar"]


def _secenekler():
    kayitlar = pdb.query_personnel()
    kayitlar.sort(key=lambda r: ui.tr_upper(r["ad_soyad"]))
    return [{"label": f"{r['ad_soyad']} — {r['tc']}", "value": r["tc"]} for r in kayitlar]


def _datalist(id_, kolon):
    try:
        degerler = pdb.distinct_values(kolon)
    except Exception:
        degerler = []
    return html.Datalist([html.Option(value=v) for v in degerler], id=id_)


def _metin(id_, deger, liste=None, **kw):
    return dcc.Input(id=id_, type="text", value=deger or "", list=liste, className="vm-input", debounce=False, **kw)


def sayfa_duzenle(tc=None):
    p = pdb.get_personnel(tc) if tc and tc != "yeni" else None
    mod = "duzenle" if p else "yeni"
    yeni_tc = "" if (p or tc in (None, "yeni")) else tc
    kayit = p or {}
    orijinal = {k: (kayit.get(k) or "") for k in FORM_ALANLARI} if p else None

    def bolum(baslik, *icerik):
        return html.Div([html.Span(baslik, className="disp", style={"fontSize": "15px", "fontWeight": 700}), *icerik],
                        className="col", style={"gap": "12px"})

    form = ui.kart([
        bolum("Kimlik", html.Div([
            html.Div([html.Label("TC Kimlik No", htmlFor="pd-tc", className="vm-lbl"),
                      dcc.Input(id="pd-tc", type="text", value=kayit.get("tc") or yeni_tc, maxLength=11,
                                disabled=mod == "duzenle", className="vm-input mono", inputMode="numeric"),
                      html.Span(id="pd-tc-msg", className="vm-field-msg")], className="vm-field"),
            ui.alan("Ad Soyad", _metin("pd-ad", kayit.get("ad_soyad")), "pd-ad"),
        ], className="g", style={"gridTemplateColumns": "260px minmax(0,1fr)"})),
        bolum("Organizasyon", html.Div([
            ui.alan("Grup", _metin("pd-grup", kayit.get("grup"), "pd-grup-list"), "pd-grup"),
            ui.alan("Bağlı olduğu ekip", _metin("pd-ekip", kayit.get("bagli_oldugu_ekip"), "pd-ekip-list"), "pd-ekip"),
            ui.alan("Görevi", _metin("pd-gorev", kayit.get("gorevi"), "pd-gorev-list"), "pd-gorev"),
        ], className="g g-3"),
            _datalist("pd-grup-list", "grup"), _datalist("pd-ekip-list", "bagli_oldugu_ekip"),
            _datalist("pd-gorev-list", "gorevi"),
            html.Span("Yazdıkça veri setindeki mevcut değerler önerilir — aynı ekip iki farklı yazımla kaydedilmez.",
                      className="small-12 muted")),
        bolum("Tarihler", html.Div([
            ui.alan("İşe giriş tarihi", dcc.Input(id="pd-giris", type="date", value=(kayit.get("ise_giris_tarihi") or "")[:10],
                                                  className="vm-input"), "pd-giris"),
            html.Div([html.Label("İşten çıkış tarihi", htmlFor="pd-cikis", className="vm-lbl"),
                      dcc.Input(id="pd-cikis", type="date", value=(kayit.get("isten_cikis_tarihi") or "")[:10], className="vm-input"),
                      html.Span("Boş = hâlâ çalışıyor", className="small-12 muted")], className="vm-field"),
        ], className="g g-3")),
        ui.alan("Notlar", dcc.Textarea(id="pd-notlar", value=kayit.get("notlar") or "", className="vm-input"), "pd-notlar"),
    ], "pad-l col", style={"gap": "22px", "padding": "22px 24px"})

    sag = html.Div([
        ui.kart([
            html.Div([ui.h2("Kaydedilecek değişiklikler"), html.Span(id="pd-fark-sayi", className="mono",
                                                                     style={"fontWeight": 600})], className="vm-card-head"),
            html.Div(id="pd-fark", className="col", style={"gap": "8px"}),
            html.Span("Her değişiklik, eski değeriyle birlikte değişiklik geçmişine yazılır.", className="small-12 muted"),
            html.Div([html.Button([ui.ikon("check-lg"), "Kaydet"], id="pd-kaydet", className="vm-btn pri grow", n_clicks=0),
                      html.Button("Geri al", id="pd-gerial", className="vm-btn", n_clicks=0)], className="row", style={"gap": "8px"}),
            html.Div(id="pd-sonuc"),
        ], "pad-l col", style={"gap": "12px"}),
        html.Div([
            html.Span("Personeli sil", className="disp", style={"fontSize": "15px", "fontWeight": 700, "color": "#8F1B12"}),
            html.Span(["Kayıt ve tüm değişiklik geçmişi kalıcı olarak silinir. Personel ayrıldıysa silmek yerine ",
                       html.Strong("işten çıkış tarihi"), " girin — geçmiş puantajı raporlarda kalır."],
                      style={"fontSize": "13px", "color": "#3E4751"}),
            html.Button("Kalıcı olarak sil", id="pd-sil", className="vm-btn danger", n_clicks=0, disabled=mod != "duzenle"),
        ], className="vm-card pad-l col", style={"gap": "10px", "borderColor": "#F0C6C0"}),
    ], className="col", style={"gap": "16px"})

    ust = html.Div([
        html.Div(dcc.Dropdown(id="pd-sec", options=_secenekler(), value=p["tc"] if p else None, searchable=True,
                              placeholder="Düzenlenecek personeli ada veya TC'ye göre ara…"), className="vm-dd",
                 style={"width": "460px"}),
        ui.segment("pd-mod", [("duzenle", "Mevcut personeli düzenle"), ("yeni", "Yeni personel")], mod),
    ], className="row")

    return ui.sayfa("Personel Düzenle", "Kayıtlar / Personel / Düzenle", [
        ust,
        html.Div([form, sag], className="g", style={"gridTemplateColumns": "minmax(0,1fr) 400px"}),
        dcc.Store(id="pd-orijinal", data=orijinal),
        dcc.ConfirmDialog(id="pd-sil-onay", message="Bu personeli ve tüm değişiklik geçmişini kalıcı olarak silmek "
                                                     "istediğinize emin misiniz? Bu işlem geri alınamaz."),
    ])


_FORM_CIKTILARI = [Output("pd-tc", "value"), Output("pd-tc", "disabled"), Output("pd-ad", "value"),
                   Output("pd-grup", "value"), Output("pd-ekip", "value"), Output("pd-gorev", "value"),
                   Output("pd-giris", "value"), Output("pd-cikis", "value"), Output("pd-notlar", "value"),
                   Output("pd-orijinal", "data"), Output("pd-sil", "disabled")]


def _form_degerleri(p):
    if not p:
        return ["", False, "", "", "", "", "", "", "", None, True]
    orijinal = {k: (p.get(k) or "") for k in FORM_ALANLARI}
    return [p["tc"], True, p["ad_soyad"] or "", p["grup"] or "", p["bagli_oldugu_ekip"] or "", p["gorevi"] or "",
            (p["ise_giris_tarihi"] or "")[:10], (p["isten_cikis_tarihi"] or "")[:10], p["notlar"] or "", orijinal, False]


@callback(
    *_FORM_CIKTILARI,
    Output("pd-mod", "value", allow_duplicate=True),
    Output("pd-sec", "value", allow_duplicate=True),
    Input("pd-sec", "value"),
    Input("pd-mod", "value"),
    Input("pd-gerial", "n_clicks"),
    State("pd-orijinal", "data"),
    State("pd-tc", "value"),
    prevent_initial_call=True,
)
def form_yukle(tc_sec, mod, _g, orijinal, tc_simdiki):
    tetik = ctx.triggered_id
    if tetik == "pd-mod":
        if mod == "yeni":
            return _form_degerleri(None) + [no_update, None]
        if tc_sec:
            return _form_degerleri(pdb.get_personnel(tc_sec)) + [no_update, no_update]
        return [no_update] * 11 + [no_update, no_update]
    if tetik == "pd-gerial":
        if orijinal:
            return _form_degerleri(pdb.get_personnel(tc_simdiki)) + [no_update, no_update]
        return _form_degerleri(None)[:1] + [False] + [""] * 7 + [None, True] + [no_update, no_update]
    # pd-sec
    if not tc_sec:
        return [no_update] * 13
    p = pdb.get_personnel(tc_sec)
    if not p:
        return [no_update] * 13
    return _form_degerleri(p) + ["duzenle", no_update]


@callback(
    Output("pd-tc-msg", "children"),
    Output("pd-tc-msg", "style"),
    Input("pd-tc", "value"),
    Input("pd-tc", "disabled"),
)
def tc_kontrol(tc, kilitli):
    if kilitli:
        return "Birincil anahtar — düzenlenemez", {"color": "#5A6572"}
    tc = (tc or "").strip()
    if not tc:
        return "11 haneli TC Kimlik No girin", {"color": "#5A6572"}
    ok, mesaj = pdb.validate_tc(tc)
    if not ok:
        return mesaj, {"color": "#B42318"}
    if pdb.get_personnel(tc):
        return "Bu TC ile zaten kayıt var — yukarıdan seçip düzenleyin.", {"color": "#9A3A0B"}
    return "Sağlama doğru · veri setinde kayıtlı değil", {"color": "#2F6B3A"}


@callback(
    Output("pd-fark", "children"),
    Output("pd-fark-sayi", "children"),
    Input("pd-ad", "value"), Input("pd-grup", "value"), Input("pd-ekip", "value"), Input("pd-gorev", "value"),
    Input("pd-giris", "value"), Input("pd-cikis", "value"), Input("pd-notlar", "value"), Input("pd-orijinal", "data"),
)
def fark(ad, grup, ekip, gorev, giris, cikis, notlar, orijinal):
    yeni = dict(zip(FORM_ALANLARI, [ad, grup, ekip, gorev, giris, cikis, notlar]))
    satirlar = []
    for k in FORM_ALANLARI:
        y = (yeni[k] or "").strip()
        e = ((orijinal or {}).get(k) or "").strip()[:10 if "tarih" in k else None]
        if y != e and (y or orijinal):
            deger = ui.tarih_tr(y) if ("tarih" in k and y) else (y or "(boş)")
            eski = ui.tarih_tr(e) if ("tarih" in k and e) else (e or "—")
            satirlar.append(html.Div([ui.lbl(ALAN_ETIKET[k]),
                                      html.Span([html.Span(eski, className="vm-strike"), " → ", html.Strong(deger)],
                                                style={"fontSize": "14px"})],
                                     className="col", style={"gap": "2px", "padding": "10px 12px", "borderRadius": "8px",
                                                             "background": "#F4F2EE"}))
    if not satirlar:
        return [html.Span("Henüz değişiklik yok. Bir alanı düzenleyince burada eski → yeni olarak görünür.",
                          className="small-12 muted")], "0"
    return satirlar, str(len(satirlar))


@callback(
    Output("pd-sonuc", "children"),
    Output("pd-sec", "options"),
    Output("pd-sec", "value"),
    Input("pd-kaydet", "n_clicks"),
    State("pd-mod", "value"), State("pd-tc", "value"), State("pd-ad", "value"), State("pd-grup", "value"),
    State("pd-ekip", "value"), State("pd-gorev", "value"), State("pd-giris", "value"), State("pd-cikis", "value"),
    State("pd-notlar", "value"),
    prevent_initial_call=True,
)
def kaydet(n, mod, tc, ad, grup, ekip, gorev, giris, cikis, notlar):
    if not n:
        return no_update, no_update, no_update
    tc = (tc or "").strip()
    if mod == "yeni":
        ok, mesaj = pdb.validate_tc(tc)
        if not ok:
            return ui.banner(f"TC Kimlik No geçersiz: {mesaj}", "err"), no_update, no_update
        if pdb.get_personnel(tc):
            return ui.banner("Bu TC ile zaten bir personel kayıtlı — yukarıdan seçip düzenleyebilirsiniz.", "warn"), \
                no_update, no_update
    if not (ad or "").strip():
        return ui.banner("Ad Soyad boş olamaz.", "err"), no_update, no_update
    if giris and cikis and giris > cikis:
        return ui.banner("İşe giriş tarihi işten çıkış tarihinden sonra olamaz.", "err"), no_update, no_update
    kayit = {"tc": tc, "ad_soyad": ad.strip(), "grup": (grup or "").strip(), "bagli_oldugu_ekip": (ekip or "").strip(),
             "gorevi": (gorev or "").strip(), "ise_giris_tarihi": giris or "", "isten_cikis_tarihi": cikis or "",
             "notlar": notlar or ""}
    try:
        sonuc = pdb.upsert_personnel(kayit, source="veri_merkezi:manuel", overwrite_blanks=True)
    except pdb.DatasetError as e:
        return ui.banner(str(e), "err"), no_update, no_update
    etiket = {"inserted": "Yeni personel eklendi.", "updated": "Değişiklikler kaydedildi.",
              "unchanged": "Değişiklik yok — kayıt zaten güncel."}[sonuc]
    return (ui.banner([etiket, " ", dcc.Link("Personel kartını aç →", href=f"/personel/{tc}")], "ok", "check2-circle"),
            _secenekler(), tc)


@callback(Output("pd-sil-onay", "displayed"), Input("pd-sil", "n_clicks"), prevent_initial_call=True)
def sil_iste(n):
    return bool(n)


@callback(
    Output("pd-sonuc", "children", allow_duplicate=True),
    Output("pd-sec", "options", allow_duplicate=True),
    Output("pd-mod", "value"),
    Input("pd-sil-onay", "submit_n_clicks"),
    State("pd-tc", "value"),
    prevent_initial_call=True,
)
def sil(n, tc):
    if not n:
        return no_update, no_update, no_update
    silindi = pdb.delete_personnel(tc)
    if silindi:
        ik.sil(tc)
    return (ui.banner("Personel silindi." if silindi else "Silinecek kayıt bulunamadı.", "ok" if silindi else "warn"),
            _secenekler(), "yeni")
