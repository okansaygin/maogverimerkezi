# -*- coding: utf-8 -*-
"""
PERSONEL VERİ SETİ MOTORU (personnel_db.py)
==============================================
Projenin kendi kalıcı personel veritabanı. Herhangi bir aylık puantaj
dosyasından bağımsız olarak, personelin kimlik bilgilerini (ad soyad, TC,
grup, alt ekip, görevi, işe giriş/çıkış tarihi) SQLite veritabanında saklar.

Özellikler:
  - TC Kimlik No birincil anahtardır -> aynı personel asla tekrar kaydedilmez,
    mevcut kayıt güncellenir (upsert).
  - Kaynak excel'den toplu içe aktarma (team_report.build_employee_roster
    ile aynı hiyerarşik grup/alt ekip/meslek filtresini kullanır).
  - Manuel ekleme / düzenleme / silme.
  - Alan bazında değişiklik geçmişi (personnel_history tablosu) - ör. bir
    personelin ekibi veya görevi değiştiğinde eski değer kaybolmaz.
  - Filtreleme (grup, alt ekip, görev, durum) ve serbest metin arama.
  - TC format/checksum doğrulama (puantaj_engine ile aynı algoritma).
  - Aktif/ayrılmış durum, istatistik özeti, excele dışa aktarma.

Bu modül GUI'den bağımsız, tamamen test edilebilir saf Python + sqlite3'tür.
"""

import datetime
import sqlite3
from contextlib import contextmanager
from pathlib import Path

import puantaj_engine as engine

SCRIPT_DIR = Path(__file__).resolve().parent
DB_PATH = SCRIPT_DIR / "personel_veritabani.db"

FIELDS = ["ad_soyad", "grup", "bagli_oldugu_ekip", "gorevi",
          "ise_giris_tarihi", "isten_cikis_tarihi"]

FIELD_LABELS = {
    "ad_soyad": "Ad Soyad",
    "grup": "Grup",
    "bagli_oldugu_ekip": "Bağlı Olduğu Ekip",
    "gorevi": "Görevi",
    "ise_giris_tarihi": "İşe Giriş Tarihi",
    "isten_cikis_tarihi": "İşten Çıkış Tarihi",
    "notlar": "Notlar",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS personnel (
    tc TEXT PRIMARY KEY,
    ad_soyad TEXT NOT NULL,
    grup TEXT,
    bagli_oldugu_ekip TEXT,
    gorevi TEXT,
    ise_giris_tarihi TEXT,
    isten_cikis_tarihi TEXT,
    notlar TEXT DEFAULT '',
    kaynak TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS personnel_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    tc TEXT NOT NULL,
    field TEXT NOT NULL,
    old_value TEXT,
    new_value TEXT,
    changed_at TEXT NOT NULL,
    source TEXT
);

CREATE INDEX IF NOT EXISTS idx_personnel_grup ON personnel(grup);
CREATE INDEX IF NOT EXISTS idx_personnel_ekip ON personnel(bagli_oldugu_ekip);
CREATE INDEX IF NOT EXISTS idx_personnel_gorev ON personnel(gorevi);
CREATE INDEX IF NOT EXISTS idx_history_tc ON personnel_history(tc);
"""


class DatasetError(Exception):
    pass


def _now():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


@contextmanager
def get_connection(db_path=None):
    db_path = db_path or DB_PATH
    # timeout: başka bir bağlantı (ör. Veri Merkezi'nde süren bir içe aktarma) yazarken hemen
    # "database is locked" hatası vermek yerine 15 sn'ye kadar bekler.
    conn = sqlite3.connect(str(db_path), timeout=15)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def init_db(db_path=None):
    with get_connection(db_path) as conn:
        conn.executescript(SCHEMA)


# ---------------------------------------------------------------- TC doğrulama

def validate_tc(tc):
    """(geçerli_mi, mesaj) döner - puantaj_engine ile aynı algoritma."""
    tc = engine.normalize_tc(tc)
    if not tc:
        return False, "TC Kimlik No boş olamaz."
    if not engine.is_valid_tc_format(tc):
        return False, "TC Kimlik No 11 haneli bir sayı olmalı."
    if not engine.tc_checksum_valid(tc):
        return False, "TC Kimlik No sağlaması tutmuyor (yazım hatası olabilir)."
    return True, ""


# ---------------------------------------------------------------- yazma (upsert)

def upsert_personnel(record, source="manuel", db_path=None, log_history=True, overwrite_blanks=False):
    """record: {"tc", "ad_soyad", "grup", "bagli_oldugu_ekip", "gorevi",
    "ise_giris_tarihi", "isten_cikis_tarihi", "notlar"(ops.)} sözlüğü.
    overwrite_blanks=False ise, gelen alan boşsa mevcut değeri SİLMEZ
    (ör. içe aktarımda çıkış tarihi olmayan biri, elle girilmiş bir notu ezmesin).
    Döner: 'inserted' | 'updated' | 'unchanged'"""
    init_db(db_path)
    tc = engine.normalize_tc(record.get("tc"))
    if not tc:
        raise DatasetError("TC Kimlik No boş olamaz.")
    ad_soyad = (record.get("ad_soyad") or "").strip()
    if not ad_soyad:
        raise DatasetError("Ad Soyad boş olamaz.")

    now = _now()
    with get_connection(db_path) as conn:
        row = conn.execute("SELECT * FROM personnel WHERE tc=?", (tc,)).fetchone()

        if row is None:
            conn.execute(
                "INSERT INTO personnel (tc, ad_soyad, grup, bagli_oldugu_ekip, gorevi, "
                "ise_giris_tarihi, isten_cikis_tarihi, notlar, kaynak, created_at, updated_at) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (tc, ad_soyad, record.get("grup"), record.get("bagli_oldugu_ekip"),
                 record.get("gorevi"), record.get("ise_giris_tarihi"),
                 record.get("isten_cikis_tarihi"), record.get("notlar", ""), source, now, now)
            )
            return "inserted"

        changed = False
        for f in ["ad_soyad"] + FIELDS[1:]:
            new_val = record.get(f)
            if f == "ad_soyad":
                new_val = ad_soyad
            if new_val is None or (new_val == "" and not overwrite_blanks):
                continue
            old_val = row[f]
            if (old_val or "") != (new_val or ""):
                conn.execute(f"UPDATE personnel SET {f}=? WHERE tc=?", (new_val, tc))
                if log_history:
                    conn.execute(
                        "INSERT INTO personnel_history (tc, field, old_value, new_value, changed_at, source) "
                        "VALUES (?,?,?,?,?,?)", (tc, f, old_val, new_val, now, source))
                changed = True

        if "notlar" in record and record["notlar"] is not None and record["notlar"] != row["notlar"]:
            conn.execute("UPDATE personnel SET notlar=? WHERE tc=?", (record["notlar"], tc))
            changed = True

        if changed:
            conn.execute("UPDATE personnel SET updated_at=?, kaynak=? WHERE tc=?", (now, source, tc))
            return "updated"
        return "unchanged"


def delete_personnel(tc, db_path=None):
    tc = engine.normalize_tc(tc)
    with get_connection(db_path) as conn:
        conn.execute("DELETE FROM personnel_history WHERE tc=?", (tc,))
        cur = conn.execute("DELETE FROM personnel WHERE tc=?", (tc,))
        return cur.rowcount > 0


# ---------------------------------------------------------------- okuma / filtreleme

def get_personnel(tc, db_path=None):
    with get_connection(db_path) as conn:
        row = conn.execute("SELECT * FROM personnel WHERE tc=?", (engine.normalize_tc(tc),)).fetchone()
        return dict(row) if row else None


def query_personnel(db_path=None, search=None, grup=None, subteam=None, gorevi=None, durum=None):
    """durum: None | 'aktif' | 'ayrilmis'"""
    init_db(db_path)
    sql = "SELECT * FROM personnel WHERE 1=1"
    params = []
    if grup:
        sql += " AND grup=?"
        params.append(grup)
    if subteam:
        sql += " AND bagli_oldugu_ekip=?"
        params.append(subteam)
    if gorevi:
        sql += " AND gorevi=?"
        params.append(gorevi)
    if durum == "aktif":
        sql += " AND (isten_cikis_tarihi IS NULL OR isten_cikis_tarihi='')"
    elif durum == "ayrilmis":
        sql += " AND isten_cikis_tarihi IS NOT NULL AND isten_cikis_tarihi<>''"
    if search:
        sql += " AND (ad_soyad LIKE ? OR tc LIKE ?)"
        like = f"%{search}%"
        params += [like, like]
    sql += " ORDER BY ad_soyad COLLATE NOCASE"
    with get_connection(db_path) as conn:
        return [dict(r) for r in conn.execute(sql, params).fetchall()]


def distinct_values(column, db_path=None):
    if column not in ("grup", "bagli_oldugu_ekip", "gorevi"):
        raise DatasetError(f"Geçersiz sütun: {column}")
    init_db(db_path)
    with get_connection(db_path) as conn:
        rows = conn.execute(
            f"SELECT DISTINCT {column} FROM personnel "
            f"WHERE {column} IS NOT NULL AND {column}<>'' ORDER BY {column} COLLATE NOCASE"
        ).fetchall()
        return [r[0] for r in rows]


def get_history(tc, db_path=None):
    with get_connection(db_path) as conn:
        rows = conn.execute(
            "SELECT * FROM personnel_history WHERE tc=? ORDER BY changed_at DESC", (engine.normalize_tc(tc),)
        ).fetchall()
        return [dict(r) for r in rows]


def get_history_for_tcs(tc_list, db_path=None):
    """Birden fazla TC için değişiklik geçmişini, ad soyad bilgisiyle
    zenginleştirilmiş şekilde döner - Mesai Raporu'nun 'Personel Değişiklikleri'
    sayfasına doğrudan verilebilir (Kademe 4)."""
    if not tc_list:
        return []
    tcs = [engine.normalize_tc(t) for t in tc_list]
    tcs = [t for t in tcs if t]
    if not tcs:
        return []
    placeholders = ",".join("?" * len(tcs))
    with get_connection(db_path) as conn:
        rows = conn.execute(
            f"SELECT h.*, p.ad_soyad FROM personnel_history h "
            f"LEFT JOIN personnel p ON p.tc = h.tc "
            f"WHERE h.tc IN ({placeholders}) ORDER BY h.changed_at DESC",
            tcs,
        ).fetchall()
        results = []
        for r in rows:
            d = dict(r)
            d["name"] = d.pop("ad_soyad", None)
            results.append(d)
        return results


def get_lookup_dict(tc_list=None, db_path=None):
    """{tc: {"grup":.., "bagli_oldugu_ekip":.., "gorevi":..}} - Mesai
    Raporu'nun dataset_lookup parametresine doğrudan verilebilir (Kademe 1).
    tc_list verilmezse veri setindeki TÜM personel döner."""
    init_db(db_path)
    sql = "SELECT tc, grup, bagli_oldugu_ekip, gorevi FROM personnel"
    params = []
    if tc_list:
        tcs = [engine.normalize_tc(t) for t in tc_list]
        tcs = [t for t in tcs if t]
        if not tcs:
            return {}
        sql += f" WHERE tc IN ({','.join('?' * len(tcs))})"
        params = tcs
    with get_connection(db_path) as conn:
        rows = conn.execute(sql, params).fetchall()
    return {r["tc"]: {"grup": r["grup"], "bagli_oldugu_ekip": r["bagli_oldugu_ekip"],
                       "gorevi": r["gorevi"]} for r in rows}


def get_stats(db_path=None):
    init_db(db_path)
    with get_connection(db_path) as conn:
        total = conn.execute("SELECT COUNT(*) FROM personnel").fetchone()[0]
        aktif = conn.execute(
            "SELECT COUNT(*) FROM personnel WHERE isten_cikis_tarihi IS NULL OR isten_cikis_tarihi=''"
        ).fetchone()[0]
        by_group = [dict(r) for r in conn.execute(
            "SELECT grup, COUNT(*) as adet FROM personnel GROUP BY grup ORDER BY adet DESC"
        ).fetchall()]
        by_role = [dict(r) for r in conn.execute(
            "SELECT gorevi, COUNT(*) as adet FROM personnel GROUP BY gorevi ORDER BY adet DESC"
        ).fetchall()]
    return {"total": total, "aktif": aktif, "ayrilmis": total - aktif,
            "by_group": by_group, "by_role": by_role}


# ---------------------------------------------------------------- kaynak excelden içe aktarma

def _read_giris_cikis_dates(cfg, source_path, tc_set):
    """Kaynak exceldeki İşe Giriş / İşten Çıkış Tarihi sütunlarını, verilen
    TC kümesi için okur (her TC'nin ilk geçtiği satır esas alınır)."""
    import openpyxl
    from openpyxl.utils import column_index_from_string

    wb = openpyxl.load_workbook(source_path, data_only=True)
    sheet_name = cfg.get("source_sheet")
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    tc_col = column_index_from_string(cfg.get("source_tc_col", "J"))
    giris_col = column_index_from_string(cfg.get("source_giris_col", "O"))
    cikis_col = column_index_from_string(cfg.get("source_cikis_col", "P"))

    def fmt(v):
        if isinstance(v, datetime.datetime):
            return v.strftime("%Y-%m-%d")
        return str(v).strip() if v else None

    result = {}
    for r in range(cfg.get("source_data_start_row", 4), cfg.get("source_data_end_row", 800) + 1):
        tc = engine.normalize_tc(ws.cell(row=r, column=tc_col).value)
        if tc and tc in tc_set and tc not in result:
            result[tc] = {
                "giris": fmt(ws.cell(row=r, column=giris_col).value),
                "cikis": fmt(ws.cell(row=r, column=cikis_col).value),
            }
    return result


def import_from_source(cfg, source_path, team_filter=None, subteam_filter=None,
                        role_filter=None, db_path=None):
    """team_report.build_employee_roster ile aynı hiyerarşik filtreyi kullanarak
    kaynak exceldeki personeli veri setine aktarır (var olanları günceller,
    yenileri ekler). Ad Soyad, Grup, Alt Ekip, Görev değişse bile TC sabit
    kaldığı için aynı kişi asla tekrar eklenmez."""
    import team_report as report

    roster = report.build_employee_roster(cfg, source_path, team_filter, subteam_filter, role_filter)
    records = roster["records"]
    tc_set = {r["tc"] for r in records}
    dates = _read_giris_cikis_dates(cfg, source_path, tc_set)

    source_label = f"excel:{Path(source_path).name}"
    counts = {"inserted": 0, "updated": 0, "unchanged": 0, "errors": 0}
    error_list = []

    for rec in records:
        tc = rec["tc"]
        d = dates.get(tc, {})
        try:
            result = upsert_personnel({
                "tc": tc,
                "ad_soyad": rec["name"],
                "grup": rec["team"],
                "bagli_oldugu_ekip": rec["subteam"] if rec["subteam"] != report.NO_SUBTEAM else "",
                "gorevi": rec["role"] if rec["role"] != report.NO_ROLE else "",
                "ise_giris_tarihi": d.get("giris"),
                "isten_cikis_tarihi": d.get("cikis"),
            }, source=source_label, db_path=db_path)
            counts[result] += 1
        except DatasetError as e:
            counts["errors"] += 1
            error_list.append((tc, rec.get("name"), str(e)))

    return {
        "total_scanned": len(records),
        "warning": roster["warning"],
        "error_list": error_list,
        **counts,
    }


# ---------------------------------------------------------------- excele dışa aktarma

# ---------------------------------------------------------------- hiyerarşi / rapor entegrasyonu

def get_hierarchy(db_path=None):
    """{grup: {ekip: [görevler]}} - team_report.HierarchyTree ile aynı yapıda,
    Mesai Raporu'nda 'Veri Seti' kaynağı seçildiğinde filtre ağacını doldurmak için."""
    init_db(db_path)
    with get_connection(db_path) as conn:
        rows = conn.execute("SELECT grup, bagli_oldugu_ekip, gorevi FROM personnel").fetchall()
    hierarchy = {}
    for r in rows:
        grup = r["grup"] or "(Grup belirtilmemiş)"
        ekip = r["bagli_oldugu_ekip"] or "(Ekip belirtilmemiş)"
        gorev = r["gorevi"] or "(Görev belirtilmemiş)"
        hierarchy.setdefault(grup, {}).setdefault(ekip, set()).add(gorev)
    return {g: {e: sorted(rs) for e, rs in sorted(subs.items())} for g, subs in sorted(hierarchy.items())}


def get_tc_set(grup_filter=None, ekip_filter=None, gorev_filter=None, durum=None, db_path=None):
    """Verilen (çoklu) filtrelere uyan personelin TC kümesini döner - Mesai
    Raporu'nun tc_filter parametresine doğrudan verilebilir."""
    init_db(db_path)
    sql = "SELECT tc, grup, bagli_oldugu_ekip, gorevi FROM personnel WHERE 1=1"
    params = []
    if durum == "aktif":
        sql += " AND (isten_cikis_tarihi IS NULL OR isten_cikis_tarihi='')"
    elif durum == "ayrilmis":
        sql += " AND isten_cikis_tarihi IS NOT NULL AND isten_cikis_tarihi<>''"
    with get_connection(db_path) as conn:
        rows = conn.execute(sql, params).fetchall()

    result = set()
    for r in rows:
        grup = r["grup"] or "(Grup belirtilmemiş)"
        ekip = r["bagli_oldugu_ekip"] or "(Ekip belirtilmemiş)"
        gorev = r["gorevi"] or "(Görev belirtilmemiş)"
        if grup_filter and grup not in grup_filter:
            continue
        if ekip_filter and ekip not in ekip_filter:
            continue
        if gorev_filter and gorev not in gorev_filter:
            continue
        result.add(r["tc"])
    return result


def export_to_excel(records, output_path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personel Veri Seti"
    ws.sheet_view.showGridLines = False

    headers = ["Ad Soyad", "TC Kimlik No", "Grup", "Bağlı Olduğu Ekip", "Görevi",
               "İşe Giriş Tarihi", "İşten Çıkış Tarihi", "Durum", "Notlar", "Kaynak", "Son Güncelleme"]
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for i, r in enumerate(records, start=2):
        durum = "Ayrılmış" if r.get("isten_cikis_tarihi") else "Aktif"
        vals = [r.get("ad_soyad"), r.get("tc"), r.get("grup"), r.get("bagli_oldugu_ekip"),
                r.get("gorevi"), r.get("ise_giris_tarihi"), r.get("isten_cikis_tarihi"),
                durum, r.get("notlar"), r.get("kaynak"), r.get("updated_at")]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=i, column=c, value=v)

    from openpyxl.utils import get_column_letter
    widths = [26, 15, 22, 22, 26, 16, 16, 12, 30, 26, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
