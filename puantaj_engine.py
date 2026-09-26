"""
PUANTAJ AKTARIM MOTORU (puantaj_engine.py)
============================================
İş mantığının tamamı burada. GUI bu modülü çağırır.

Bu sürümde eklenenler:
  - TC no format/checksum kontrolü
  - Anormal değer uyarısı (negatif / eşik üstü saat)
  - "Kaynakta var, hedefte yok" raporu
  - İsimle yedek eşleştirme önerileri (manuel onay için)
  - Manuel eşleştirme desteği (dışarıdan TC ataması)
  - Özet sayfası (opsiyonel)
  - Önizleme modu (dry_run) - hiçbir şeyi diske yazmadan sonuçları hesaplar
"""

import datetime
import difflib
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

# Kaynaktaki izin kodlarının hedefe yazılacak tam açıklamaları.
# Anahtarlar normalize_name() ile normalize edilmiş haldedir (Türkçe harf
# farkları / büyük-küçük harf önemsizdir) - örn. "Yİ", "yi", "YI" hepsi "YI"a normalize olur.
LEAVE_CODE_MAP = {
    "UCS": "ÜCRETSİZ İZİN",
    "YI": "YILLIK İZİN",
    "RP": "SAĞLIK RAPORU",
    "BI": "BABALIK İZNİ",
    "VI": "VEFAT İZNİ",
    "HT": "HAFTA TATİLİ",
    "DV": "DEVAMSIZLIK",      # günlük mesai formundaki Devamsızlık kodu
}

DEFAULT_CONFIG = {
    "source_sheet": "PERSONEL.0526",
    "source_date_header_row": 3,
    "source_data_start_row": 4,
    "source_data_end_row": 757,
    "source_tc_col": "J",
    "source_type_col": "R",
    "source_name_col": "I",
    "source_date_col_start": "S",
    "source_date_col_end": "AV",

    "target_sheet": "Sheet1",
    "target_data_start_row": 2,
    "target_tc_col": "A",
    "target_name_col": "B",
    "target_write_col": "M",
    "target_note_col": "N",
    "target_status_col": "L",   # "ÇALIŞTI" yazan sütun - izinli günlerde açıklamayla değiştirilir

    "max_reasonable_hours": 24,   # bu değeri aşan/negatif günlük toplamlar "anormal" sayılır
    "write_summary_sheet": True,
    "summary_sheet_name": "Özet",
}


class TransferError(Exception):
    """Kullanıcıya doğrudan gösterilecek, beklenen/anlaşılır hatalar için."""
    pass


# ---------------------------------------------------------------- yardımcılar

def normalize_tc(value):
    if value is None:
        return None
    if isinstance(value, float):
        value = int(value)
    return str(value).strip()


def is_valid_tc_format(tc):
    """11 haneli, sadece rakam, ilk hane 0 değil."""
    return bool(tc) and tc.isdigit() and len(tc) == 11 and tc[0] != "0"


def tc_checksum_valid(tc):
    """Resmi TC Kimlik No algoritması (format geçerliyse çağrılmalı)."""
    digits = [int(c) for c in tc]
    odd_sum = sum(digits[0:9:2])
    even_sum = sum(digits[1:8:2])
    d10_ok = ((odd_sum * 7) - even_sum) % 10 == digits[9]
    d11_ok = sum(digits[0:10]) % 10 == digits[10]
    return d10_ok and d11_ok


def check_tc(tc, side, row, name=None):
    """None -> sorun yok. Aksi halde bir uyarı sözlüğü döner."""
    if not is_valid_tc_format(tc):
        return {"side": side, "row": row, "tc": tc, "name": name, "reason": "format hatalı (11 hane değil / rakam değil)"}
    if not tc_checksum_valid(tc):
        return {"side": side, "row": row, "tc": tc, "name": name, "reason": "checksum uyuşmuyor (yazım hatası olabilir)"}
    return None


TR_MAP = str.maketrans({"İ": "I", "I": "I", "ı": "I", "Ş": "S", "ş": "s",
                         "Ğ": "G", "ğ": "g", "Ü": "U", "ü": "u", "Ö": "O", "ö": "o", "Ç": "C", "ç": "c"})


def normalize_name(s):
    if not s:
        return ""
    s = str(s).translate(TR_MAP).upper()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def parse_date(s):
    if isinstance(s, (datetime.date, datetime.datetime)):
        return s.date() if isinstance(s, datetime.datetime) else s
    return datetime.datetime.strptime(s, "%Y-%m-%d").date()


def find_all_xlsx(folder: Path, label: str, allow_empty=True):
    folder.mkdir(parents=True, exist_ok=True)
    candidates = sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith("~$"))
    if not candidates and not allow_empty:
        raise TransferError(
            f"{label} klasöründe ({folder}) hiç .xlsx dosyası bulunamadı. "
            f"Lütfen ilgili excel dosyasını bu klasöre koy."
        )
    return candidates


def find_single_xlsx(folder: Path, label: str) -> Path:
    files = find_all_xlsx(folder, label, allow_empty=False)
    if len(files) > 1:
        names = ", ".join(p.name for p in files)
        raise TransferError(
            f"{label} klasöründe birden fazla .xlsx dosyası var ({names}). "
            f"Lütfen sadece bir tane bırak."
        )
    return files[0]


def resolve_sheet(wb, configured_name, label):
    if configured_name in wb.sheetnames:
        return configured_name, None
    for name in wb.sheetnames:
        if isinstance(wb[name], Worksheet):
            warning = (f"{label} sayfası '{configured_name}' bulunamadı, "
                       f"otomatik olarak '{name}' sayfası kullanıldı.")
            return name, warning
    raise TransferError(f"{label} için uygun bir çalışma sayfası bulunamadı.")


def find_date_column(ws, header_row, col_start, col_end, target_date):
    start = column_index_from_string(col_start)
    end = column_index_from_string(col_end)
    for c in range(start, end + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, datetime.datetime):
            v = v.date()
        if v == target_date:
            return c
    return None


def load_threaded_comments(xlsx_path):
    ns = {"tc": "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"}
    comments = {}
    with zipfile.ZipFile(xlsx_path) as z:
        names = [n for n in z.namelist() if re.match(r"xl/threadedComments/threadedComment\d+\.xml$", n)]
        for name in names:
            root = ET.fromstring(z.read(name))
            for tc in root.findall("tc:threadedComment", ns):
                ref = tc.attrib.get("ref")
                text_el = tc.find("tc:text", ns)
                text = text_el.text if text_el is not None else None
                if not ref or not text:
                    continue
                comments.setdefault(ref, []).append(text)
    return {ref: "\n".join(texts) for ref, texts in comments.items()}


# ---------------------------------------------------------------- kaynak okuma

def build_day_data(cfg, source_path, target_date, warnings):
    wb = openpyxl.load_workbook(source_path, data_only=True)
    sheet_name, w = resolve_sheet(wb, cfg["source_sheet"], "Kaynak")
    if w:
        warnings.append(w)
    ws = wb[sheet_name]
    threaded_comments = load_threaded_comments(source_path)

    day_col = find_date_column(
        ws, cfg["source_date_header_row"],
        cfg["source_date_col_start"], cfg["source_date_col_end"],
        target_date,
    )
    if day_col is None:
        raise TransferError(f"Kaynak excelde {target_date} tarihine ait bir sütun bulunamadı.")

    tc_col = column_index_from_string(cfg["source_tc_col"])
    type_col = column_index_from_string(cfg["source_type_col"])
    name_col = column_index_from_string(cfg["source_name_col"])

    totals, row_counts, leave_codes, notes, blank_tcs = {}, {}, {}, {}, set()
    leave_status = {}       # tc -> tanınan izin açıklaması (NORMAL MESAİ hücresinden)
    source_names = {}       # tc -> isim (kaynaktaki tüm personel)
    source_tc_issues = []   # kaynaktaki TC format/checksum sorunları

    seen_tc_for_issue_check = set()

    for r in range(cfg["source_data_start_row"], cfg["source_data_end_row"] + 1):
        raw_tc = ws.cell(row=r, column=tc_col).value
        tc = normalize_tc(raw_tc)
        if not tc:
            continue

        name = ws.cell(row=r, column=name_col).value
        if tc not in source_names and name:
            source_names[tc] = str(name).strip()

        if tc not in seen_tc_for_issue_check:
            seen_tc_for_issue_check.add(tc)
            issue = check_tc(tc, "kaynak", r, name)
            if issue:
                source_tc_issues.append(issue)

        cell = ws.cell(row=r, column=day_col)
        val = cell.value
        if isinstance(val, (int, float)):
            totals[tc] = totals.get(tc, 0) + val
            row_counts[tc] = row_counts.get(tc, 0) + 1
        elif val not in (None, ""):
            leave_codes.setdefault(tc, []).append((r, val))
            row_counts[tc] = row_counts.get(tc, 0) + 1
            totals.setdefault(tc, totals.get(tc, 0))
        else:
            blank_tcs.add(tc)

        tip = ws.cell(row=r, column=type_col).value
        if tip and "NORMAL" in str(tip).upper():
            note = threaded_comments.get(cell.coordinate)
            if not note and cell.comment and "[Yorum yazışması]" not in (cell.comment.text or ""):
                note = cell.comment.text
            if note:
                notes[tc] = note

            if val is not None and not isinstance(val, (int, float)):
                code = LEAVE_CODE_MAP.get(normalize_name(str(val)))
                if code:
                    leave_status[tc] = code

    return {
        "totals": totals,
        "row_counts": row_counts,
        "leave_codes": leave_codes,
        "notes": notes,
        "blank_tcs": blank_tcs,
        "leave_status": leave_status,
        "day_col": day_col,
        "source_names": source_names,
        "source_tc_issues": source_tc_issues,
    }


# ---------------------------------------------------------------- isimle eşleştirme önerisi

def suggest_name_matches(target_name, source_names, already_used_tcs, top_n=3, cutoff=0.55):
    """target_name: hedefteki isim (ham). source_names: {tc: isim}.
    already_used_tcs: zaten eşleşmiş TC'ler (öneriye dahil edilmesin diye).
    Döner: [(tc, isim, skor), ...] skor büyükten küçüğe sıralı."""
    target_norm = normalize_name(target_name)
    if not target_norm:
        return []

    candidates = [(tc, name) for tc, name in source_names.items() if tc not in already_used_tcs]
    scored = []
    for tc, name in candidates:
        score = difflib.SequenceMatcher(None, target_norm, normalize_name(name)).ratio()
        if score >= cutoff:
            scored.append((tc, name, round(score, 3)))
    scored.sort(key=lambda x: x[2], reverse=True)
    return scored[:top_n]


# ---------------------------------------------------------------- özet sayfası

def write_summary_sheet(wb, cfg, meta, result):
    sheet_name = cfg.get("summary_sheet_name", "Özet")
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 0)

    bold = Font(bold=True)
    rows = [
        ("Tarih", str(meta["target_date"])),
        ("Kaynak dosya", meta["source_name"]),
        ("Hedef dosya", meta["target_name"]),
        ("İşlem zamanı", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("Hedefteki toplam personel", meta["target_row_count"]),
        ("Eşleşen", len(result["matched"])),
        ("Manuel eşleştirilen", len(result.get("manual_matched", []))),
        ("Eşleşmeyen / kırmızı işaretli", len(result["flagged"])),
        ("İzinli/raporlu görünen (eşleşen içinde)", len(result["on_leave"])),
        ("Açıklama notu aktarılan", len(result["notes_written"])),
        ("İzinli olarak işaretlenen (sarı)", len(result.get("leave_marked", []))),
        ("Anormal değerli (negatif / eşik üstü)", len(result["abnormal"])),
        ("TC format/checksum sorunlu (hedef)", len([i for i in result["tc_issues"] if i["side"] == "hedef"])),
        ("TC format/checksum sorunlu (kaynak)", len([i for i in result["tc_issues"] if i["side"] == "kaynak"])),
        ("Kaynakta var, hedefte yok", len(result["source_not_in_target"])),
    ]
    for i, (label, value) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=label).font = bold
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 40


# ---------------------------------------------------------------- ana işlem

def run_transfer(cfg, source_path, target_path, output_path, target_date_str,
                  manual_mapping=None, dry_run=False):
    """manual_mapping: {hedef_satır: kaynak_tc} - eşleşmeyen satırlar için
    elle atanan kaynak TC'si. dry_run=True ise hiçbir şey yazılmaz/kaydedilmez,
    sadece önizleme sonucu döner (output_path None olur)."""
    warnings = []
    manual_mapping = manual_mapping or {}
    target_date = parse_date(target_date_str)

    src = build_day_data(cfg, source_path, target_date, warnings)
    totals, leave_codes, notes, blank_tcs = src["totals"], src["leave_codes"], src["notes"], src["blank_tcs"]
    leave_status = src["leave_status"]
    source_names = src["source_names"]

    wb = openpyxl.load_workbook(target_path)
    sheet_name, w = resolve_sheet(wb, cfg["target_sheet"], "Hedef")
    if w:
        warnings.append(w)
    ws = wb[sheet_name]

    tc_col = column_index_from_string(cfg["target_tc_col"])
    name_col = column_index_from_string(cfg["target_name_col"])
    write_col = column_index_from_string(cfg["target_write_col"])
    note_col = column_index_from_string(cfg["target_note_col"])
    status_col = column_index_from_string(cfg["target_status_col"])
    last_col = ws.max_column
    max_hours = cfg.get("max_reasonable_hours", 24)

    matched, unmatched, on_leave, notes_written, flagged = [], [], [], [], []
    manual_matched, abnormal, leave_marked = [], [], []
    tc_issues = list(src["source_tc_issues"])
    name_suggestions = {}
    target_tcs_seen = set()
    target_row_count = 0

    for r in range(cfg["target_data_start_row"], ws.max_row + 1):
        raw_tc = ws.cell(row=r, column=tc_col).value
        tc = normalize_tc(raw_tc)
        if not tc:
            continue
        target_row_count += 1
        target_tcs_seen.add(tc)
        target_name = ws.cell(row=r, column=name_col).value

        issue = check_tc(tc, "hedef", r, target_name)
        if issue:
            tc_issues.append(issue)

        effective_tc = tc
        used_manual = False
        if tc not in totals and r in manual_mapping:
            effective_tc = manual_mapping[r]
            used_manual = True

        needs_flag = False
        leave_label = None

        if effective_tc in totals:
            leave_label = leave_status.get(effective_tc)
            value = 0 if leave_label else totals[effective_tc]
            if not dry_run:
                ws.cell(row=r, column=write_col).value = value
            matched.append((r, tc, target_name, value))
            if used_manual:
                manual_matched.append((r, tc, target_name, effective_tc, source_names.get(effective_tc)))
            if isinstance(value, (int, float)) and (value < 0 or value > max_hours):
                abnormal.append((r, tc, target_name, value))
            if effective_tc in leave_codes:
                on_leave.append((r, tc, target_name, leave_codes[effective_tc]))

            note_text = notes.get(effective_tc)
            combined_note = None
            if leave_label and note_text:
                combined_note = f"{leave_label} ({note_text})"
            elif leave_label:
                combined_note = leave_label
            elif note_text:
                combined_note = note_text

            if combined_note:
                if not dry_run:
                    ws.cell(row=r, column=note_col).value = combined_note
                notes_written.append((r, tc, target_name, combined_note))

            if leave_label:
                if not dry_run:
                    ws.cell(row=r, column=status_col).value = leave_label
                leave_marked.append((r, tc, target_name, leave_label))

            if effective_tc in blank_tcs:
                needs_flag = True
        else:
            unmatched.append((r, tc, target_name))
            needs_flag = True
            name_suggestions[r] = suggest_name_matches(target_name, source_names, target_tcs_seen)

        if not dry_run:
            if needs_flag:
                for c in range(1, last_col + 1):
                    ws.cell(row=r, column=c).fill = RED_FILL
            elif leave_label:
                for c in range(1, last_col + 1):
                    ws.cell(row=r, column=c).fill = YELLOW_FILL
        if needs_flag:
            flagged.append((r, tc, target_name))

    source_not_in_target = []
    for tc, value in totals.items():
        if tc in target_tcs_seen:
            continue
        has_leave = tc in leave_codes
        if value or has_leave:
            source_not_in_target.append((tc, source_names.get(tc, "?"), value, leave_codes.get(tc)))

    result = {
        "target_date": target_date,
        "day_col": src["day_col"],
        "matched": matched,
        "manual_matched": manual_matched,
        "unmatched": unmatched,
        "on_leave": on_leave,
        "notes_written": notes_written,
        "flagged": flagged,
        "leave_marked": leave_marked,
        "abnormal": abnormal,
        "tc_issues": tc_issues,
        "source_not_in_target": source_not_in_target,
        "name_suggestions": name_suggestions,
        "odd_row_counts": {tc: c for tc, c in src["row_counts"].items() if c not in (1, 2)},
        "warnings": warnings,
        "output_path": None,
        "source_names": source_names,
    }

    meta = {
        "target_date": target_date,
        "source_name": Path(source_path).name,
        "target_name": Path(target_path).name,
        "target_row_count": target_row_count,
    }

    if not dry_run:
        if cfg.get("write_summary_sheet", True):
            write_summary_sheet(wb, cfg, meta, result)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        result["output_path"] = output_path

    return result
