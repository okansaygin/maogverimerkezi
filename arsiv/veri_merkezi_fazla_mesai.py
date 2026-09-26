# -*- coding: utf-8 -*-
"""
VERİ MERKEZİ - FAZLA MESAİ RAPORU (veri_merkezi_fazla_mesai.py)
===================================================================
Puantaj Suite'teki "Ekip Fazla Mesai Raporu"nun Veri Merkezi karşılığı.

FARK: Suite kaynak Excel'i her seferinde yeniden okur (team_report). Bu modül
aynı analizi KALICI DEPODAN (gunluk_puantaj + personnel) yapar - Excel açılmaz.
Katsayılar ve 9 saatlik standart gün veri_merkezi_sorgu ile aynıdır; aylık
fazla mesai eşiği team_report.CONFIG["overtime_threshold"] değerinden okunur
(Suite'te hangi eşik ayarlıysa burada da o kullanılır).

Excel çıktısı team_report'un biçim yardımcılarıyla (banner, başlık/satır/toplam
stilleri, yazdırma ayarı) yazılır - Suite raporlarıyla aynı görsel dil.
"""

__version__ = "2026-09-24.1"

import datetime
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import team_report as report
import veri_merkezi_sema as sema
import veri_merkezi_sorgu as sorgu

YILLIK_FM_SINIRI = 270.0   # 4857 sayılı İş Kanunu md. 41: yılda en fazla 270 saat fazla çalışma
NO_GRUP = "(Grup belirtilmemiş)"
NO_EKIP = report.NO_SUBTEAM
NO_GOREV = report.NO_ROLE
AY_KISA = ["Oca", "Şub", "Mar", "Nis", "May", "Haz", "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]


def aylik_esik():
    try:
        return float(report.CONFIG.get("overtime_threshold", 52.0))
    except Exception:
        return 52.0


def _periyot_anahtari(d, periyot):
    if periyot == "gun":
        return d, d.strftime("%d.%m")
    if periyot == "ay":
        return (d.year, d.month), f"{AY_KISA[d.month - 1]} {d.year}"
    pzt = d - datetime.timedelta(days=d.weekday())
    return pzt, None   # hafta etiketi aralıkla sonradan kırpılarak üretilir


def _satirlar(bas, bit, gruplar, ekipler, gorevler, db_path=None):
    return sorgu.get_daily_hours_filtered(gruplar, ekipler, gorevler, None, bas, bit, db_path=db_path)


def rapor_verisi(bas, bit, periyot="hafta", gruplar=None, ekipler=None, gorevler=None, db_path=None):
    """bas/bit: ISO tarih. periyot: 'gun' | 'hafta' | 'ay'. Döner: ekran + Excel için tek sözlük."""
    d1, d2 = datetime.date.fromisoformat(bas), datetime.date.fromisoformat(bit)
    rows = _satirlar(bas, bit, gruplar, ekipler, gorevler, db_path)

    # Kırılım: tek grup seçiliyse alt ekip, değilse grup (team_report ile aynı kural)
    kirilim = "ekip" if (gruplar and len(gruplar) == 1) else "grup"

    periyotlar = {}      # anahtar -> {etiket, bas, bit, normal, fazla}
    kirilimlar = {}      # isim -> {normal, fazla, kisi:set}
    kisiler = {}         # tc -> {...}
    normal = fazla = fazla_pz = 0.0
    for r in rows:
        d = datetime.date.fromisoformat(r["tarih"])
        n, f = r["normal_saat"] or 0.0, r["fazla_saat"] or 0.0
        pazar = d.weekday() == 6
        normal += n
        fazla += f
        if pazar:
            fazla_pz += f
        key, etiket = _periyot_anahtari(d, periyot)
        p = periyotlar.setdefault(key, {"etiket": etiket, "bas": d, "bit": d, "normal": 0.0, "fazla": 0.0})
        p["bas"], p["bit"] = min(p["bas"], d), max(p["bit"], d)
        p["normal"] += n
        p["fazla"] += f

        grup = r["grup"] or NO_GRUP
        ekip = r["bagli_oldugu_ekip"] or NO_EKIP
        k = kirilimlar.setdefault(ekip if kirilim == "ekip" else grup, {"normal": 0.0, "fazla": 0.0, "kisi": set()})
        k["normal"] += n
        k["fazla"] += f
        k["kisi"].add(r["tc"])

        ks = kisiler.setdefault(r["tc"], {"tc": r["tc"], "ad_soyad": r["ad_soyad"], "grup": grup, "ekip": ekip,
                                          "gorev": r["gorevi"] or NO_GOREV, "normal": 0.0, "fazla": 0.0,
                                          "fazla_pazar": 0.0, "fazla_hi": 0.0, "gun": 0, "periyot": {}})
        ks["normal"] += n
        ks["fazla"] += f
        ks["fazla_pazar" if pazar else "fazla_hi"] += f
        ks["gun"] += 1
        ks["periyot"][key] = ks["periyot"].get(key, 0.0) + f

    # hafta etiketleri
    for key, p in periyotlar.items():
        if p["etiket"] is None:
            lo = max(key, d1)
            hi = min(key + datetime.timedelta(days=6), d2)
            p["etiket"] = f"{lo:%d}–{hi:%d.%m}" if lo.month == hi.month else f"{lo:%d.%m}–{hi:%d.%m}"
    per_sirali = [dict(periyotlar[k], anahtar=k) for k in sorted(periyotlar)]

    # önceki eşit uzunluktaki dönem
    gun_sayisi = (d2 - d1).days + 1
    onceki_bit = d1 - datetime.timedelta(days=1)
    onceki_bas = onceki_bit - datetime.timedelta(days=gun_sayisi - 1)
    onceki = _satirlar(onceki_bas.isoformat(), onceki_bit.isoformat(), gruplar, ekipler, gorevler, db_path)
    onceki_fazla = sum(r["fazla_saat"] or 0.0 for r in onceki)
    degisim = ((fazla - onceki_fazla) / onceki_fazla * 100) if onceki_fazla else None

    # yıllık fazla mesai (seçili dönemin bitiş yılı, yıl başından dönem sonuna)
    yil = d2.year
    yillik = {}
    if kisiler:
        conn = sema.get_connection(db_path)
        try:
            tcs = list(kisiler)
            for i in range(0, len(tcs), 900):
                parca = tcs[i:i + 900]
                for r in conn.execute(
                        f"SELECT tc, COALESCE(SUM(fazla_saat),0) AS f FROM gunluk_puantaj WHERE tarih BETWEEN ? AND ? "
                        f"AND tc IN ({','.join('?' * len(parca))}) GROUP BY tc",
                        [f"{yil}-01-01", bit] + parca):
                    yillik[r["tc"]] = r["f"]
        finally:
            conn.close()

    esik_donem = aylik_esik() * gun_sayisi / 30.0
    kisi_liste = []
    for tc, k in kisiler.items():
        k["yillik_fazla"] = yillik.get(tc, k["fazla"])
        k["yevmiye"] = sorgu.yevmiye_hesapla(k["normal"], k["fazla_hi"], k["fazla_pazar"])
        k["periyot_liste"] = [k["periyot"].get(p["anahtar"], 0.0) for p in per_sirali]
        k["esik_ustu"] = k["fazla"] > esik_donem
        kisi_liste.append(k)
    kisi_liste.sort(key=lambda k: -k["fazla"])

    kir_liste = [{"ad": ad, "normal": v["normal"], "fazla": v["fazla"], "kisi": len(v["kisi"]),
                  "oran": (v["fazla"] / (v["normal"] + v["fazla"]) * 100) if (v["normal"] + v["fazla"]) else 0.0}
                 for ad, v in kirilimlar.items()]
    kir_liste.sort(key=lambda x: -x["oran"])

    toplam = normal + fazla
    return {
        "bas": d1, "bit": d2, "gun_sayisi": gun_sayisi, "periyot": periyot, "kirilim": kirilim,
        "gruplar": gruplar or [], "ekipler": ekipler or [], "gorevler": gorevler or [],
        "normal": normal, "fazla": fazla, "fazla_pazar": fazla_pz, "toplam": toplam,
        "oran": (fazla / toplam * 100) if toplam else 0.0,
        "kisi_sayisi": len(kisiler), "kisi_basi_fazla": (fazla / len(kisiler)) if kisiler else 0.0,
        "yevmiye": sorgu.yevmiye_hesapla(normal, fazla - fazla_pz, fazla_pz),
        "onceki_fazla": onceki_fazla, "onceki_aralik": (onceki_bas, onceki_bit), "degisim": degisim,
        "periyotlar": per_sirali, "kirilimlar": kir_liste, "kisiler": kisi_liste,
        "esik_donem": esik_donem, "esik_aylik": aylik_esik(),
        "esik_ustu_sayisi": sum(1 for k in kisi_liste if k["esik_ustu"]),
        "yil": yil, "kayit_sayisi": len(rows),
    }


def kapsam_adi(data):
    if data["ekipler"]:
        liste, birim = data["ekipler"], "ALT EKİP"
    elif data["gruplar"]:
        liste, birim = data["gruplar"], "GRUP"
    else:
        return "TÜM GRUPLAR"
    return " / ".join(liste) if len(liste) <= 3 else f"{len(liste)} {birim}"


# ---------------------------------------------------------------- Excel

def excel_yaz(data, output_path=None):
    if not data["kisiler"]:
        raise ValueError("Seçilen dönem ve kapsamda puantaj kaydı yok.")
    kapsam = kapsam_adi(data)
    donem = f"{data['bas']:%d.%m.%Y} – {data['bit']:%d.%m.%Y}"
    if output_path is None:
        report.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = report.OUTPUT_DIR / f"{report._slugify(kapsam)}_fazla_mesai_{data['bas']:%Y%m%d}_{data['bit']:%Y%m%d}.xlsx"
    wb = openpyxl.Workbook()

    # ---- Özet
    ws = wb.active
    ws.title = "Özet"
    ws.sheet_view.showGridLines = False
    report.banner(ws, "A1:D1", f"{kapsam} — FAZLA MESAİ RAPORU", height=30, size=15)
    ws["A2"] = f"Dönem: {donem}  ·  Kaynak: Veri Merkezi kalıcı deposu"
    ws["A2"].font = report.SUBTITLE_FONT
    ws.merge_cells("A2:D2")
    satirlar = [
        ("Personel sayısı", data["kisi_sayisi"], report.INT_FMT),
        ("Normal mesai (saat)", data["normal"], report.HOURS_FMT),
        ("Fazla mesai (saat)", data["fazla"], report.HOURS_FMT),
        ("  — Pazar fazla mesai (saat)", data["fazla_pazar"], report.HOURS_FMT),
        ("Fazla mesai oranı (%)", data["oran"], report.HOURS_FMT),
        ("Kişi başı fazla mesai (saat)", data["kisi_basi_fazla"], report.HOURS_FMT),
        ("Yevmiye günü", data["yevmiye"], report.HOURS_FMT),
        ("Önceki dönem fazla mesai (saat)", data["onceki_fazla"], report.HOURS_FMT),
        ("Değişim (%)", data["degisim"] if data["degisim"] is not None else "—", report.PCT_FMT),
        (f"Dönem eşiğini aşan kişi (aylık {report.format_tr(data['esik_aylik'], 0)} sa)", data["esik_ustu_sayisi"], report.INT_FMT),
    ]
    for i, (lbl, val, fmt) in enumerate(satirlar, start=4):
        ws.cell(row=i, column=1, value=lbl).font = report.KPI_LABEL_FONT
        c = ws.cell(row=i, column=2, value=round(val, 2) if isinstance(val, float) else val)
        c.number_format = fmt
        c.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    report.setup_print(ws, landscape=False)

    # ---- Periyot
    ws2 = wb.create_sheet("Periyot Bazlı")
    ws2.sheet_view.showGridLines = False
    report.banner(ws2, "A1:E1", f"{kapsam} — PERİYOT BAZLI MESAİ", height=26)
    heads = ["Periyot", "Normal (sa)", "Fazla (sa)", "Toplam (sa)", "FM oranı (%)"]
    for c, h in enumerate(heads, start=1):
        ws2.cell(row=3, column=c, value=h)
    report.style_header_row(ws2, 3, 1, 5)
    for i, p in enumerate(data["periyotlar"]):
        r = 4 + i
        top = p["normal"] + p["fazla"]
        for c, v in enumerate([p["etiket"], p["normal"], p["fazla"], top, (p["fazla"] / top * 100) if top else 0], start=1):
            cell = ws2.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
            if c > 1:
                cell.number_format = report.HOURS_FMT
        report.style_data_row(ws2, r, 1, 5, zebra=(i % 2 == 1))
    son = 3 + len(data["periyotlar"])
    for col, w in zip("ABCDE", (18, 14, 14, 14, 14)):
        ws2.column_dimensions[col].width = w
    if data["periyotlar"]:
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "stacked"
        ch.overlap = 100
        ch.title = "Periyot bazlı normal + fazla mesai"
        ch.add_data(Reference(ws2, min_col=2, max_col=3, min_row=3, max_row=son), titles_from_data=True)
        ch.set_categories(Reference(ws2, min_col=1, min_row=4, max_row=son))
        ch.height, ch.width = 9, 20
        ws2.add_chart(ch, "G3")
    report.setup_print(ws2)

    # ---- Ekip
    ws3 = wb.create_sheet("Ekip Bazlı")
    ws3.sheet_view.showGridLines = False
    etiket = "Alt Ekip" if data["kirilim"] == "ekip" else "Grup"
    report.banner(ws3, "A1:F1", f"{kapsam} — {etiket.upper()} BAZLI FAZLA MESAİ", height=26)
    for c, h in enumerate([etiket, "Personel", "Normal (sa)", "Fazla (sa)", "FM oranı (%)", "Kişi başı FM (sa)"], start=1):
        ws3.cell(row=3, column=c, value=h)
    report.style_header_row(ws3, 3, 1, 6)
    for i, k in enumerate(data["kirilimlar"]):
        r = 4 + i
        vals = [k["ad"], k["kisi"], k["normal"], k["fazla"], k["oran"], k["fazla"] / k["kisi"] if k["kisi"] else 0]
        for c, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
            if c == 2:
                cell.number_format = report.INT_FMT
            elif c > 2:
                cell.number_format = report.HOURS_FMT
        report.style_data_row(ws3, r, 1, 6, zebra=(i % 2 == 1))
    if data["kirilimlar"]:
        ws3.conditional_formatting.add(f"E4:E{3 + len(data['kirilimlar'])}",
                                       DataBarRule(start_type="num", start_value=0, end_type="max", color="E8762C"))
    for col, w in zip("ABCDEF", (28, 12, 14, 14, 14, 16)):
        ws3.column_dimensions[col].width = w
    report.setup_print(ws3)

    # ---- Personel
    ws4 = wb.create_sheet("Personel Detay")
    ws4.sheet_view.showGridLines = False
    per_etiket = [p["etiket"] for p in data["periyotlar"]]
    heads = (["Ad Soyad", "TC Kimlik No", "Grup", "Alt Ekip", "Görevi", "Çalışılan gün", "Normal (sa)",
              "Fazla (sa)", "Pazar FM (sa)"] + [f"FM {e}" for e in per_etiket] +
             [f"{data['yil']} yıllık FM (sa)", "Yıllık sınıra kalan (sa)", "Uyarı"])
    last = len(heads)
    report.banner(ws4, f"A1:{get_column_letter(last)}1", f"{kapsam} — PERSONEL FAZLA MESAİ DETAYI", height=26)
    for c, h in enumerate(heads, start=1):
        ws4.cell(row=3, column=c, value=h)
    report.style_header_row(ws4, 3, 1, last)
    ws4.row_dimensions[3].height = 32
    for i, k in enumerate(data["kisiler"]):
        r = 4 + i
        uyarilar = []
        if k["esik_ustu"]:
            uyarilar.append("Dönem eşiği üstü")
        if k["yillik_fazla"] >= YILLIK_FM_SINIRI * 0.85:
            uyarilar.append("Yıllık 270 saat sınırına yakın")
        vals = ([k["ad_soyad"], k["tc"], k["grup"], k["ekip"], k["gorev"], k["gun"], k["normal"], k["fazla"],
                 k["fazla_pazar"]] + k["periyot_liste"] +
                [k["yillik_fazla"], max(0.0, YILLIK_FM_SINIRI - k["yillik_fazla"]), "; ".join(uyarilar)])
        for c, v in enumerate(vals, start=1):
            cell = ws4.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
            if 6 <= c < last:
                cell.number_format = report.INT_FMT if c == 6 else report.HOURS_FMT
        report.style_data_row(ws4, r, 1, last, fill=report.WARN_FILL if uyarilar else None, zebra=(i % 2 == 1))
    for c in range(1, last + 1):
        ws4.column_dimensions[get_column_letter(c)].width = [26, 14, 18, 20, 22][c - 1] if c <= 5 else 12
    ws4.column_dimensions[get_column_letter(last)].width = 34
    ws4.freeze_panes = "C4"
    report.setup_print(ws4)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
