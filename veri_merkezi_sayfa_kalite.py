# -*- coding: utf-8 -*-
"""VERİ MERKEZİ - SAYFA: Veri Kalitesi Merkezi.

Yeni: uyarılar tek tek ya da toplu olarak 'çözüldü' işaretlenebilir (çözüm
notuyla), çözülmüş uyarılar ayrıca görüntülenebilir, içe aktarma bazında
filtrelenebilir ve Excel'e aktarılabilir."""

import datetime

from dash import dcc, html, dash_table, Input, Output, State, ALL, no_update, callback, ctx
import openpyxl

import team_report as report
import veri_merkezi_ekip_listesi as el
import veri_merkezi_sorgu as sorgu
import veri_merkezi_ui as ui

SUTUNLAR = {
    "tc_eksik": [("kaynak_satir", "Satır"), ("ad_soyad", "Ad Soyad")],
    "ayni_tc_farkli_isim": [("tc", "TC"), ("ad_soyad", "Çakışan isimler")],
    "olasi_farkli_tc": [("tc", "Yeni TC"), ("ad_soyad", "Ad Soyad")],
    "tc_format_hatali": [("kaynak_satir", "Satır"), ("tc", "Hatalı TC"), ("ad_soyad", "Ad Soyad")],
    "tcsiz_saat_ozeti": [("ad_soyad", "Etkilenen isimler")],
    "giris_cikis_tutarsiz": [("tc", "TC"), ("ad_soyad", "Ad Soyad")],
    "personel_kaydi_yok": [("tc", "TC"), ("ad_soyad", "Ad Soyad")],
    "tarihsiz_satir": [("kaynak_satir", "Satır"), ("tc", "TC")],
}
VARSAYILAN = [("kaynak_satir", "Satır"), ("tc", "TC"), ("ad_soyad", "Ad Soyad")]


def _batch_secenekleri():
    return [{"label": f"#{b['id']} · {b['dosya_adi']} · {ui.tarih_tr(b['baslama_zamani'])}", "value": b["id"]}
            for b in sorgu.get_import_history_v2(limit=100) if b["durum"] == "tamamlandi"]


def sayfa(tip=None):
    return ui.sayfa("Veri Kalitesi Merkezi", "Denetim / Veri Kalitesi", html.Div([
        # ---- sol: filtreler
        html.Div([
            ui.segment("kl-durum", [("acik", "Açık"), ("cozuldu", "Çözülmüş")], "acik", blok=True),
            html.Div([html.P("Kategori", className="vm-navsec", style={"color": "#5A6572", "padding": 0}),
                      html.Div(id="kl-kategoriler", className="col", style={"gap": "2px"})], className="col",
                     style={"gap": 0}),
            html.Div([html.P("Kaynak içe aktarma", className="vm-navsec", style={"color": "#5A6572", "padding": 0}),
                      ui.dropdown("kl-batch", _batch_secenekleri(), None, yer_tutucu="Tümü")], className="col"),
            html.Div(id="kl-aciklama", style={"marginTop": "auto"}),
        ], className="vm-pane-l"),
        # ---- orta: liste
        html.Div([
            html.H2(id="kl-baslik", className="vm-h2", style={"fontSize": "20px"}),
            html.Div([html.Span("Kutucuklarla birden fazla satır seçip toplu çözebilirsiniz; ayrıntı için satıra tıklayın.",
                                className="small-12 muted grow"),
                      html.Button([ui.ikon("file-earmark-excel"), "Excel'e aktar"], id="kl-excel", className="vm-btn sm", n_clicks=0),
                      html.Button([ui.ikon("check2-all"), "Seçilenleri çözüldü işaretle"], id="kl-toplu",
                                  className="vm-btn sm", n_clicks=0)], className="row", style={"gap": "8px"}),
            html.Div(id="kl-mesaj"),
            html.Div(dash_table.DataTable(
                id="kl-tablo", columns=[], data=[], row_selectable="multi", selected_rows=[],
                page_size=15, sort_action="native", filter_action="native",
                style_table={"overflowX": "auto"},
                style_cell={"whiteSpace": "normal", "height": "auto", "maxWidth": "520px"},
                style_cell_conditional=[{"if": {"column_id": "detay"}, "minWidth": "320px"}],
                style_data_conditional=[{"if": {"state": "active"}, "backgroundColor": "#EEF3F8", "border": "none"}],
            ), className="vm-card clip vm-dt"),
            html.Div(id="kl-trend"),
        ], className="vm-pane-c"),
        # ---- sağ: detay
        html.Aside([
            html.Div(id="kl-detay", className="col", style={"gap": "14px"}),
            html.Div([
                html.Label("Çözüm notu", htmlFor="kl-not", className="vm-lbl"),
                dcc.Textarea(id="kl-not", className="vm-input",
                             placeholder="Örn. Formen ile görüşüldü, satır 688'in TC'si kaynakta düzeltilecek."),
                html.Button([ui.ikon("check-lg"), "Çözüldü olarak işaretle"], id="kl-coz", className="vm-btn pri",
                            n_clicks=0, disabled=True),
                html.Button([ui.ikon("arrow-counterclockwise"), "Yeniden aç"], id="kl-yeniden", className="vm-btn",
                            n_clicks=0, style={"display": "none"}),
            ], className="col", style={"gap": "6px", "marginTop": "auto"}),
        ], className="vm-pane-r"),
        dcc.Store(id="kl-tip", data=tip),
        dcc.Store(id="kl-secili-id"),
        dcc.Store(id="kl-yenile", data=0),
        dcc.Store(id="kl-dummy"),
    ], className="vm-split", style={"gridTemplateColumns": "250px minmax(0,1fr) 360px"}), govde_cls=None)


def _kategori_listesi(sayimlar, secili):
    if not sayimlar:
        return [html.Span("Bu görünümde uyarı yok.", className="small-12 muted", style={"padding": "0 12px"})]
    return [html.Button([ui.dot(a["seviye"]), html.Span(a["etiket"], className="grow"),
                         html.Span(ui.sayi(a["adet"]), className="n")],
                        id={"type": "kl-kat", "tip": a["tip"]}, n_clicks=0,
                        className="vm-catbtn on" if a["tip"] == secili else "vm-catbtn") for a in sayimlar]


def _trend_karti(tip):
    t = sorgu.get_alert_type_trend(tip, 6)
    if not t:
        return None
    mx = max([x["adet"] for x in t] + [1])
    sutunlar = [html.Div([html.Div(className="b last" if i == len(t) - 1 else "b",
                                   style={"height": f"{max(4, x['adet'] / mx * 60):.0f}px"}),
                          html.Span(f"#{x['batch_id']} · {x['adet']}", className="mono", style={"fontSize": "11px", "color": "#5A6572"})])
                for i, x in enumerate(t)]
    return ui.kart([ui.lbl("Bu kategoride son içe aktarmalar"),
                    html.Div(sutunlar, className="vm-spark", style={"gridTemplateColumns": f"repeat({len(t)}, minmax(0,1fr))"})],
                   "pad col", style={"gap": "10px"})


@callback(
    Output("kl-kategoriler", "children"),
    Output("kl-baslik", "children"),
    Output("kl-tablo", "columns"),
    Output("kl-tablo", "data"),
    Output("kl-tablo", "selected_rows"),
    Output("kl-trend", "children"),
    Output("kl-aciklama", "children"),
    Output("kl-tip", "data"),
    Input("kl-durum", "value"),
    Input("kl-tip", "data"),
    Input("kl-batch", "value"),
    Input("kl-yenile", "data"),
)
def listeyi_ciz(durum, tip, batch_id, _y):
    cozuldu = durum == "cozuldu"
    sayimlar = sorgu.get_alert_counts(cozuldu)
    if batch_id:
        sayimlar = [dict(a, adet=len(sorgu.get_alerts(a["tip"], cozuldu, batch_id))) for a in sayimlar]
        sayimlar = [a for a in sayimlar if a["adet"]]
    tipler = [a["tip"] for a in sayimlar]
    if tip not in tipler:
        tip = tipler[0] if tipler else None
    if tip is None:
        return (_kategori_listesi([], None), "Açık veri kalitesi uyarısı yok" if not cozuldu else "Çözülmüş uyarı yok",
                [], [], [], None, ui.banner("Her şey temiz görünüyor.", "ok", "check2-circle"), None)
    kayitlar = sorgu.get_alerts(tip, cozuldu, batch_id)
    kolonlar = SUTUNLAR.get(tip, VARSAYILAN) + [("detay", "Detay"), ("dosya_adi", "Kaynak"), ("tarih", "Tarih")]
    if cozuldu:
        kolonlar += [("cozum_notu", "Çözüm notu")]
    for k in kayitlar:
        k["tarih"] = ui.tarih_tr(k["olusturma_zamani"])
    seviye = sorgu.uyari_seviyesi(tip)
    aciklama = sorgu.UYARI_ACIKLAMA.get(tip)
    kutu = None
    if aciklama:
        renk = {"kritik": ("#F9E1DE", "#8F1B12", "Neden kritik?"), "uyari": ("#FDF3EA", "#6E2A07", "Neden önemli?"),
                "bilgi": ("#ECE8E1", "#3E4751", "Bilgi")}[seviye]
        kutu = html.Div([html.Strong(renk[2]), html.Span(aciklama)], className="col",
                        style={"gap": "4px", "padding": "12px", "borderRadius": "8px", "background": renk[0],
                               "color": renk[1], "fontSize": "13px"})
    return (_kategori_listesi(sayimlar, tip), sorgu.UYARI_ETIKET.get(tip, tip),
            [{"name": ad, "id": k} for k, ad in kolonlar], kayitlar, [], _trend_karti(tip), kutu, tip)


@callback(
    Output("kl-tip", "data", allow_duplicate=True),
    Output("kl-secili-id", "data", allow_duplicate=True),
    Input({"type": "kl-kat", "tip": ALL}, "n_clicks"),
    prevent_initial_call=True,
)
def kategori_secildi(_n):
    if not ctx.triggered_id or not (ctx.triggered and ctx.triggered[0]["value"]):
        return no_update, no_update
    return ctx.triggered_id["tip"], None


@callback(
    Output("kl-secili-id", "data"),
    Input("kl-tablo", "active_cell"),
    State("kl-tablo", "derived_viewport_data"),
    prevent_initial_call=True,
)
def satir_secildi(aktif, gorunen):
    if not aktif:
        return no_update
    if aktif.get("row_id") is not None:
        return aktif["row_id"]
    try:
        return gorunen[aktif["row"]]["id"]
    except (TypeError, IndexError, KeyError):
        return no_update


@callback(
    Output("kl-detay", "children"),
    Output("kl-coz", "disabled"),
    Output("kl-yeniden", "style"),
    Output("kl-not", "value"),
    Input("kl-secili-id", "data"),
    Input("kl-tip", "data"),
)
def detay(alert_id, tip):
    gizli = {"display": "none"}
    a = sorgu.get_alert(alert_id) if alert_id else None
    if not a:
        return [html.Span("Ayrıntıları görmek için listeden bir uyarıya tıklayın. Birden fazla uyarıyı soldaki kutularla "
                          "seçip toplu çözebilirsiniz.", className="muted", style={"fontSize": "14px"})], True, gizli, ""
    seviye = sorgu.uyari_seviyesi(a["tip"], a["onem"])
    parca = [
        html.Div([ui.seviye_chip(seviye), html.Span(f"#A-{a['id']} · içe aktarma #{a['import_batch_id'] or '—'}",
                                                     className="vm-sub")], className="row", style={"gap": "8px"}),
        html.H3(sorgu.UYARI_ETIKET.get(a["tip"], a["tip"]) + (f" — TC {a['tc']}" if a["tc"] else ""), className="vm-h3"),
    ]
    if a["tip"] == "ayni_tc_farkli_isim" and a["ad_soyad_goster"] and " / " in a["ad_soyad_goster"]:
        kutular = []
        for i, isim in enumerate(a["ad_soyad_goster"].split(" / ")[:2]):
            kayitli = a.get("kayitli_isim") and isim.strip() == (a["kayitli_isim"] or "").strip()
            kutular.append(html.Div([
                ui.lbl(f"İsim {'AB'[i]}"), html.Span(isim, style={"fontWeight": 600}),
                html.Span("Veri setindeki isim" if kayitli else "Kaynakta ayrıca geçiyor",
                          className="small-12 " + ("c-green" if kayitli else "c-orange"), style={"fontWeight": 600}),
            ], className="col", style={"gap": "4px", "border": "1px solid #DDD8CF", "borderRadius": "8px", "padding": "12px"}))
        parca.append(html.Div(kutular, className="g g-2 gap-10"))
    else:
        parca.append(html.Div([
            ui.bilgi_satiri("Ad soyad", a["ad_soyad_goster"] or "—"),
            ui.bilgi_satiri("TC", a["tc"] or "—", True),
            ui.bilgi_satiri("Kaynak satır", str(a["kaynak_satir"]) if a["kaynak_satir"] else "—", True),
            ui.bilgi_satiri("Dosya", a["dosya_adi"] or "—", True),
            ui.bilgi_satiri("Tespit", ui.tarih_tr(a["olusturma_zamani"])),
        ], className="col", style={"gap": "6px"}))
    parca.append(html.P(a["detay"], style={"margin": 0, "fontSize": "14px", "lineHeight": 1.5, "color": "#3E4751"}))
    linkler = []
    if a["tc"] and a["personel_var"]:
        linkler.append(dcc.Link([f"{a['kayitli_isim'] or 'Personel'} kaydını aç", ui.ikon("chevron-right")],
                                href=f"/personel/{a['tc']}", className="vm-btn between"))
    elif a["tc"] and a["tip"] in ("personel_kaydi_yok", "olasi_farkli_tc"):
        linkler.append(dcc.Link(["Bu TC ile personel kaydı aç", ui.ikon("chevron-right")],
                                href=f"/personel-duzenle/{a['tc']}", className="vm-btn between"))
    if a["import_batch_id"]:
        linkler.append(dcc.Link(["İçe aktarma kaydını gör", ui.ikon("chevron-right")], href="/gecmis",
                                className="vm-btn between"))
    if linkler:
        parca.append(html.Div(linkler, className="col", style={"gap": "8px"}))
    if a["cozuldu"]:
        parca.append(ui.banner([html.Strong(f"Çözüldü · {ui.tarih_tr(a['cozulme_zamani'])}. "), a["cozum_notu"] or ""],
                               "ok", "check2-circle"))
        return parca, True, {}, ""
    return parca, False, gizli, ""


@callback(
    Output("kl-yenile", "data"),
    Output("kl-secili-id", "data", allow_duplicate=True),
    Output("kl-mesaj", "children"),
    Output("nav-uyari-sayi", "children", allow_duplicate=True),
    Input("kl-coz", "n_clicks"),
    Input("kl-toplu", "n_clicks"),
    Input("kl-yeniden", "n_clicks"),
    State("kl-secili-id", "data"),
    State("kl-not", "value"),
    State("kl-tablo", "selected_rows"),
    State("kl-tablo", "data"),
    State("kl-yenile", "data"),
    prevent_initial_call=True,
)
def coz(_a, _b, _c, secili, notu, secilen_satirlar, veri, sayac):
    if not ctx.triggered or not ctx.triggered[0]["value"]:
        return no_update, no_update, no_update, no_update
    tetik = ctx.triggered_id
    if tetik == "kl-yeniden" and secili:
        sorgu.reopen_alert(secili)
        mesaj = ui.banner("Uyarı yeniden açıldı.", "info", "arrow-counterclockwise")
    elif tetik == "kl-coz" and secili:
        sorgu.resolve_alerts([secili], (notu or "").strip())
        mesaj = ui.banner("Uyarı çözüldü olarak işaretlendi.", "ok", "check2-circle")
    elif tetik == "kl-toplu":
        idler = [veri[i]["id"] for i in (secilen_satirlar or []) if i < len(veri or [])]
        if not idler:
            return no_update, no_update, ui.banner("Önce tablodan en az bir satır seçin (soldaki kutucuklar).", "warn"), no_update
        n = sorgu.resolve_alerts(idler, (notu or "").strip() or "Toplu olarak çözüldü")
        mesaj = ui.banner(f"{n} uyarı çözüldü olarak işaretlendi.", "ok", "check2-all")
    else:
        return no_update, no_update, no_update, no_update
    acik = sorgu.get_alert_totals()["acik"]
    return (sayac or 0) + 1, None, mesaj, (ui.sayi(acik) if acik else "")


@callback(
    Output("kl-mesaj", "children", allow_duplicate=True),
    Input("kl-excel", "n_clicks"),
    State("kl-tablo", "data"),
    State("kl-tablo", "columns"),
    State("kl-tip", "data"),
    prevent_initial_call=True,
)
def excele_aktar(n, veri, kolonlar, tip):
    if not n or not veri:
        return no_update
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Veri Kalitesi"
    for c, k in enumerate(kolonlar, start=1):
        ws.cell(row=1, column=c, value=k["name"])
    report.style_header_row(ws, 1, 1, len(kolonlar))
    for r, satir in enumerate(veri, start=2):
        for c, k in enumerate(kolonlar, start=1):
            ws.cell(row=r, column=c, value=satir.get(k["id"]))
        report.style_data_row(ws, r, 1, len(kolonlar), zebra=(r % 2 == 1))
    for c, k in enumerate(kolonlar, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 60 if k["id"] == "detay" else 20
    ws.freeze_panes = "A2"
    report.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    yol = report.OUTPUT_DIR / f"veri_kalitesi_{tip}_{datetime.date.today():%Y%m%d}.xlsx"
    try:
        wb.save(yol)
    except PermissionError:
        return ui.banner("Dosya yazılamadı: aynı adlı Excel şu an açık olabilir. Kapatıp tekrar deneyin.", "err")
    el.open_path(yol)
    return ui.banner(["Excel oluşturuldu: ", html.Code(str(yol))], "ok", "check2-circle")
