# -*- coding: utf-8 -*-
"""VERİ MERKEZİ - SAYFA: Puantaj Geçmişi (personel × gün matrisi + liste)."""

import datetime

from dash import dcc, html, dash_table, Input, Output, State, no_update, callback, ctx
import openpyxl

import personnel_db as pdb
import team_report as report
import veri_merkezi_ekip_listesi as el
import veri_merkezi_hafta_kurali as hk
import veri_merkezi_sorgu as sorgu
import veri_merkezi_ui as ui

MATRIS_MAX_GUN = 31
MATRIS_MAX_KISI = 250


def _kisaltma(durum):
    kelimeler = [k for k in ui.tr_upper(durum).split() if k]
    return "".join(k[0] for k in kelimeler[:2]) or "—"


def _varsayilan_aralik():
    bugun = datetime.date.today()
    aylar = sorgu.get_covered_periods()
    bu_ay = bugun.strftime("%Y-%m")
    if aylar and bu_ay not in aylar:
        bas, bit = sorgu.ay_araligi(aylar[-1])
        return bas, bit
    return bugun.replace(day=1).isoformat(), max(bugun.replace(day=1), bugun - datetime.timedelta(days=1)).isoformat()


def _hiyerarsi():
    try:
        return pdb.get_hierarchy()
    except Exception:
        return {}


def _ekipler(h, gruplar):
    return sorted({e for g, ek in h.items() if not gruplar or g in gruplar for e in ek}, key=ui.tr_upper)


def _gorevler(h, gruplar, ekipler):
    return sorted({r for g, ek in h.items() if not gruplar or g in gruplar
                   for e, roller in ek.items() if not ekipler or e in ekipler for r in roller}, key=ui.tr_upper)


def sayfa():
    h = _hiyerarsi()
    bas, bit = _varsayilan_aralik()
    kisiler = [{"label": f"{k['ad_soyad']} · {k['tc']}", "value": k["tc"]} for k in sorgu.get_all_personnel_brief()]
    return ui.sayfa("Puantaj Geçmişi", "Kayıtlar / Puantaj Geçmişi", [
        ui.kart([
            html.Div([
                ui.alan("Grup", ui.dropdown("pg-grup", sorted(h, key=ui.tr_upper), [], coklu=True)),
                ui.alan("Alt ekip", ui.dropdown("pg-ekip", _ekipler(h, None), [], coklu=True)),
                ui.alan("Görevi", ui.dropdown("pg-gorev", _gorevler(h, None, None), [], coklu=True)),
                ui.alan("Personel", ui.dropdown("pg-kisi", kisiler, None, yer_tutucu="Ad veya TC")),
                ui.alan("Tarih aralığı", html.Div(dcc.DatePickerRange(id="pg-tarih", start_date=bas, end_date=bit,
                                                                      display_format="DD.MM.YYYY", first_day_of_week=1,
                                                                      minimum_nights=0), className="vm-dd")),
                html.Button([ui.ikon("search"), "Sorgula"], id="pg-sorgula", className="vm-btn pri", n_clicks=0,
                            style={"alignSelf": "end"}),
            ], className="g", style={"gridTemplateColumns": "repeat(4, minmax(0,1fr)) auto auto", "gap": "10px",
                                     "alignItems": "end"}),
            html.Div([html.Span("Hızlı:", className="small-12 muted"),
                      ui.segment("pg-hizli", [("hafta", "Bu hafta"), ("ay", "Bu ay"), ("gecen", "Geçen ay"),
                                              ("uc", "Son 3 ay")], None)], className="row", style={"gap": "8px"}),
        ], "pad col", style={"gap": "12px"}),
        html.Div([
            html.Div(id="pg-kpi", className="grow"),
            ui.segment("pg-gorunum", [("matris", "Matris"), ("liste", "Liste")], "matris"),
            html.Button([ui.ikon("file-earmark-excel"), "Excel'e aktar"], id="pg-excel", className="vm-btn", n_clicks=0),
        ], className="row"),
        html.Div(id="pg-mesaj"),
        dcc.Loading(html.Div(id="pg-sonuc"), type="default", color="#1F4E78"),
    ])


@callback(
    Output("pg-ekip", "options"), Output("pg-ekip", "value"),
    Output("pg-gorev", "options"), Output("pg-gorev", "value"),
    Input("pg-grup", "value"), Input("pg-ekip", "value"),
    State("pg-gorev", "value"),
    prevent_initial_call=True,
)
def filtre_zinciri(gruplar, ekipler, gorevler):
    h = _hiyerarsi()
    eks = _ekipler(h, gruplar)
    ekipler = [e for e in (ekipler or []) if e in eks]
    gors = _gorevler(h, gruplar, ekipler)
    return ([{"label": e, "value": e} for e in eks], ekipler,
            [{"label": g, "value": g} for g in gors], [g for g in (gorevler or []) if g in gors])


@callback(
    Output("pg-tarih", "start_date"), Output("pg-tarih", "end_date"),
    Input("pg-hizli", "value"),
    prevent_initial_call=True,
)
def hizli_aralik(secim):
    b = datetime.date.today()
    dun = b - datetime.timedelta(days=1)
    if secim == "hafta":
        bas = b - datetime.timedelta(days=b.weekday())
        return bas.isoformat(), max(bas, dun).isoformat()
    if secim == "ay":
        bas = b.replace(day=1)
        return bas.isoformat(), max(bas, dun).isoformat()
    if secim == "gecen":
        son = b.replace(day=1) - datetime.timedelta(days=1)
        return son.replace(day=1).isoformat(), son.isoformat()
    if secim == "uc":
        y, m = b.year, b.month - 2
        if m <= 0:
            y, m = y - 1, m + 12
        return datetime.date(y, m, 1).isoformat(), dun.isoformat()
    return no_update, no_update


def _veri(gruplar, ekipler, gorevler, kisi, bas, bit):
    return sorgu.get_daily_hours_filtered(gruplar or None, ekipler or None, gorevler or None, kisi, bas, bit, limit=50000)


def _kpi(rows, bas, bit):
    # Yevmiye hafta tatili / pazar kuralıyla: süzülen kişilerin dönem sınırındaki haftaları da tam okunur.
    tcs = {r["tc"] for r in rows}
    t = hk.db_toplami(bas, bit, tcs=tcs)
    return html.Section([
        ui.kpi("Kişi", ui.sayi(len(tcs))),
        ui.kpi("Normal", f"{ui.sayi(t['normal'])} sa", "navy"),
        ui.kpi("Fazla", f"{ui.sayi(t['fazla'])} sa", "orange"),
        ui.kpi("Yevmiye günü", ui.sayi(t["yevmiye"], 1), None,
               f"pazar ×2,5: {ui.sayi(t['pazar_katli'])} sa" if t["pazar_katli"] else None),
    ], className="g", style={"gridTemplateColumns": "repeat(4, minmax(0, 180px))", "gap": "10px"})


def _matris(rows, bas, bit):
    d1, d2 = datetime.date.fromisoformat(bas), datetime.date.fromisoformat(bit)
    gunler = [d1 + datetime.timedelta(days=i) for i in range((d2 - d1).days + 1)]
    kisiler = {}
    for r in rows:
        k = kisiler.setdefault(r["tc"], {"ad": r["ad_soyad"], "gorev": r["gorevi"] or "—", "gun": {}, "n": 0.0, "f": 0.0})
        k["gun"][r["tarih"]] = r
        k["n"] += r["normal_saat"] or 0
        k["f"] += r["fazla_saat"] or 0
    notlar = []
    liste = sorted(kisiler.items(), key=lambda kv: ui.tr_upper(kv[1]["ad"]))
    if len(liste) > MATRIS_MAX_KISI:
        notlar.append(ui.banner(f"Matriste ilk {MATRIS_MAX_KISI} kişi gösteriliyor ({ui.sayi(len(liste))} kişi var). "
                                "Filtreyi daraltın ya da Liste görünümünü kullanın.", "info", "info-circle"))
        liste = liste[:MATRIS_MAX_KISI]
    bugun = datetime.date.today()
    sablon = {"gridTemplateColumns": f"190px repeat({len(gunler)}, minmax(0, 1fr)) 58px 58px"}
    bas_satir = [ui.lbl("Personel")] + [
        html.Span(str(d.day), className="mono", title=f"{d:%d.%m} {ui.GUN_KISA[d.weekday()]}",
                  style={"fontSize": "11px", "textAlign": "center", "fontWeight": 600,
                         "color": "#9A3A0B" if d.weekday() == 6 else "#5A6572"}) for d in gunler] + [
        ui.lbl("Normal", style={"textAlign": "right"}), ui.lbl("Fazla", style={"textAlign": "right"})]
    satirlar = [html.Div(bas_satir, className="vm-mx", style=sablon)]
    for tc, k in liste:
        hucreler = [dcc.Link([html.Span(k["ad"], style={"fontSize": "13px", "fontWeight": 600, "whiteSpace": "nowrap",
                                                         "overflow": "hidden", "textOverflow": "ellipsis"}),
                              html.Span(k["gorev"], style={"fontSize": "11px", "color": "#5A6572", "whiteSpace": "nowrap",
                                                           "overflow": "hidden", "textOverflow": "ellipsis"})],
                             href=f"/personel/{tc}", className="col", style={"gap": 0, "minWidth": 0, "color": "#16202B",
                                                                              "textDecoration": "none"})]
        for d in gunler:
            r = k["gun"].get(d.isoformat())
            if not r:
                cls, v, t = ("mx-0", "", "") if d >= bugun else ("mx-yok", "", "Veri yok")
            else:
                h = (r["normal_saat"] or 0) + (r["fazla_saat"] or 0)
                durum = (r["durum"] or "").strip()
                if h == 0 and ui.tr_upper(durum) not in ("", "ÇALIŞTI", "CALISTI"):
                    cls, v, t = "mx-iz", _kisaltma(durum), durum
                elif d.weekday() == 6 and h:
                    cls, v, t = "mx-pz", ui.saat(h), f"Pazar · {ui.saat(h)} sa"
                else:
                    cls = "mx-1" if h <= 9 else ("mx-2" if h <= 11 else "mx-3")
                    v, t = ui.saat(h), f"{ui.saat(h)} sa"
            hucreler.append(html.Div(v, className=f"vm-mx-c {cls}", title=f"{d:%d.%m} · {t}" if t else f"{d:%d.%m}"))
        hucreler += [html.Span(ui.saat(k["n"]), className="mono", style={"fontSize": "13px", "textAlign": "right"}),
                     html.Span(ui.saat(k["f"]), className="mono c-orange", style={"fontSize": "13px", "textAlign": "right",
                                                                                  "fontWeight": 600})]
        satirlar.append(html.Div(hucreler, className="vm-mx", style=sablon))
    lejant = html.Div([ui.lejant("#D6E2EE", "≤ 9 sa"), ui.lejant("#8FAACB", "9–11 sa"), ui.lejant("#1F4E78", "> 11 sa"),
                       ui.lejant("#E8762C", "Pazar (FM ×2,5)"), ui.lejant("#EDE3F5", "İzin / HT"),
                       html.Span([html.I(style={"border": "1.5px dashed #B8410C"}), "Veri yok"], className="vm-legend")],
                      className="row", style={"gap": "14px", "paddingTop": "8px", "flexWrap": "wrap"})
    return html.Div([*notlar, ui.kart(html.Div(satirlar + [lejant], className="col", style={"gap": "4px"}),
                                      "pad", style={"overflowX": "auto"})], className="col", style={"gap": "10px"})


def _liste(rows):
    veri = [{"ad_soyad": r["ad_soyad"], "grup": r["grup"] or "", "ekip": r["bagli_oldugu_ekip"] or "",
             "gorevi": r["gorevi"] or "", "tarih": ui.tarih_tr(r["tarih"]), "tarih_iso": r["tarih"],
             "durum": r["durum"] or "", "normal": r["normal_saat"], "fazla": r["fazla_saat"], "toplam": r["toplam_saat"]}
            for r in rows]
    kolonlar = [("ad_soyad", "Ad soyad"), ("grup", "Grup"), ("ekip", "Alt ekip"), ("gorevi", "Görevi"), ("tarih", "Tarih"),
                ("durum", "Çalışma durumu"), ("normal", "Normal"), ("fazla", "Fazla"), ("toplam", "Toplam")]
    return html.Div(dash_table.DataTable(
        columns=[{"name": a, "id": k, "type": "numeric" if k in ("normal", "fazla", "toplam") else "text"} for k, a in kolonlar],
        data=veri, page_size=25, sort_action="native", filter_action="native", style_table={"overflowX": "auto"},
        style_data_conditional=[{"if": {"column_id": "fazla", "filter_query": "{fazla} > 0"}, "color": "#9A3A0B",
                                 "fontWeight": 600}],
    ), className="vm-card clip vm-dt")


@callback(
    Output("pg-sonuc", "children"),
    Output("pg-kpi", "children"),
    Input("pg-sorgula", "n_clicks"),
    Input("pg-gorunum", "value"),
    State("pg-grup", "value"), State("pg-ekip", "value"), State("pg-gorev", "value"), State("pg-kisi", "value"),
    State("pg-tarih", "start_date"), State("pg-tarih", "end_date"),
)
def sorgula(_n, gorunum, gruplar, ekipler, gorevler, kisi, bas, bit):
    if not bas or not bit:
        return ui.banner("Tarih aralığı seçin.", "warn"), None
    bas, bit = bas[:10], bit[:10]
    rows = _veri(gruplar, ekipler, gorevler, kisi, bas, bit)
    if not rows:
        return ui.bos("Bu filtrelerle puantaj kaydı bulunamadı."), None
    gun = (datetime.date.fromisoformat(bit) - datetime.date.fromisoformat(bas)).days + 1
    if gorunum == "matris":
        if gun > MATRIS_MAX_GUN:
            return html.Div([ui.banner(f"Matris en fazla {MATRIS_MAX_GUN} gün gösterir ({gun} gün seçildi) — liste "
                                       "görünümüne geçildi.", "info", "info-circle"), _liste(rows)],
                            className="col", style={"gap": "10px"}), _kpi(rows, bas, bit)
        return _matris(rows, bas, bit), _kpi(rows, bas, bit)
    return _liste(rows), _kpi(rows, bas, bit)


@callback(
    Output("pg-mesaj", "children"),
    Input("pg-excel", "n_clicks"),
    State("pg-grup", "value"), State("pg-ekip", "value"), State("pg-gorev", "value"), State("pg-kisi", "value"),
    State("pg-tarih", "start_date"), State("pg-tarih", "end_date"),
    prevent_initial_call=True,
)
def excel(n, gruplar, ekipler, gorevler, kisi, bas, bit):
    if not n or not bas or not bit:
        return no_update
    rows = _veri(gruplar, ekipler, gorevler, kisi, bas[:10], bit[:10])
    if not rows:
        return ui.banner("Aktarılacak kayıt yok.", "warn")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Puantaj"
    basliklar = ["Ad Soyad", "TC Kimlik No", "Grup", "Alt Ekip", "Görevi", "Tarih", "Çalışma Durumu", "Normal (sa)",
                 "Fazla (sa)", "Toplam (sa)"]
    report.banner(ws, "A1:J1", f"PUANTAJ — {ui.tarih_tr(bas)} – {ui.tarih_tr(bit)}", height=26)
    for c, h in enumerate(basliklar, start=1):
        ws.cell(row=3, column=c, value=h)
    report.style_header_row(ws, 3, 1, len(basliklar))
    for i, r in enumerate(rows):
        vals = [r["ad_soyad"], r["tc"], r["grup"], r["bagli_oldugu_ekip"], r["gorevi"],
                datetime.date.fromisoformat(r["tarih"]), r["durum"], r["normal_saat"], r["fazla_saat"], r["toplam_saat"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=4 + i, column=c, value=v)
            if c == 6:
                cell.number_format = "DD.MM.YYYY"
            elif c >= 8:
                cell.number_format = report.HOURS_FMT
        report.style_data_row(ws, 4 + i, 1, len(basliklar), zebra=(i % 2 == 1))
    for c, w in enumerate([26, 14, 18, 20, 22, 12, 18, 11, 11, 11], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:J{3 + len(rows)}"
    report.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    yol = report.OUTPUT_DIR / f"puantaj_{bas[:10].replace('-', '')}_{bit[:10].replace('-', '')}.xlsx"
    try:
        wb.save(yol)
    except PermissionError:
        return ui.banner("Dosya yazılamadı: aynı adlı Excel şu an açık olabilir.", "err")
    el.open_path(yol)
    return ui.banner(["Excel oluşturuldu: ", html.Code(str(yol)), f" · {ui.sayi(len(rows))} satır"], "ok", "check2-circle")
