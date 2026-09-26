# -*- coding: utf-8 -*-
"""
VERİ MERKEZİ v5 - OTOMATİK TESTLER
====================================
Çalıştırma (uygulama klasöründe):
    python -m unittest discover -s testler -v

Her test kendi geçici klasöründe, kendi veritabanıyla çalışır; gerçek veritabanınıza,
ayar dosyanıza ve yedeklerinize DOKUNMAZ. Gereken paketler: pip install -r requirements-dev.txt
"""

import datetime
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))

import logging  # noqa: E402
import openpyxl  # noqa: E402

# Test çıktısını uygulamanın günlük mesajlarıyla doldurma.
_g = logging.getLogger("veri_merkezi")
_g.addHandler(logging.NullHandler())
_g.propagate = False

import veri_merkezi_ayarlar as ayarlar  # noqa: E402
import veri_merkezi_sema as sema  # noqa: E402
import personnel_db as pdb  # noqa: E402
import veri_merkezi_ice_aktarma as ie  # noqa: E402
import veri_merkezi_mesai_ice_aktarma as mie  # noqa: E402
import veri_merkezi_geri_alma as ga  # noqa: E402
import veri_merkezi_yedek as yedek  # noqa: E402
import veri_merkezi_sorgu as sorgu  # noqa: E402


# ---------------------------------------------------------------- yardımcılar

def gecerli_tc(n):
    """n'den geçerli (sağlaması tutan) bir TC üretir."""
    ilk9 = [int(c) for c in f"{100000000 + n:09d}"]
    ilk9[0] = ilk9[0] or 1
    tek = sum(ilk9[0::2])
    cift = sum(ilk9[1::2])
    d10 = (tek * 7 - cift) % 10
    d11 = (sum(ilk9) + d10) % 10
    return "".join(map(str, ilk9)) + str(d10) + str(d11)


def kimlik_dosyasi(yol, kisiler):
    """kisiler: [(tc, ad, grup, ekip, gorev)]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PERSONEL"
    for c, h in enumerate(["GRUPLAR", "AD SOYAD", "TC KİMLİK NO", "BAĞLI OLDUĞU EKİP", "GÖREVİ",
                           "İŞE GİRİŞ TARİHİ", "İŞTEN ÇIKIŞ TARİHİ"], start=1):
        ws.cell(row=3, column=c, value=h)
    for i, (tc, ad, grup, ekip, gorev) in enumerate(kisiler):
        for c, v in enumerate([grup, ad, tc, ekip, gorev, datetime.date(2026, 1, 5), None], start=1):
            ws.cell(row=4 + i, column=c, value=v)
    wb.save(yol)
    return yol


def mesai_dosyasi(yol, satirlar):
    """satirlar: [(tc, ad, tarih(date), durum, toplam_saat)]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(["TC Kimlik No", "Adı Soyadı", "Tarih", "Çalışma Durumu",
                           "Rönesans Çalışma Saati Girişi"], start=1):
        ws.cell(row=1, column=c, value=h)
    for i, satir in enumerate(satirlar):
        for c, v in enumerate(satir, start=1):
            ws.cell(row=2 + i, column=c, value=v)
    wb.save(yol)
    return yol


class Temel(unittest.TestCase):
    def setUp(self):
        self.klasor = Path(tempfile.mkdtemp(prefix="vm_test_"))
        self.db = self.klasor / "personel_veritabani.db"
        self._ayar_eski = ayarlar.AYAR_DOSYASI
        ayarlar.AYAR_DOSYASI = self.klasor / "veri_merkezi_ayarlar.json"
        ayarlar._onbellek["anahtar"] = None
        pdb.init_db(self.db)
        sema.migrate(self.db)
        self.tc = [gecerli_tc(i) for i in range(1, 6)]
        self.sayac = 0

    def tearDown(self):
        ayarlar.AYAR_DOSYASI = self._ayar_eski
        ayarlar._onbellek["anahtar"] = None
        ayarlar.uygula()
        shutil.rmtree(self.klasor, ignore_errors=True)

    def _yol(self, ad):
        self.sayac += 1
        return self.klasor / f"{self.sayac:02d}_{ad}.xlsx"

    def kimlik_yukle(self, kisiler, kararlar=None):
        on = ie.preview_import(kimlik_dosyasi(self._yol("kimlik"), kisiler), db_path=self.db)
        self.assertEqual(on["durum"], "hazir", on.get("engelleyiciler"))
        return ie.commit_import(on, db_path=self.db, field_resolutions=kararlar)

    def mesai_yukle(self, satirlar):
        on = mie.preview_mesai_import(mesai_dosyasi(self._yol("mesai"), satirlar), db_path=self.db)
        self.assertEqual(on["durum"], "hazir", on.get("engelleyiciler"))
        return mie.commit_mesai_import(on, db_path=self.db)

    def sql(self, q, *p):
        conn = sema.get_connection(self.db)
        try:
            return [dict(r) for r in conn.execute(q, p).fetchall()]
        finally:
            conn.close()

    def gun(self, tc, tarih):
        r = self.sql("SELECT * FROM gunluk_puantaj WHERE tc=? AND tarih=?", tc, tarih)
        return r[0] if r else None


D1, D2, D3 = datetime.date(2026, 9, 1), datetime.date(2026, 9, 2), datetime.date(2026, 9, 3)


# ---------------------------------------------------------------- şema

class SemaTesti(Temel):
    def test_surum_ve_tablolar(self):
        self.assertEqual(sema.get_schema_version(self.db), 5)
        tablolar = {r["name"] for r in self.sql("SELECT name FROM sqlite_master WHERE type='table'")}
        self.assertTrue({"puantaj_history", "personel_degisiklik_izi"} <= tablolar)
        sutunlar = {r["name"] for r in self.sql("PRAGMA table_info(import_batches)")}
        self.assertTrue({"izli", "geri_alma_zamani", "geri_alma_notu", "geri_alma_ozeti"} <= sutunlar)

    def test_wal(self):
        self.assertEqual(sema.gunluk_kipi(self.db), "wal")

    def test_migrate_tekrar_dosyayi_degistirmez(self):
        sema.wal_birlestir(self.db)
        once = (self.db.stat().st_mtime_ns, self.db.stat().st_size)
        wal = self.db.with_name(self.db.name + "-wal")
        wal_once = wal.stat().st_size if wal.exists() else 0
        sema.migrate(self.db)
        sema.migrate(self.db)
        self.assertEqual(wal.stat().st_size if wal.exists() else 0, wal_once,
                         "migrate() gereksiz yazma yaptı (analiz önbelleği boşa geçersiz olur)")
        self.assertEqual((self.db.stat().st_mtime_ns, self.db.stat().st_size), once)

    def test_yazma_surerken_okuma_beklemez(self):
        """WAL: bir bağlantı yazma işlemindeyken diğeri okuyabilmeli (eskiden 'database is locked')."""
        yazan = sema.get_connection(self.db)
        yazan.execute("BEGIN IMMEDIATE")
        yazan.execute("INSERT INTO personnel (tc, ad_soyad, created_at, updated_at) VALUES ('1','X','a','a')")
        okuyan = sqlite3.connect(self.db, timeout=0.2)
        try:
            self.assertEqual(okuyan.execute("SELECT COUNT(*) FROM personnel").fetchone()[0], 0)
        finally:
            okuyan.close()
            yazan.rollback()
            yazan.close()

    def test_yeni_sema_eski_surume_dusurulmez(self):
        conn = sema.get_connection(self.db)
        conn.execute("UPDATE schema_meta SET value='9' WHERE key='version'")
        conn.commit()
        conn.close()
        sema.migrate(self.db)
        self.assertEqual(sema.get_schema_version(self.db), 9)


# ---------------------------------------------------------------- puantaj geçmişi + geri alma (mesai)

class MesaiGeriAlmaTesti(Temel):
    def setUp(self):
        super().setUp()
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA"),
                           (self.tc[1], "AYŞE KARA", "G1", "E1", "OPERATÖR")])

    def test_gecmis_yazilir(self):
        s1 = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 11),
                               (self.tc[1], "AYŞE KARA", D1, "ÇALIŞTI", 9)])
        self.assertEqual(s1["yeni_kayit"], 2)
        s2 = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 8),      # değişti
                               (self.tc[1], "AYŞE KARA", D1, "ÇALIŞTI", 9),     # aynı
                               (self.tc[1], "AYŞE KARA", D2, "ÇALIŞTI", 10)])   # yeni
        self.assertEqual((s2["yeni_kayit"], s2["guncellenen_kayit"], s2["degismeyen_kayit"]), (1, 1, 1))
        h = self.sql("SELECT * FROM puantaj_history WHERE import_batch_id=? ORDER BY id", s2["batch_id"])
        self.assertEqual(len(h), 2, "değişmeyen kayıt için iz yazılmamalı")
        g = [x for x in h if x["islem"] == "guncelle"][0]
        self.assertEqual((g["eski_normal"], g["eski_fazla"], g["yeni_normal"], g["yeni_fazla"]), (9, 2, 8, 0))
        self.assertEqual(g["eski_batch_id"], s1["batch_id"])
        b = self.sql("SELECT izli FROM import_batches WHERE id=?", s2["batch_id"])[0]
        self.assertEqual(b["izli"], 1)
        self.assertEqual(len(sorgu.get_puantaj_history(self.tc[0], db_path=self.db)), 2)

    def test_geri_al_eski_degerlere_doner(self):
        s1 = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 11)])
        s2 = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "SAĞLIK RAPORU", 0),
                               (self.tc[0], "ALİ VELİ", D2, "ÇALIŞTI", 9)])
        on = ga.onizle(s2["batch_id"], db_path=self.db)
        self.assertTrue(on["izin"], on["neden"])
        self.assertEqual(on["sayilar"]["silinecek"], 1)
        self.assertEqual(on["sayilar"]["geri_yuklenecek"], 1)
        # önizleme hiçbir şey değiştirmemeli
        self.assertEqual(self.gun(self.tc[0], "2026-09-01")["durum"], "SAĞLIK RAPORU")

        sonuc = ga.geri_al(s2["batch_id"], "yanlış dosya", db_path=self.db)
        self.assertEqual(sonuc["ozet"]["silinen_gun"], 1)
        self.assertEqual(sonuc["ozet"]["geri_yuklenen_gun"], 1)
        g1 = self.gun(self.tc[0], "2026-09-01")
        self.assertEqual((g1["normal_saat"], g1["fazla_saat"], g1["durum"], g1["import_batch_id"]),
                         (9, 2, "ÇALIŞTI", s1["batch_id"]))
        self.assertIsNone(self.gun(self.tc[0], "2026-09-02"))
        b = self.sql("SELECT * FROM import_batches WHERE id=?", s2["batch_id"])[0]
        self.assertEqual(b["durum"], "geri_alindi")
        self.assertEqual(b["geri_alma_notu"], "yanlış dosya")
        self.assertEqual(json.loads(b["geri_alma_ozeti"])["silinen_gun"], 1)
        # geri alma da iz bırakır
        self.assertEqual(len(self.sql("SELECT * FROM puantaj_history WHERE geri_alma=1")), 2)
        # yedek alındı
        self.assertTrue(sonuc["yedek"])
        self.assertTrue(any(y["tur"] == "geri_alma" for y in yedek.yedekleri_listele(self.db)))

    def test_ikinci_kez_geri_alinamaz(self):
        s = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9)])
        ga.geri_al(s["batch_id"], db_path=self.db)
        on = ga.onizle(s["batch_id"], db_path=self.db)
        self.assertFalse(on["izin"])
        with self.assertRaises(ga.GeriAlmaHatasi):
            ga.geri_al(s["batch_id"], db_path=self.db)

    def test_sonraki_ice_aktarma_engeller_ve_sira_ile_calisir(self):
        s1 = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9)])
        s2 = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 12)])
        on = ga.onizle(s1["batch_id"], db_path=self.db)
        self.assertFalse(on["izin"])
        self.assertEqual([e["id"] for e in on["engelleyenler"]], [s2["batch_id"]])
        with self.assertRaises(ga.GeriAlmaHatasi):
            ga.geri_al(s1["batch_id"], db_path=self.db)
        self.assertEqual(self.gun(self.tc[0], "2026-09-01")["fazla_saat"], 3, "engellenen geri alma veri değiştirmemeli")
        # en yeniden eskiye: önce s2, sonra s1
        ga.geri_al(s2["batch_id"], db_path=self.db)
        self.assertEqual(self.gun(self.tc[0], "2026-09-01")["fazla_saat"], 0)
        ga.geri_al(s1["batch_id"], db_path=self.db)
        self.assertIsNone(self.gun(self.tc[0], "2026-09-01"))

    def test_ayni_dosya_geri_almadan_sonra_yeniden_yuklenebilir(self):
        yol = mesai_dosyasi(self._yol("mesai"), [(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9)])
        s = mie.commit_mesai_import(mie.preview_mesai_import(yol, db_path=self.db), db_path=self.db)
        self.assertEqual(mie.preview_mesai_import(yol, db_path=self.db)["durum"], "engellendi")
        ga.geri_al(s["batch_id"], db_path=self.db)
        self.assertEqual(mie.preview_mesai_import(yol, db_path=self.db)["durum"], "hazir")

    def test_uyarilar_kapatilir(self):
        s = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9),
                              (gecerli_tc(77), "BİLİNMEYEN", D1, "ÇALIŞTI", 9)])   # personel_kaydi_yok uyarısı
        acik = self.sql("SELECT COUNT(*) AS n FROM quality_alerts WHERE import_batch_id=? AND cozuldu=0", s["batch_id"])
        self.assertGreater(acik[0]["n"], 0)
        sonuc = ga.geri_al(s["batch_id"], db_path=self.db)
        self.assertEqual(sonuc["ozet"]["kapatilan_uyari"], acik[0]["n"])
        kalan = self.sql("SELECT COUNT(*) AS n FROM quality_alerts WHERE import_batch_id=? AND cozuldu=0", s["batch_id"])
        self.assertEqual(kalan[0]["n"], 0)
        notlar = self.sql("SELECT cozum_notu FROM quality_alerts WHERE import_batch_id=?", s["batch_id"])
        self.assertTrue(all(n["cozum_notu"].startswith("İçe aktarma geri alındı") for n in notlar))

    def test_v5_oncesi_ice_aktarma_geri_alinamaz(self):
        s = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9)])
        conn = sema.get_connection(self.db)
        conn.execute("UPDATE import_batches SET izli=NULL WHERE id=?", (s["batch_id"],))
        conn.commit()
        conn.close()
        on = ga.onizle(s["batch_id"], db_path=self.db)
        self.assertFalse(on["izin"])
        self.assertIn("v5", on["neden"])

    def test_hatali_ve_olmayan_kayit(self):
        self.assertFalse(ga.onizle(99999, db_path=self.db)["izin"])
        conn = sema.get_connection(self.db)
        conn.execute("INSERT INTO import_batches (dosya_adi, dosya_hash, baslama_zamani, durum, tur) "
                     "VALUES ('x.xlsx','h','2026-09-01 10:00:00','hata','mesai')")
        bid = conn.execute("SELECT MAX(id) FROM import_batches").fetchone()[0]
        conn.commit()
        conn.close()
        self.assertFalse(ga.onizle(bid, db_path=self.db)["izin"])

    def test_basarisiz_geri_alma_hicbir_sey_degistirmez(self):
        s = self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9),
                              (self.tc[1], "AYŞE KARA", D1, "ÇALIŞTI", 9)])
        eski = ga._mesai_uygula

        def bozuk(conn, batch_id, plan, now):
            eski(conn, batch_id, plan, now)
            raise RuntimeError("yapay hata")
        ga._mesai_uygula = bozuk
        try:
            with self.assertRaises(RuntimeError):
                ga.geri_al(s["batch_id"], db_path=self.db)
        finally:
            ga._mesai_uygula = eski
        self.assertIsNotNone(self.gun(self.tc[0], "2026-09-01"))
        self.assertIsNotNone(self.gun(self.tc[1], "2026-09-01"))
        self.assertEqual(self.sql("SELECT durum FROM import_batches WHERE id=?", s["batch_id"])[0]["durum"], "tamamlandi")
        self.assertEqual(len(self.sql("SELECT * FROM puantaj_history WHERE geri_alma=1")), 0)


# ---------------------------------------------------------------- personel bilgileri geri alma (kimlik)

class KimlikGeriAlmaTesti(Temel):
    def test_eklenen_kisiler_silinir_alanlar_geri_doner(self):
        s1 = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        s2 = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G2", "E1", "USTA"),
                                (self.tc[1], "AYŞE KARA", "G1", "E1", "OPERATÖR")])
        self.assertEqual((s2["yeni_personel"], s2["guncellenen_personel"]), (1, 1))
        on = ga.onizle(s2["batch_id"], db_path=self.db)
        self.assertTrue(on["izin"], on["neden"])
        self.assertEqual(on["sayilar"]["silinecek_kisi"], 1)
        self.assertEqual(on["sayilar"]["alan_geri_yuklenecek"], 1)
        sonuc = ga.geri_al(s2["batch_id"], db_path=self.db)
        self.assertEqual(sonuc["ozet"]["silinen_kisi"], 1)
        p0 = pdb.get_personnel(self.tc[0], self.db)
        self.assertEqual(p0["grup"], "G1")
        self.assertEqual(p0["kaynak"], f"veri_merkezi:{Path(self.klasor / '01_kimlik.xlsx').name}")
        self.assertIsNone(pdb.get_personnel(self.tc[1], self.db))
        # silinen kişinin tam kaydı izde saklanır
        sil = self.sql("SELECT * FROM personel_degisiklik_izi WHERE islem='sil'")[0]
        self.assertEqual(json.loads(sil["eski_deger"])["ad_soyad"], "AYŞE KARA")
        # geri alma personel kartındaki geçmişe de yazılır
        h = pdb.get_history(self.tc[0], self.db)
        self.assertTrue(any(x["source"] == f"veri_merkezi:geri_alma#{s2['batch_id']}" for x in h))

    def test_elle_duzeltme_korunur(self):
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        s2 = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G2", "E2", "USTA")])
        pdb.upsert_personnel({"tc": self.tc[0], "ad_soyad": "ALİ VELİ", "grup": "G9"}, source="veri_merkezi:manuel",
                             db_path=self.db)
        on = ga.onizle(s2["batch_id"], db_path=self.db)
        self.assertEqual(on["sayilar"]["alan_elle_degismis"], 1)     # grup
        self.assertEqual(on["sayilar"]["alan_geri_yuklenecek"], 1)   # ekip
        ga.geri_al(s2["batch_id"], db_path=self.db)
        p = pdb.get_personnel(self.tc[0], self.db)
        self.assertEqual((p["grup"], p["bagli_oldugu_ekip"]), ("G9", "E1"))

    def test_puantaji_olan_kisi_silinmez(self):
        s1 = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        self.mesai_yukle([(self.tc[0], "ALİ VELİ", D1, "ÇALIŞTI", 9)])
        on = ga.onizle(s1["batch_id"], db_path=self.db)
        self.assertTrue(on["izin"])
        self.assertEqual(on["sayilar"]["korunan_kisi"], 1)
        ga.geri_al(s1["batch_id"], db_path=self.db)
        self.assertIsNotNone(pdb.get_personnel(self.tc[0], self.db))

    def test_sonraki_kimlik_ice_aktarmasi_engeller(self):
        s1 = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        s2 = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G2", "E1", "USTA")])
        on = ga.onizle(s1["batch_id"], db_path=self.db)
        self.assertFalse(on["izin"])
        self.assertEqual([e["id"] for e in on["engelleyenler"]], [s2["batch_id"]])
        ga.geri_al(s2["batch_id"], db_path=self.db)
        self.assertTrue(ga.onizle(s1["batch_id"], db_path=self.db)["izin"])
        ga.geri_al(s1["batch_id"], db_path=self.db)
        self.assertIsNone(pdb.get_personnel(self.tc[0], self.db))

    def test_etkilenen_personel_geri_almadan_sonra_da_gorunur(self):
        s = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        ga.geri_al(s["batch_id"], db_path=self.db)
        self.assertIn(self.tc[0], sorgu.get_batch_affected_tcs(s["batch_id"], db_path=self.db))


# ---------------------------------------------------------------- yedekleme

class YedekTesti(Temel):
    def test_yedek_al_listele_incele(self):
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        y = yedek.yedek_al("elle", self.db)
        self.assertTrue(Path(y["yol"]).exists())
        self.assertFalse(Path(y["yol"] + "-wal").exists(), "yedek tek başına bir dosya olmalı")
        inc = yedek.yedek_incele(y["yol"])
        self.assertTrue(inc["gecerli"], inc["hata"])
        self.assertEqual(inc["personel"], 1)
        self.assertEqual(inc["sema"], 5)
        self.assertIn(y["ad"], [x["ad"] for x in yedek.yedekleri_listele(self.db)])

    def test_ice_aktarma_oncesi_yedek(self):
        s = self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        self.assertTrue(s["yedek"])
        inc = yedek.yedek_incele(yedek.yedek_bul(s["yedek"], self.db))
        self.assertEqual(inc["personel"], 0, "yedek içe aktarmadan ÖNCEKİ hâli içermeli")

    def test_geri_yukle(self):
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        y = yedek.yedek_al("elle", self.db)
        self.kimlik_yukle([(self.tc[1], "AYŞE KARA", "G1", "E1", "OPERATÖR")])
        self.assertIsNotNone(pdb.get_personnel(self.tc[1], self.db))
        sonuc = yedek.geri_yukle(y["yol"], self.db)
        self.assertIsNone(pdb.get_personnel(self.tc[1], self.db))
        self.assertIsNotNone(pdb.get_personnel(self.tc[0], self.db))
        self.assertTrue(sonuc["onceki_yedek"])
        # geri yüklemeden önceki hâl de yedeklendi -> o yedekte 2 kişi var
        self.assertEqual(yedek.yedek_incele(yedek.yedek_bul(sonuc["onceki_yedek"], self.db))["personel"], 2)
        self.assertEqual(sema.gunluk_kipi(self.db), "wal")
        self.assertEqual(sema.get_schema_version(self.db), 5)

    def test_eski_semali_yedek_geri_yuklenince_yukseltilir(self):
        eski = self.klasor / "eski.db"
        conn = sqlite3.connect(eski)
        conn.executescript(pdb.SCHEMA)
        conn.execute("CREATE TABLE schema_meta (key TEXT PRIMARY KEY, value TEXT)")
        conn.execute("INSERT INTO schema_meta VALUES ('version','3')")
        conn.commit()
        conn.close()
        yedek.geri_yukle(eski, self.db)
        self.assertEqual(sema.get_schema_version(self.db), 5)
        self.assertIn("puantaj_history", {r["name"] for r in self.sql("SELECT name FROM sqlite_master")})

    def test_bozuk_yedek_reddedilir_ve_veri_degismez(self):
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])
        bozuk = self.klasor / "bozuk.db"
        bozuk.write_bytes(b"bu bir veritabani degil" * 100)
        with self.assertRaises(yedek.YedekHatasi):
            yedek.geri_yukle(bozuk, self.db)
        yabanci = self.klasor / "yabanci.db"
        c = sqlite3.connect(yabanci)
        c.execute("CREATE TABLE baska (x)")
        c.commit()
        c.close()
        with self.assertRaises(yedek.YedekHatasi):
            yedek.geri_yukle(yabanci, self.db)
        self.assertIsNotNone(pdb.get_personnel(self.tc[0], self.db))

    def test_saklama_ve_elle_yedekler(self):
        for _ in range(4):
            yedek.yedek_al("ice_aktarma", self.db)
        for _ in range(2):
            yedek.yedek_al("elle", self.db)
        silinen = yedek.temizle(self.db, saklama=2)
        self.assertEqual(len(silinen), 2)
        liste = yedek.yedekleri_listele(self.db)
        self.assertEqual(sum(1 for y in liste if y["tur"] == "ice_aktarma"), 2)
        self.assertEqual(sum(1 for y in liste if y["tur"] == "elle"), 2)

    def test_acilis_yedegi_gunde_bir(self):
        self.assertIsNotNone(yedek.acilis_yedegi(self.db))
        self.assertIsNone(yedek.acilis_yedegi(self.db))
        ayarlar.kaydet({"acilis_yedegi": False}, uygula_=False)
        self.assertIsNone(yedek.acilis_yedegi(self.db, bugun=datetime.date.today() + datetime.timedelta(days=1)))

    def test_yedek_bul_klasor_disina_cikamaz(self):
        self.assertIsNone(yedek.yedek_bul("../personel_veritabani.db", self.db))
        self.assertIsNone(yedek.yedek_bul("/etc/passwd", self.db))
        self.assertIsNone(yedek.yedek_bul("rastgele.db", self.db))


# ---------------------------------------------------------------- ayarlar

class AyarTesti(Temel):
    def test_varsayilanlar(self):
        self.assertEqual(ayarlar.al("yillik_fm_siniri"), 270.0)
        self.assertEqual(ayarlar.al("proje_adi"), "MAOG Projesi")

    def test_kaydet_dogrula_uygula(self):
        import veri_merkezi_verimlilik as vm
        import veri_merkezi_mesai_formu as mf
        ok, hatalar = ayarlar.kaydet({"yillik_fm_siniri": "250", "proje_adi": "  Test   Projesi ",
                                      "fm_katsayi_haftaici": "1,5", "devamsizlik_esigi": 4})
        self.assertTrue(ok, hatalar)
        self.assertEqual(vm.YILLIK_FM_SINIRI, 250.0)
        self.assertEqual(vm.DEVAMSIZLIK_ESIGI, 4)
        self.assertEqual(mf.PROJE_ADI, "Test Projesi")
        self.assertEqual(vm.etiket_yillik(), "250")
        self.assertIn("uygulanan eşik 250", vm.esik_notu(vm.YILLIK_FM_SINIRI, 270))
        veri = json.loads(ayarlar.AYAR_DOSYASI.read_text(encoding="utf-8"))
        self.assertEqual(veri["yillik_fm_siniri"], 250.0)

    def test_gecersiz_deger_hicbir_sey_yazmaz(self):
        ok, hatalar = ayarlar.kaydet({"yillik_fm_siniri": "abc", "proje_adi": "Yeni", "devamsizlik_esigi": 2.5,
                                      "gunluk_azami_saat": 30})
        self.assertFalse(ok)
        self.assertEqual(set(hatalar), {"yillik_fm_siniri", "devamsizlik_esigi", "gunluk_azami_saat"})
        self.assertFalse(ayarlar.AYAR_DOSYASI.exists())
        self.assertEqual(ayarlar.al("proje_adi"), "MAOG Projesi")

    def test_bozuk_dosya_varsayilana_duser(self):
        ayarlar.AYAR_DOSYASI.write_text("{bozuk json", encoding="utf-8")
        self.assertEqual(ayarlar.al("yillik_fm_siniri"), 270.0)
        ayarlar.AYAR_DOSYASI.write_text(json.dumps({"yillik_fm_siniri": -5, "proje_adi": "X", "bilinmeyen": 1}),
                                        encoding="utf-8")
        self.assertEqual(ayarlar.al("yillik_fm_siniri"), 270.0)
        self.assertEqual(ayarlar.al("proje_adi"), "X")


# ---------------------------------------------------------------- çevrimdışı varlıklar

class VarlikTesti(Temel):
    def test_indir_sahte_ag_ile(self):
        import veri_merkezi_varliklar as va
        sayfalar = {
            "https://cdn.example/bs/css/bootstrap.min.css": b".a{background:url(data:image/svg+xml,x)}",
            "https://cdn.example/icons/font/bootstrap-icons.min.css":
                b'@font-face{src:url("fonts/bootstrap-icons.woff2?abc") format("woff2"),'
                b'url(\'fonts/bootstrap-icons.woff?abc\') format("woff")}',
            "https://cdn.example/icons/font/fonts/bootstrap-icons.woff2?abc": b"W2",
            "https://cdn.example/icons/font/fonts/bootstrap-icons.woff?abc": b"W1",
            "https://fonts.example/css2?family=X": b"/* latin */@font-face{src:url(https://gs.example/s/archivo/v1/a.woff2)}"
                                                   b"@font-face{src:url(https://gs.example/s/plex/v2/a.woff2)}",
            "https://gs.example/s/archivo/v1/a.woff2": b"A",
            "https://gs.example/s/plex/v2/a.woff2": b"P",
        }

        def getir(url, zaman_asimi):
            if url not in sayfalar:
                raise OSError(f"yok: {url}")
            return sayfalar[url]
        kaynaklar = [("bootstrap", "https://cdn.example/bs/css/bootstrap.min.css"),
                     ("ikonlar", "https://cdn.example/icons/font/bootstrap-icons.min.css"),
                     ("yazitipleri", "https://fonts.example/css2?family=X")]
        hedef = self.klasor / "yerel_varliklar"
        rapor = va.indir(hedef=hedef, getir=getir, kaynaklar=kaynaklar)
        self.assertTrue(all(r["ok"] for r in rapor), rapor)
        ikon_css = (hedef / "ikonlar" / "stil.css").read_text(encoding="utf-8")
        self.assertNotIn("https://", ikon_css)
        self.assertIn("url(\"dosya/", ikon_css)
        yt = (hedef / "yazitipleri" / "stil.css").read_text(encoding="utf-8")
        self.assertNotIn("gs.example", yt)
        dosyalar = sorted(p.name for p in (hedef / "yazitipleri" / "dosya").iterdir())
        self.assertEqual(len(dosyalar), 2, "aynı ada sahip iki yazı tipi dosyası birbirini ezmemeli")
        d = va.durum(hedef)
        self.assertTrue(all(x["yerel"] for x in d))
        stiller = va.stil_listesi(hedef, kaynaklar)
        self.assertTrue(all(s.startswith("/yerel/") for s in stiller))

    def test_ag_hatasinda_eskisi_korunur(self):
        import veri_merkezi_varliklar as va
        hedef = self.klasor / "yerel_varliklar"
        kaynaklar = [("bootstrap", "https://cdn.example/bs.css")]
        va.indir(hedef=hedef, getir=lambda u, z: b".a{}", kaynaklar=kaynaklar)

        def hata(u, z):
            raise OSError("internet yok")
        rapor = va.indir(hedef=hedef, getir=hata, kaynaklar=kaynaklar)
        self.assertFalse(rapor[0]["ok"])
        self.assertEqual((hedef / "bootstrap" / "stil.css").read_text(encoding="utf-8"), ".a{}")

    def test_yerel_yoksa_cdn(self):
        import veri_merkezi_varliklar as va
        kaynaklar = [("bootstrap", "https://cdn.example/bs.css")]
        self.assertEqual(va.stil_listesi(self.klasor / "yok", kaynaklar), ["https://cdn.example/bs.css"])


# ---------------------------------------------------------------- v5.1: hafta tatili / pazar kuralı

import veri_merkezi_hafta_kurali as hk  # noqa: E402

PZT = datetime.date(2026, 9, 7)          # pazartesi; pazar = 13.09.2026


def gun(i):
    return PZT + datetime.timedelta(days=i)


class HaftaKuraliTesti(Temel):
    def setUp(self):
        super().setUp()
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA")])

    def hafta(self, durumlar):
        """durumlar: 7 öğe (pzt..paz); sayı = çalışılan toplam saat, metin = durum, None = kayıt yok."""
        satir = []
        for i, v in enumerate(durumlar):
            if v is None:
                continue
            if isinstance(v, (int, float)):
                satir.append((self.tc[0], "ALİ VELİ", gun(i), "ÇALIŞTI", v))
            else:
                satir.append((self.tc[0], "ALİ VELİ", gun(i), v, 0))
        self.mesai_yukle(satir)
        kural = hk.kural_uygula(hk.db_kayitlari(gun(0).isoformat(), gun(6).isoformat(), db_path=self.db))
        return {i: kural.get((self.tc[0], gun(i).isoformat())) for i in range(7)}

    def toplam(self, bas=0, bit=6):
        return hk.db_toplami(gun(bas).isoformat(), gun(bit).isoformat(), db_path=self.db)

    def test_alti_gun_tam_pazar_katli(self):
        k = self.hafta([9, 9, 9, 9, 9, 9, 11])
        self.assertEqual(k[6], ("calisti", True))
        t = self.toplam()
        self.assertAlmostEqual(t["yevmiye"], (54 + 11 * 2.5) / 9)
        self.assertEqual(t["pazar_katli"], 11)

    def test_cuma_ucretsiz_pazar_calisti_cuma_hafta_tatiline_doner(self):
        k = self.hafta([9, 9, 9, 9, "ÜCRETSİZ İZİN", 9, 11])
        self.assertEqual(k[4], ("hafta_tatili", False))
        self.assertEqual(k[6], ("calisti", False))
        t = self.toplam()
        self.assertAlmostEqual(t["yevmiye"], (54 + 2 * 1.5) / 9)   # pazar normal gün: 9 normal + 2 FM ×1,5
        self.assertEqual(t["pazar_katli"], 0)

    def test_cuma_ucretsiz_pazar_calismadi_ikisi_de_ucretsiz(self):
        k = self.hafta([9, 9, 9, 9, "ÜCRETSİZ İZİN", 9, "HAFTA TATİLİ"])
        self.assertEqual(k[4], ("ucretsiz_izin", False))
        self.assertEqual(k[6], ("ucretsiz_izin", False))

    def test_raporlu_gunler_calismis_sayilir(self):
        k = self.hafta(["SAĞLIK RAPORU", "SAĞLIK RAPORU", "SAĞLIK RAPORU", 9, 9, 9, 11])
        self.assertEqual(k[6], ("calisti", True))
        self.assertEqual(k[0], ("rapor", False))
        self.assertAlmostEqual(self.toplam()["yevmiye"], (27 + 11 * 2.5) / 9)

    def test_babalik_vefat_yillik_izin_calismis_sayilir(self):
        k = self.hafta(["BABALIK İZNİ", "VEFAT İZNİ", "YILLIK İZİN", 9, 9, 9, 9])
        self.assertEqual(k[6], ("calisti", True))

    def test_devamsizlik_da_mazeretsiz(self):
        k = self.hafta([9, "DEVAMSIZ", 9, 9, 9, 9, 9])
        self.assertEqual(k[1], ("hafta_tatili", False))
        self.assertEqual(k[6], ("calisti", False))

    def test_iki_eksik_gun_hicbiri_donmez(self):
        k = self.hafta([9, 9, "ÜCRETSİZ İZİN", 9, "ÜCRETSİZ İZİN", 9, 9])
        self.assertEqual(k[2], ("ucretsiz_izin", False))
        self.assertEqual(k[4], ("ucretsiz_izin", False))
        self.assertEqual(k[6], ("calisti", False))

    def test_hafta_ici_hafta_tatili_varsa_pazar_normal(self):
        k = self.hafta([9, 9, "HAFTA TATİLİ", 9, 9, 9, 9])
        self.assertEqual(k[6], ("calisti", False))
        self.assertEqual(k[2], ("hafta_tatili", False))

    def test_kayit_olmayan_gun_kurali_tetiklemez(self):
        k = self.hafta([None, None, 9, 9, 9, 9, 9])      # çarşamba işe başladı
        self.assertEqual(k[6], ("calisti", True))

    def test_donem_siniri_haftanin_tamamina_bakar(self):
        self.hafta(["ÜCRETSİZ İZİN", 9, 9, 9, 9, 9, 11])
        t = self.toplam(bas=2)                              # dönem çarşamba başlıyor; pazartesi dönem dışı
        self.assertEqual(t["pazar_katli"], 0, "dönem dışındaki ücretsiz izin de pazarı etkilemeli")
        self.assertAlmostEqual(t["yevmiye"], (45 + 2 * 1.5) / 9)   # çar-cmt 36 + pazar 9 normal, 2 FM ×1,5

    def test_verimlilik_ve_ekranlar_ayni_sonucu_verir(self):
        import veri_merkezi_verimlilik as vm
        self.hafta([9, 9, 9, 9, "ÜCRETSİZ İZİN", 9, 11])
        d = vm.analiz(gun(0).isoformat(), gun(6).isoformat(), db_path=self.db)
        self.assertAlmostEqual(d["genel"]["yevmiye"], self.toplam()["yevmiye"])
        self.assertEqual(d["genel"]["gun"]["hafta_tatili"], 1)
        self.assertEqual(d["genel"]["gun"]["ucretsiz_izin"], 0)
        self.assertEqual(d["genel"]["pazar_katli"], 0)
        kisi = sorgu.get_person_period_totals(self.tc[0], gun(0).isoformat(), gun(6).isoformat(), db_path=self.db)
        self.assertAlmostEqual(kisi["yevmiye_gunu"], round(self.toplam()["yevmiye"], 1))
        ozet = sorgu.get_period_summary(gun(0).isoformat(), gun(6).isoformat(), db_path=self.db)
        self.assertAlmostEqual(ozet["yevmiye"], self.toplam()["yevmiye"])

    def test_verimlilik_katli_pazar(self):
        import veri_merkezi_verimlilik as vm
        self.hafta([9, 9, 9, 9, 9, 9, 11])
        g = vm.analiz(gun(0).isoformat(), gun(6).isoformat(), db_path=self.db)["genel"]
        self.assertAlmostEqual(g["yevmiye"], (54 + 27.5) / 9)
        self.assertAlmostEqual(g["fm_prim"], (54 + 27.5) / 9 - 65 / 9)


# ---------------------------------------------------------------- v5.2: İK bilgileri ve Personel Kartı

import veri_merkezi_ik as ik  # noqa: E402
import veri_merkezi_personel_karti as pk  # noqa: E402


def ik_dosyasi(yol, satirlar):
    """satirlar: [{başlık: değer}] - eksik başlıklar boş."""
    wb = openpyxl.Workbook()
    ws = wb.active
    basliklar = ["TC Kimlik No"] + [h for h, _s, _t in ik.ESLEME]
    ws.append(basliklar)
    for s in satirlar:
        ws.append([s.get(h) for h in basliklar])
    wb.save(yol)
    return yol


class IkTesti(Temel):
    def test_ice_aktar_yalniz_bos_alanlar_ve_eslesenler(self):
        self.kimlik_yukle([(self.tc[0], "ALİ VELİ", "G1", "E1", "USTA"), (self.tc[1], "AYŞE KARA", "G1", "E1", "OP")])
        conn = sema.get_connection(self.db)
        conn.execute("UPDATE personnel SET ise_giris_tarihi=NULL WHERE tc=?", (self.tc[0],))
        conn.commit()
        conn.close()
        eski_giris = pdb.get_personnel(self.tc[1], self.db)["ise_giris_tarihi"]
        ortak = {"Adı": "ALİ", "Soyadı": "VELİ", "Kan Grubu": "0 RH+", "Telefon No": "05000000000",
                 "Acil Ad Soyad": "EŞE VELİ- ANNESİ", "SGK Giriş Tarihi": datetime.datetime(2025, 3, 1),
                 "İşe giriş trh.": datetime.datetime(2025, 3, 3), "Sözleşme": "X", "Normal Kadro": "true"}
        yol = ik_dosyasi(self.klasor / "ik.xlsx", [
            dict(ortak, **{"TC Kimlik No": self.tc[0], "Çalışan Statüsü": "Etkin", "Görev": "USTA"}),
            dict(ortak, **{"TC Kimlik No": self.tc[0], "Çalışan Statüsü": "İşten Çıktı", "Görev": "DÜZ İŞÇİ",
                           "İşe giriş trh.": datetime.datetime(2024, 1, 1)}),
            dict(ortak, **{"TC Kimlik No": self.tc[1], "Çalışan Statüsü": "Etkin",
                           "SGK Giriş Tarihi": datetime.datetime(2020, 1, 1)}),
            dict(ortak, **{"TC Kimlik No": gecerli_tc(55), "Çalışan Statüsü": "Etkin"}),   # veritabanında yok
        ])
        o = ik.ice_aktar(yol, self.db)
        self.assertEqual((o["eslesen"], o["giris_doldurulan"]), (2, 1))
        self.assertEqual(pdb.get_personnel(self.tc[0], self.db)["ise_giris_tarihi"], "2025-03-01")
        self.assertEqual(pdb.get_personnel(self.tc[1], self.db)["ise_giris_tarihi"], eski_giris, "dolu alan değişmemeli")
        k = ik.getir(self.tc[0], self.db)
        self.assertEqual(k["calisan_statusu"], "Etkin")
        self.assertEqual(k["kayit_sayisi"], 2)
        self.assertEqual(k["onceki_kayitlar"][0]["gorev"], "DÜZ İŞÇİ")
        self.assertEqual(k["kan_grubu"], "0 Rh+")
        self.assertIsNone(k["telefon"], "yer tutucu telefon boş bırakılmalı")
        self.assertEqual((k["evrak_sozlesme"], k["evrak_diploma"], k["normal_kadro"]), (1, 0, 1))
        self.assertIsNone(ik.getir(gecerli_tc(55), self.db))
        self.assertTrue(any(h["source"].startswith("ik_listesi:") for h in pdb.get_history(self.tc[0], self.db)))
        ik.sil(self.tc[0], self.db)
        self.assertIsNone(ik.getir(self.tc[0], self.db))

    def test_kart_yardimcilari(self):
        b = datetime.date(2026, 9, 25)
        self.assertEqual(pk.tarih_durumu("2026-09-01", b)[0], "gecti")
        self.assertEqual(pk.tarih_durumu("2026-10-10", b)[0], "yakin")
        self.assertEqual(pk.tarih_durumu("2027-10-10", b)[0], "tamam")
        self.assertEqual(pk.tarih_durumu("9999-12-31", b)[0], "suresiz")
        self.assertEqual(pk.tarih_durumu(None, b)[0], "yok")
        self.assertEqual(pk.sure_metni("2024-06-30", b), "2 yıl 2 ay")
        self.assertEqual(pk.sure_metni("2026-09-20", b), "5 gün")
        self.assertEqual(pk.yas("1995-09-26", b), 30)
        self.assertEqual(pk._acil_kisi("EŞE AYAR- ANNESİ"), ("EŞE AYAR", "annesi"))
        self.assertEqual(pk._acil_kisi("ASAD FAYZILLOEV ABİSİ"), ("ASAD FAYZILLOEV", "abisi"))
        self.assertEqual(pk._acil_kisi("KSK"), ("KSK", None))
        self.assertEqual(pk._tel("05321234567"), "0532 123 45 67")


if __name__ == "__main__":
    unittest.main()
