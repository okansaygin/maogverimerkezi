# -*- coding: utf-8 -*-
"""
VERİ MERKEZİ - MESAİ (ÇALIŞMA SAATİ) İÇE AKTARMA (veri_merkezi_mesai_ice_aktarma.py)
========================================================================================
Personel BİLGİLERİ (kimlik) ile personel ÇALIŞMA BİLGİLERİ (mesai saati) ARTIK
İKİ AYRI kaynaktan gelir:
  - Kimlik (ad soyad, grup, alt ekip, görev, giriş/çıkış tarihi):
    aylık, çok-gün-sütunlu ESKİ format -> veri_merkezi_ice_aktarma.py
  - Çalışma saati (BU DOSYA): GÜNLÜK SAP/Rönesans roster excel'i - HER SATIR
    bir kişinin BİR GÜNÜNÜ temsil eder (çoklu gün sütunu YOK, çoklu-gün
    formatının aksine). Sütunlar SABİT HARF DEĞİL, başlık metnine göre
    bulunur:
        - "TC Kimlik No"                      -> tc
        - "Tarih"                             -> tarih (TAM eşleşme aranır;
          "İşe Giriş Tarihi" gibi başka başlıklar da "TARİH" içerdiği için
          alt-dize eşleşmesi YANLIŞ sütunu seçebilir)
        - "Çalışma Durumu"                    -> durum (ÇALIŞTI/İZİNLİ/RAPORLU..)
          ("Taşeron Çalışma Durumu" ile KARIŞTIRILMAZ - "TAŞERON" hariç tutulur)
        - "Rönesans Çalışma Saati Girişi"      -> toplam_saat (o günün NORMAL+
          FAZLA TOPLAMI - kaynakta zaten toplanmış tek değer)

Toplam saat, standart 9 saatlik iş günü baz alınarak normal/fazla mesaiye
bölünür: normal_saat = min(toplam, 9), fazla_saat = max(0, toplam-9). Bu,
team_report.py / cost_report.py'nin kullandığı 9 saatlik standart iş günü
varsayımıyla TUTARLIDIR.

Kalıcı depoda personel BAŞINA GÜN BAŞINA TEK SATIR tutulur (gunluk_puantaj
tablosu, UNIQUE(tc, tarih)) - normal_saat ve fazla_saat AYNI satırda, ayrı
sütunlar olarak (kişiler asla 2 satırda görünmez).

İki aşamalı: preview_mesai_import() SADECE okur; commit_mesai_import() TEK
transaction'da kalıcı depoya yazar (personel bilgisi importundakiyle aynı
"ya hep ya hiç" garantisi).

v5: yazmadan önce veritabanının yedeği alınır (veri_merkezi_yedek) ve eklenen /
değişen HER gün puantaj_history tablosuna eski ve yeni değeriyle yazılır. Böylece
bir günün değeri değiştiğinde eskisi kaybolmaz ve içe aktarma geri alınabilir
(veri_merkezi_geri_alma).
"""

__version__ = "2026-09-25.1"

import datetime
from pathlib import Path

import openpyxl

import puantaj_engine as engine
import veri_merkezi_sema as sema
import veri_merkezi_dogrulama as dogrulama
from veri_merkezi_ice_aktarma import (
    ImportError_, _file_hash, _pick_sheet, _existing_file_hashes,
)

STANDART_GUN_SAATI = 9.0

# NOT (Türkçe İ/ı/I tuzağı): Python'un std str.upper()'ı Türkçe kurallarına
# göre DEĞİL, Unicode'un genel kuralına göre büyütür - küçük "i" HER ZAMAN
# noktasız "I"ya (ASCII) dönüşür, noktalı "İ"ye DEĞİL. Kaynak başlık TÜMÜ
# BÜYÜK HARFLE yazılmışsa (eski aylık kaynak gibi) bu sorun yaratmaz (İ zaten
# İ kalır); ama bu YENİ dosyadaki gibi "Tarih", "İşe Giriş Tarihi" gibi
# Baş Harf Büyük başlıklarda "i" -> ASCII "I" olur. Bu yüzden başlık
# metnindeki TÜM İ/ı/I varyantlarını TEK bir harfe (ASCII I) katlayıp öyle
# karşılaştırıyoruz - hangi yazım stiliyle gelirse gelsin doğru eşleşsin diye.
_I_KATLA = str.maketrans({"İ": "I", "ı": "I", "I": "I"})


def _baslik_normalize(v):
    return dogrulama.normalize_header(v).translate(_I_KATLA)


# Alt-dize kuralları (any_of / none_of) - "TARIH" hariç, o TAM eşleşmeyle
# ayrıca aranır (aşağıya bkz.) çünkü "İşe Giriş Tarihi" gibi başlıklar da
# "TARIH" alt-dizesini içerir. Tüm literaller ASCII "I" ile yazılmıştır
# (yukarıdaki katlama sayesinde İ/ı/I hepsi buna eşlenir).
FIELD_HEADER_RULES = [
    ("tc", ["TC KIMLIK NO"], []),
    ("ad_soyad", ["ADI SOYADI", "AD SOYAD"], []),
    ("durum", ["ÇALIŞMA DURUMU"], ["TAŞERON"]),
    ("toplam_saat", ["RÖNESANS ÇALIŞMA SAATI", "ÇALIŞMA SAATI GIRIŞI"], ["TAŞERON", "PDKS"]),
]
TARIH_TAM_ESLESME = "TARIH"
REQUIRED_FIELDS = ["tc", "toplam_saat", "tarih"]


def build_column_map(ws, header_rows=(1, 2, 3)):
    """Başlık metnine göre {alan: sütun_no}. 'Tarih' TAM eşleşmeyle, diğerleri
    alt-dize (any_of/none_of) ile aranır - team_report/veri_merkezi_dogrulama
    ile aynı felsefe: SABİT sütun harfi kullanılmaz."""
    texts = {}
    for r in header_rows:
        for c in range(1, ws.max_column + 1):
            if c not in texts:
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    texts[c] = _baslik_normalize(v)

    colmap = {}
    claimed = set()
    for c in sorted(texts):
        if texts[c] == TARIH_TAM_ESLESME:
            colmap["tarih"] = c
            claimed.add(c)
            break
    for key, any_of, none_of in FIELD_HEADER_RULES:
        for c in sorted(texts):
            if c in claimed:
                continue
            t = texts[c]
            if any(a in t for a in any_of) and not any(n in t for n in none_of):
                colmap[key] = c
                claimed.add(c)
                break
    return colmap


def kural_baslik_eslesmesi(colmap):
    """ENGELLEYİCİ: TC/Tarih/Toplam Saat sütunlarından biri bulunamazsa içe
    aktarma tamamen durur."""
    eksikler = [f for f in REQUIRED_FIELDS if f not in colmap]
    if eksikler:
        etiket = {"tc": "TC Kimlik No", "toplam_saat": "Rönesans Çalışma Saati Girişi",
                   "tarih": "Tarih"}
        return [{"tip": "baslik_eslesmedi", "onem": "engelleyici", "tc": None,
                   "detay": f"Mesai dosyasında şu sütunların başlığı bulunamadı: "
                            f"{', '.join(etiket.get(f, f) for f in eksikler)}. "
                            f"Dosya yapısı beklenenden farklı olabilir."}]
    return []


def _parse_date_cell(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def _to_number(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        v = v.strip().replace(",", ".")
        try:
            return float(v)
        except ValueError:
            return None
    return None


def preview_mesai_import(dosya_yolu, sheet_name=None, db_path=None, max_saat=24, colmap_override=None):
    """Dosyayı okur, doğrular, SADECE bir önizleme sözlüğü döner. Veritabanına
    hiçbir şey yazmaz."""
    sema.migrate(db_path)
    import personnel_db as pdb
    pdb.init_db(db_path)  # "personnel" tablosu yoksa oluşturur (sadece YAPI, satır yazmaz)
    dosya_yolu = Path(dosya_yolu)
    dosya_hash = _file_hash(dosya_yolu)

    wb = openpyxl.load_workbook(dosya_yolu, data_only=True)
    ws = _pick_sheet(wb, sheet_name)

    colmap = build_column_map(ws)
    if colmap_override:            # kullanıcının elle seçtiği sütunlar başlık eşleşmesine üstün gelir
        colmap.update({k: v for k, v in colmap_override.items() if v})
    engelleyiciler = list(kural_baslik_eslesmesi(colmap))
    dosya_tekrar = dogrulama.kural_dosya_tekrar(dosya_hash, _existing_file_hashes(db_path))
    engelleyiciler.extend(a for a in dosya_tekrar if a["onem"] == "engelleyici")

    if engelleyiciler:
        return {
            "durum": "engellendi", "dosya_adi": dosya_yolu.name, "dosya_hash": dosya_hash,
            "sheet_name": ws.title, "engelleyiciler": engelleyiciler, "uyarilar": [],
            "gunluk_kayitlar": [], "taranan_satir": 0,
            "donem_baslangic": None, "donem_bitis": None,
        }

    conn = sema.get_connection(db_path)
    try:
        bilinen_tcler = {r["tc"] for r in conn.execute("SELECT tc FROM personnel").fetchall()}
    finally:
        conn.close()

    uyarilar = []
    kayit_index = {}          # (tc, tarih_iso) -> kayıt - dosya İÇİ tekrar tespiti
    tcsiz_satirlar = []
    tcsiz_isimler = set()
    bilinmeyen_tc = set()     # (tc, ad) - personel tablosunda yok, bilgilendirme
    taranan_satir = 0

    for r in range(2, ws.max_row + 1):
        raw_tc = ws.cell(row=r, column=colmap["tc"]).value
        raw_saat = ws.cell(row=r, column=colmap["toplam_saat"]).value
        name = ws.cell(row=r, column=colmap["ad_soyad"]).value if "ad_soyad" in colmap else None
        name = str(name).strip() if name else None

        if raw_tc in (None, "") and raw_saat in (None, ""):
            continue  # tamamen boş satır - veri değil
        taranan_satir += 1

        tc = engine.normalize_tc(raw_tc)
        tarih = _parse_date_cell(ws.cell(row=r, column=colmap["tarih"]).value)
        durum = ws.cell(row=r, column=colmap["durum"]).value if "durum" in colmap else None
        durum = str(durum).strip() if durum else None
        toplam = _to_number(raw_saat)

        if not tc:
            tcsiz_satirlar.append(r)
            if name:
                tcsiz_isimler.add(name)
            uyarilar.append({"tip": "tc_eksik", "onem": "uyari", "tc": None, "ad_soyad": name,
                              "kaynak_satir": r,
                              "detay": f"Satır {r}: {name or '(isim de yok)'} - TC boş, bu satırın "
                                       f"saati kimseye atfedilemedi."})
            continue

        issue = engine.check_tc(tc, "mesai", r, name)
        if issue:
            uyarilar.append({"tip": "tc_format_hatali", "onem": "uyari", "tc": tc, "ad_soyad": name,
                              "kaynak_satir": r,
                              "detay": f"Satır {r}: TC '{tc}' ({name or '?'}) -> {issue['reason']}"})

        if tc not in bilinen_tcler:
            bilinmeyen_tc.add((tc, name))

        if tarih is None:
            uyarilar.append({"tip": "tarihsiz_satir", "onem": "uyari", "tc": tc, "ad_soyad": name,
                              "kaynak_satir": r,
                              "detay": f"Satır {r} (TC {tc}): 'Tarih' sütunu boş/okunamadı, "
                                       f"bu satırın saati kaydedilmedi."})
            continue

        if toplam is None:
            toplam = 0.0
        if toplam < 0 or toplam > max_saat:
            uyarilar.append({"tip": "anormal_saat", "onem": "uyari", "tc": tc, "ad_soyad": name,
                              "kaynak_satir": r,
                              "detay": f"Satır {r}: {name or tc} - {tarih.isoformat()} tarihinde "
                                       f"{toplam} saat (beklenen aralık: 0-{max_saat})."})

        normal_saat = round(min(toplam, STANDART_GUN_SAATI), 2)
        fazla_saat = round(max(0.0, toplam - STANDART_GUN_SAATI), 2)

        key = (tc, tarih.isoformat())
        if key in kayit_index:
            onceki = kayit_index[key]
            uyarilar.append({"tip": "gunluk_kayit_cakismasi", "onem": "uyari", "tc": tc, "ad_soyad": name,
                              "kaynak_satir": r,
                              "detay": f"TC {tc} ({name or '(isimsiz)'}), {tarih.isoformat()}: satır "
                                       f"{onceki['kaynak_satir']} ile satır {r} aynı gün için "
                                       f"tekrar ediyor - bu dosyadaki SON değer ({toplam} sa) esas alındı."})
        kayit_index[key] = {
            "tc": tc, "tarih": tarih.isoformat(), "normal_saat": normal_saat,
            "fazla_saat": fazla_saat, "durum": durum, "kaynak_satir": r, "ad_soyad": name,
        }

    if tcsiz_satirlar:
        uyarilar.append({"tip": "tcsiz_saat_ozeti", "onem": "uyari", "tc": None,
                          "ad_soyad": ", ".join(sorted(tcsiz_isimler)) or None, "kaynak_satir": None,
                          "detay": f"TC'si olmayan {len(tcsiz_satirlar)} satırın saati kimseye "
                                   f"atfedilemedi (satırlar: {sorted(tcsiz_satirlar)[:20]}"
                                   f"{'...' if len(tcsiz_satirlar) > 20 else ''})."})

    for tc, ad in sorted(bilinmeyen_tc, key=lambda x: x[0]):
        uyarilar.append({"tip": "personel_kaydi_yok", "onem": "uyari", "tc": tc, "ad_soyad": ad,
                          "kaynak_satir": None,
                          "detay": f"TC {tc} ({ad or '?'}) 'Personel Bilgileri' veri setinde kayıtlı "
                                   f"değil - saatler yine de yazıldı, ama isim/grup/görev boş "
                                   f"görünecek. Personel Bilgileri içe aktarımı yapın veya "
                                   f"Personel Düzenle'den elle ekleyin."})

    gunluk_kayitlar = list(kayit_index.values())
    tarihler = [k["tarih"] for k in gunluk_kayitlar]

    return {
        "durum": "hazir", "dosya_adi": dosya_yolu.name, "dosya_hash": dosya_hash,
        "sheet_name": ws.title, "engelleyiciler": [], "uyarilar": uyarilar,
        "gunluk_kayitlar": gunluk_kayitlar, "taranan_satir": taranan_satir,
        "donem_baslangic": min(tarihler) if tarihler else None,
        "donem_bitis": max(tarihler) if tarihler else None,
    }


def commit_mesai_import(preview, db_path=None):
    """preview_mesai_import()'un ürettiği veriyi TEK bir transaction'da kalıcı
    depoya yazar (gunluk_puantaj: kişi+gün başına TEK satır, upsert). Ya
    TAMAMEN başarılı olur ya da hiçbir şey yazılmaz."""
    if preview["durum"] != "hazir":
        raise ImportError_("Bu önizleme içe aktarılamaz durumda (engelleyici sorunlar var).")

    sema.migrate(db_path)
    import veri_merkezi_yedek as yedek
    yedek_bilgisi = yedek.islem_oncesi_yedek("ice_aktarma", db_path)
    conn = sema.get_connection(db_path)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    try:
        cur = conn.execute(
            "INSERT INTO import_batches (dosya_adi, dosya_hash, donem_baslangic, donem_bitis, "
            "baslama_zamani, durum, tur, izli) VALUES (?,?,?,?,?, 'basladi', 'mesai', 1)",
            (preview["dosya_adi"], preview["dosya_hash"], preview["donem_baslangic"],
             preview["donem_bitis"], now),
        )
        batch_id = cur.lastrowid

        yeni = guncellenen = degismeyen = 0
        for kayit in preview["gunluk_kayitlar"]:
            mevcut = conn.execute(
                "SELECT normal_saat, fazla_saat, durum, kaynak_satir, import_batch_id FROM gunluk_puantaj "
                "WHERE tc=? AND tarih=?",
                (kayit["tc"], kayit["tarih"]),
            ).fetchone()
            if mevcut is None:
                conn.execute(
                    "INSERT INTO gunluk_puantaj (tc, tarih, normal_saat, fazla_saat, durum, "
                    "kaynak_satir, import_batch_id) VALUES (?,?,?,?,?,?,?)",
                    (kayit["tc"], kayit["tarih"], kayit["normal_saat"], kayit["fazla_saat"],
                     kayit["durum"], kayit["kaynak_satir"], batch_id),
                )
                conn.execute(
                    "INSERT INTO puantaj_history (tc, tarih, islem, yeni_normal, yeni_fazla, yeni_durum, "
                    "yeni_kaynak_satir, import_batch_id, zaman) VALUES (?,?, 'ekle', ?,?,?,?,?,?)",
                    (kayit["tc"], kayit["tarih"], kayit["normal_saat"], kayit["fazla_saat"], kayit["durum"],
                     kayit["kaynak_satir"], batch_id, now),
                )
                yeni += 1
            elif (round(mevcut["normal_saat"], 2) != kayit["normal_saat"]
                    or round(mevcut["fazla_saat"], 2) != kayit["fazla_saat"]
                    or (mevcut["durum"] or None) != kayit["durum"]):
                conn.execute(
                    "UPDATE gunluk_puantaj SET normal_saat=?, fazla_saat=?, durum=?, "
                    "kaynak_satir=?, import_batch_id=? WHERE tc=? AND tarih=?",
                    (kayit["normal_saat"], kayit["fazla_saat"], kayit["durum"],
                     kayit["kaynak_satir"], batch_id, kayit["tc"], kayit["tarih"]),
                )
                conn.execute(
                    "INSERT INTO puantaj_history (tc, tarih, islem, eski_normal, eski_fazla, eski_durum, "
                    "eski_kaynak_satir, eski_batch_id, yeni_normal, yeni_fazla, yeni_durum, yeni_kaynak_satir, "
                    "import_batch_id, zaman) VALUES (?,?, 'guncelle', ?,?,?,?,?,?,?,?,?,?,?)",
                    (kayit["tc"], kayit["tarih"], mevcut["normal_saat"], mevcut["fazla_saat"], mevcut["durum"],
                     mevcut["kaynak_satir"], mevcut["import_batch_id"], kayit["normal_saat"],
                     kayit["fazla_saat"], kayit["durum"], kayit["kaynak_satir"], batch_id, now),
                )
                guncellenen += 1
            else:
                degismeyen += 1

        engelleyici_sayisi = sum(1 for a in preview["uyarilar"] if a["onem"] == "engelleyici")
        for alert in preview["uyarilar"]:
            conn.execute(
                "INSERT INTO quality_alerts (import_batch_id, tip, onem, tc, ad_soyad, kaynak_satir, "
                "detay, olusturma_zamani) VALUES (?,?,?,?,?,?,?,?)",
                (batch_id, alert["tip"], alert["onem"], alert.get("tc"), alert.get("ad_soyad"),
                 alert.get("kaynak_satir"), alert["detay"], now),
            )

        conn.execute(
            "UPDATE import_batches SET tamamlanma_zamani=?, durum='tamamlandi', taranan_satir=?, "
            "yazilan_gun_kaydi=?, engelleyici_sayisi=?, uyari_sayisi=? WHERE id=?",
            (now, preview.get("taranan_satir", 0), yeni + guncellenen, engelleyici_sayisi,
             len(preview["uyarilar"]), batch_id),
        )
        conn.commit()
        return {"batch_id": batch_id, "yeni_kayit": yeni, "guncellenen_kayit": guncellenen,
                "degismeyen_kayit": degismeyen, "uyari_sayisi": len(preview["uyarilar"]),
                "yedek": yedek_bilgisi["ad"] if yedek_bilgisi else None}
    except Exception as e:
        conn.rollback()
        conn.execute(
            "INSERT INTO import_batches (dosya_adi, dosya_hash, baslama_zamani, tamamlanma_zamani, "
            "durum, hata_mesaji, tur) VALUES (?,?,?,?, 'hata', ?, 'mesai')",
            (preview["dosya_adi"], preview["dosya_hash"], now, now, str(e)),
        )
        conn.commit()
        raise
    finally:
        conn.close()
