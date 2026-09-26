# -*- coding: utf-8 -*-
"""
VERİ MERKEZİ - EKİP LİSTESİ EXCEL ÇIKTISI (veri_merkezi_ekip_excel.py)
=========================================================================
Ekip Listesi'nin Veri Merkezi'ne özel Excel çıktısı. TEK sayfada, yukarıdan
aşağıya üç bölüm:

  1) Personel listesi      (görev gruplarına bölünmüş; Formen > Ekip Başı >
                            diğer görevler > Bayrakçı - team_report.sort_roster)
  2) Göreve göre özet      (görev / personel sayısı / toplam)
  3) Ay içinde çıkış yapan (ad soyad, TC, görev, çıkış tarihi)

Görsel dil team_report ile aynıdır (banner, başlık/satır/toplam stilleri).
team_report.write_roster_report'a (Puantaj Suite'in kullandığı iki sayfalı
çıktı) DOKUNULMAZ.
"""

__version__ = "2026-09-25.1"

import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import puantaj_engine as engine
import team_report as report
import veri_merkezi_ekip_listesi as el
import veri_merkezi_sema as sema

SON_SUTUN = 5          # A..E
ROL_ZEMIN = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
ROL_FONT = Font(bold=True, color=report.NAVY_DARK, size=10)
BOLUM_FONT = Font(bold=True, color="FFFFFF", size=12)
AYLAR = ["OCAK", "ŞUBAT", "MART", "NİSAN", "MAYIS", "HAZİRAN", "TEMMUZ", "AĞUSTOS", "EYLÜL", "EKİM", "KASIM", "ARALIK"]


AY_ADLARI = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım",
             "Aralık"]


def ay_adi(ay):
    """'2026-04' -> 'Nisan 2026' (str.title() Türkçe İ harfini bozduğu için tabloyla)."""
    y, m = ay.split("-")
    return f"{AY_ADLARI[int(m) - 1]} {y}"


def ay_etiketi(ay):
    y, m = ay.split("-")
    return f"{AYLAR[int(m) - 1]} {y}"


def ay_araligi(ay):
    y, m = (int(x) for x in ay.split("-"))
    bas = datetime.date(y, m, 1)
    bit = datetime.date(y + (m == 12), m % 12 + 1, 1) - datetime.timedelta(days=1)
    return bas.isoformat(), bit.isoformat()


def son_aylar(n=12, bugun=None):
    bugun = bugun or datetime.date.today()
    y, m = bugun.year, bugun.month
    out = []
    for _ in range(n):
        out.append(f"{y}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return out


def _tarih(v):
    s = str(v or "").strip()[:10]
    try:
        return datetime.date.fromisoformat(s)
    except ValueError:
        return None


def cikis_yapanlar(ay, grup_filter=None, ekip_filter=None, gorev_filter=None, db_path=None):
    """Seçili kapsamda, verilen ay (YYYY-MM) içinde işten çıkış tarihi olan personel.
    Grup/alt ekip/görev filtreleri Ekip Listesi ile aynı kuralla (boş alan etiketleri dahil) uygulanır."""
    bas, bit = ay_araligi(ay)
    conn = sema.get_connection(db_path)
    try:
        rows = conn.execute(
            "SELECT tc, ad_soyad, grup, bagli_oldugu_ekip, gorevi, isten_cikis_tarihi FROM personnel "
            "WHERE isten_cikis_tarihi IS NOT NULL AND isten_cikis_tarihi <> '' "
            "AND substr(isten_cikis_tarihi, 1, 10) BETWEEN ? AND ?", (bas, bit)).fetchall()
    finally:
        conn.close()
    grup_set = set(grup_filter or [])
    ekip_set = set(ekip_filter or [])
    gorev_set = set(gorev_filter or [])
    out = []
    for r in rows:
        k = {"tc": r["tc"], "name": (r["ad_soyad"] or "").strip(),
             "team": el._bos_ise(r["grup"], el.NO_GRUP), "subteam": el._bos_ise(r["bagli_oldugu_ekip"], el.NO_EKIP),
             "role": el._bos_ise(r["gorevi"], el.NO_GOREV), "cikis": _tarih(r["isten_cikis_tarihi"])}
        if grup_set and k["team"] not in grup_set:
            continue
        if ekip_set and k["subteam"] not in ekip_set:
            continue
        if gorev_set and k["role"] not in gorev_set:
            continue
        out.append(k)
    out.sort(key=lambda k: (k["cikis"] or datetime.date.max, engine.normalize_name(k["name"])))
    return out


def _bolum_basligi(ws, satir, metin):
    ws.merge_cells(start_row=satir, start_column=1, end_row=satir, end_column=SON_SUTUN)
    c = ws.cell(row=satir, column=1, value=metin)
    c.font = BOLUM_FONT
    c.fill = PatternFill(start_color=report.NAVY, end_color=report.NAVY, fill_type="solid")
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[satir].height = 22


def yaz(roster_data, kapsam, cikanlar, ay, output_path):
    """roster_data: el.build_roster() sonucu (gizlenenler çıkarılmış olabilir)."""
    sirali, rol_sayi = report.sort_roster(roster_data["records"])
    if not sirali:
        raise engine.TransferError("Listede kimse kalmadı: seçilen filtrelerle eşleşen personel yok ya da hepsi gizlendi.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ekip Listesi"
    ws.sheet_view.showGridLines = False

    # ---- başlık
    report.banner(ws, f"A1:{get_column_letter(SON_SUTUN)}1", f"{kapsam} EKİP LİSTESİ", height=28)
    alt = f"Toplam {len(sirali)} personel"
    if roster_data.get("excluded_exited"):
        alt += f"   ·   {roster_data['excluded_exited']} ayrılmış personel listeye dahil edilmedi"
    alt += f"   ·   {ay_adi(ay)} ayında çıkış yapan: {len(cikanlar)}"
    c = ws.cell(row=2, column=1, value=alt)
    c.font = report.SUBTITLE_FONT
    ws.merge_cells(f"A2:{get_column_letter(SON_SUTUN)}2")

    # ---- 1) personel listesi
    r = 4
    for col, h in enumerate(["Sıra No", "Ad Soyad", "TC Kimlik No", "Görevi"], start=1):
        ws.cell(row=r, column=col, value=h)
    report.style_header_row(ws, r, 1, 4)
    r += 1
    sira, onceki = 1, None
    for rec in sirali:
        if rec["role"] != onceki:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            hc = ws.cell(row=r, column=1, value=f"{rec['role']}  ({rol_sayi[rec['role']]} kişi)")
            hc.font = ROL_FONT
            hc.alignment = Alignment(vertical="center", indent=1)
            for col in range(1, 5):
                ws.cell(row=r, column=col).border = report.THIN_BORDER
                ws.cell(row=r, column=col).fill = ROL_ZEMIN
            onceki = rec["role"]
            r += 1
        ws.cell(row=r, column=1, value=sira).number_format = report.INT_FMT
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=rec["name"])
        ws.cell(row=r, column=3, value=rec["tc"])
        ws.cell(row=r, column=4, value=rec["role"])
        report.style_data_row(ws, r, 1, 4, zebra=(sira % 2 == 0))
        sira += 1
        r += 1

    # ---- 2) göreve göre özet
    r += 2
    _bolum_basligi(ws, r, "GÖREVE GÖRE PERSONEL SAYISI")
    r += 1
    ws.cell(row=r, column=1, value="Sıra")
    ws.cell(row=r, column=2, value="Görevi")
    ws.cell(row=r, column=3, value="Personel Sayısı")
    report.style_header_row(ws, r, 1, 3)
    r += 1
    roller = []
    for rec in sirali:
        if rec["role"] not in roller:
            roller.append(rec["role"])
    for i, rol in enumerate(roller, start=1):
        ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=rol)
        ws.cell(row=r, column=3, value=rol_sayi[rol]).number_format = report.INT_FMT
        report.style_data_row(ws, r, 1, 3, zebra=(i % 2 == 0))
        r += 1
    ws.cell(row=r, column=2, value="TOPLAM")
    ws.cell(row=r, column=3, value=len(sirali)).number_format = report.INT_FMT
    report.style_total_row(ws, r, 1, 3)
    r += 1

    # ---- 3) ay içinde çıkış yapanlar
    r += 2
    _bolum_basligi(ws, r, f"{ay_etiketi(ay)} AYINDA ÇIKIŞ YAPAN PERSONEL ({len(cikanlar)} kişi)")
    r += 1
    for col, h in enumerate(["Sıra No", "Ad Soyad", "TC Kimlik No", "Görevi", "Çıkış Tarihi"], start=1):
        ws.cell(row=r, column=col, value=h)
    report.style_header_row(ws, r, 1, SON_SUTUN)
    r += 1
    if cikanlar:
        for i, k in enumerate(cikanlar, start=1):
            ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=k["name"])
            ws.cell(row=r, column=3, value=k["tc"])
            ws.cell(row=r, column=4, value=k["role"])
            tc = ws.cell(row=r, column=5, value=k["cikis"])
            tc.number_format = "DD.MM.YYYY"
            tc.alignment = Alignment(horizontal="center")
            report.style_data_row(ws, r, 1, SON_SUTUN, zebra=(i % 2 == 0))
            r += 1
    else:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=SON_SUTUN)
        bc = ws.cell(row=r, column=1, value=f"{ay_adi(ay)} ayında bu kapsamda çıkış yapan personel yok.")
        bc.font = Font(italic=True, color="546E7A")
        bc.alignment = Alignment(horizontal="center")
        report.style_data_row(ws, r, 1, SON_SUTUN)

    for col, w in zip("ABCDE", (10, 30, 18, 30, 14)):
        ws.column_dimensions[col].width = w
    report.setup_print(ws, landscape=False)
    ws.oddHeader.center.text = f"{kapsam} — Ekip Listesi"   # setup_print "Mesai Raporu" yazıyor

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
